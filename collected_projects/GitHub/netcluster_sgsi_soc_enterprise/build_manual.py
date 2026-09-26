# -*- coding: utf-8 -*-
"""
Generador del Manual de Usuario y Operación Oficial SGSI & SOC (SERMIG 2026).
Produce un documento Microsoft Word (.docx) formal, estilizado y estructurado.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_manual(output_path):
    doc = docx.Document()

    # Márgenes de página (2 cm aprox)
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Colores corporativos
    C_NAVY = RGBColor(15, 23, 42)       # Slate 900
    C_CYAN = RGBColor(14, 116, 144)     # Cyan 700
    C_BLUE = RGBColor(29, 78, 216)      # Blue 700
    C_MUTED = RGBColor(71, 85, 105)     # Slate 600
    C_TEXT = RGBColor(30, 41, 59)       # Slate 800

    def set_cell_background(cell, fill_hex):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    # -------------------------------------------------------------
    # PORTADA Y ENCABEZADOS
    # -------------------------------------------------------------
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("REPÚBLICA DE CHILE • SERVICIO NACIONAL DE MIGRACIONES (SERMIG)")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(9.5)
    r_inst.font.bold = True
    r_inst.font.color.rgb = C_CYAN

    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_after = Pt(16)
    r_dept = p_dept.add_run("UNIDAD DE SEGURIDAD DE LA INFORMACIÓN Y CIBERSEGURIDAD (CISO)")
    r_dept.font.name = "Arial"
    r_dept.font.size = Pt(8.5)
    r_dept.font.bold = True
    r_dept.font.color.rgb = C_MUTED

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("MANUAL DE USUARIO Y OPERACIÓN")
    r_title.font.name = "Arial Black"
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = C_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    r_sub = p_sub.add_run("PLATAFORMA INSTITUCIONAL SGSI & CENTRO SOC EN VIVO\nConformidad ISO/IEC 27001:2022, ISO 27005:2022 y Ley N° 21.663")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = C_CYAN

    # Ficha Técnica en Tabla
    tbl_meta = doc.add_table(rows=5, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Versión del Documento:", "3.5 (Edición Integrada OneFirewall CTI, Wazuh SIEM, Lansweeper ITAM & GLPI / GLPI-Agent)"),
        ("Ámbito de Aplicación:", "Infraestructura, Sistemas y Servicios Digitales de SERMIG"),
        ("Marco Normativo:", "ISO/IEC 27001:2022, ISO 27005:2022, Ley Marco N° 21.663, MITRE ATT&CK v15"),
        ("Clasificación de Seguridad:", "USO INTERNO RESTRINGIDO / CONFIDENCIAL"),
        ("Fecha de Actualización:", "Septiembre 2026")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        c0 = tbl_meta.cell(row_idx, 0)
        c1 = tbl_meta.cell(row_idx, 1)
        c0.width = Inches(2.3)
        c1.width = Inches(4.7)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 80, 80, 100, 100)
        set_cell_margins(c1, 80, 80, 100, 100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.name = "Arial"
        r0.font.bold = True
        r0.font.size = Pt(9)
        r0.font.color.rgb = C_NAVY
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = "Arial"
        r1.font.size = Pt(9)
        r1.font.color.rgb = C_TEXT

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # HELPERS DE CONTENIDO
    # -------------------------------------------------------------
    def add_h1(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.name = "Arial Black"
        r.font.size = Pt(13.5)
        r.font.color.rgb = C_NAVY

    def add_h2(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(11.5)
        r.font.color.rgb = C_CYAN

    def add_h3(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = C_BLUE

    def add_p(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rp = p.add_run(bold_prefix)
            rp.font.name = "Arial"
            rp.font.bold = True
            rp.font.size = Pt(9.5)
            rp.font.color.rgb = C_NAVY
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = C_TEXT

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        if bold_prefix:
            rp = p.add_run(bold_prefix)
            rp.font.name = "Arial"
            rp.font.bold = True
            rp.font.size = Pt(9.5)
            rp.font.color.rgb = C_NAVY
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = C_TEXT

    def add_callout(title, text, bg_hex="F0FDF4", border_hex="16A34A"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_t = p.add_run(f"🛡️ {title}\n")
        r_t.font.name = "Arial"
        r_t.font.bold = True
        r_t.font.size = Pt(10)
        r_t.font.color.rgb = C_NAVY
        
        r_b = p.add_run(text)
        r_b.font.name = "Arial"
        r_b.font.size = Pt(9)
        r_b.font.color.rgb = C_TEXT
        
        p_af = doc.add_paragraph()
        p_af.paragraph_format.space_before = Pt(2)
        p_af.paragraph_format.space_after = Pt(2)

    # -------------------------------------------------------------
    # 1. INTRODUCCIÓN Y ARQUITECTURA
    # -------------------------------------------------------------
    add_h1("1. INTRODUCCIÓN Y ARQUITECTURA DEL SISTEMA")
    add_p("La plataforma SGSI & SOC SERMIG constituye el núcleo operativo y analítico para la gestión integral de la ciberseguridad, el cumplimiento normativo internacional ISO/IEC 27001:2022 y la gobernanza de riesgos bajo ISO 27005:2022 y la Ley Marco de Ciberseguridad N° 21.663.")
    
    add_h2("1.1 Modelos de Despliegue Disponibles")
    add_bullet(" Servidor FastAPI ligero, almacenamiento estructurado JSON, WebSocket bidireccional nativo para telemetría en tiempo real y despliegue rápido.", bold_prefix="Edición Web Ligera (sgsi-soc-web):")
    add_bullet(" Arquitectura basada en SQLAlchemy ORM, soporte dual PostgreSQL y SQLite de alta concurrencia, pistas de auditoría inmutables y soporte para clústeres de alta disponibilidad.", bold_prefix="Edición Enterprise (sgsi_soc_enterprise):")

    add_h2("1.2 Matriz de Roles y Control de Acceso (RBAC)")
    add_bullet(" Acceso total irrestricto. Configuración de parámetros institucionales, credenciales CTI, gestión de usuarios, edición de SoA, mitigación de incidentes y aprobación de políticas.", bold_prefix="Administrador CISO:")
    add_bullet(" Operación del Centro SOC, ingesta forense de logs, ejecución del Copiloto IA, consulta de inteligencia OneFirewall, actualización de riesgos y mitigación técnica.", bold_prefix="Analista SOC / Especialista SGSI:")
    add_bullet(" Acceso de solo lectura con capacidades de exportación forense, visualización de pistas de auditoría (Audit Log) y descarga de matrices de cumplimiento para acreditaciones externas.", bold_prefix="Auditor de Seguridad:")
    add_bullet(" Acceso restringido a paneles ejecutivos y tableros de comando sin capacidad de modificación de datos.", bold_prefix="Visor Institucional:")

    # -------------------------------------------------------------
    # 2. TABLERO DE COMANDO CENTRAL (DASHBOARD - /)
    # -------------------------------------------------------------
    add_h1("2. TABLERO DE COMANDO CENTRAL (CISO DASHBOARD - /)")
    add_p("El Tablero Principal condensa la postura de seguridad institucional en una vista unificada que correlaciona controles de gobierno, gestión de riesgos y eventos de ciberseguridad en tiempo real.")

    add_h2("2.1 Indicadores Clave de Rendimiento (KPIs)")
    add_bullet(" Porcentaje promedio ponderado de implementación de los 93 controles de la norma ISO 27001:2022.", bold_prefix="Índice de Madurez SoA:")
    add_bullet(" Porcentaje de riesgos residuales controlados frente al universo de riesgos inherentes evaluados.", bold_prefix="Eficacia de Mitigación de Riesgos:")
    add_bullet(" Tiempo promedio transcurrido desde la ocurrencia del evento hasta su correlación en el SOC (medido en minutos).", bold_prefix="MTTD (Mean Time to Detect):")
    add_bullet(" Tiempo medio de respuesta y contención de incidentes mediante playbooks automáticos.", bold_prefix="MTTR (Mean Time to Respond):")
    add_bullet(" Cuantificación de activos calificados como esenciales o de importancia vital bajo la Ley N° 21.663.", bold_prefix="Activos Críticos SERMIG:")

    add_h2("2.2 Geointeligencia de Amenazas Globales en Tiempo Real")
    add_p("Incorpora un mapa de telemetría geoespacial que geolocaliza el origen de los ataques recibidos contra la infraestructura perimetral de SERMIG, mostrando:")
    add_bullet("Banderas oficiales del país atacante (Alemania, Países Bajos, Rusia, China, Estados Unidos, etc.).")
    add_bullet("Distribución porcentual del tráfico anómalo por cuadrante geográfico.")
    add_bullet("Correlación instantánea de vectores de ataque más recurrentes (Botnet C2, Escaneo de Puertos, Fuerza Bruta).")

    # -------------------------------------------------------------
    # 3. CENTRO SOC EN TIEMPO REAL (/soc)
    # -------------------------------------------------------------
    add_h1("3. CENTRO DE OPERACIONES DE CIBERSEGURIDAD (SOC EN VIVO - /soc)")
    add_p("El Centro SOC es el módulo operativo de primera línea para la detección, correlación, investigación y mitigación de amenazas informáticas.")

    add_h2("3.1 Telemetría y Registro de Incidentes (11 Columnas Completas)")
    add_p("La tabla principal ofrece una visualización densa y altamente legible con encabezados fijos (sticky header) y barra de desplazamiento responsiva:")
    add_bullet("Identificador único alfanumérico (ej: INC-2026-0842).", bold_prefix="1. ID:")
    add_bullet("Marca temporal con precisión de segundos.", bold_prefix="2. Fecha/Hora:")
    add_bullet("Identificación del hallazgo y resumen descriptivo del evento.", bold_prefix="3. Título / Descripción:")
    add_bullet("Insignia de severidad codificada por color (Crítica, Alta, Media, Baja).", bold_prefix="4. Severidad:")
    add_bullet("Dirección IP atacante. Es interactiva: al hacer clic abre automáticamente el inspector de OneFirewall CTI.", bold_prefix="5. IP Origen:")
    add_bullet("Nombre del host o usuario institucional involucrado.", bold_prefix="6. Host Origen:")
    add_bullet("Dirección IP del servidor o servicio interno afectado.", bold_prefix="7. IP Destino:")
    add_bullet("Servidor de base de datos, API o equipo terminal SERMIG.", bold_prefix="8. Host Destino:")
    add_bullet("Código de técnica oficial según matriz MITRE ATT&CK (ej: T1110.001, T1059, T1190).", bold_prefix="9. MITRE ATT&CK:")
    add_bullet("Estado del flujo de trabajo (Abierto, En Investigación, En Mitigación, Cerrado).", bold_prefix="10. Estado:")
    add_bullet("Botonera de acciones rápidas (Copiloto IA, CTI OneFirewall y Mitigación).", bold_prefix="11. Acciones:")

    add_h2("3.2 Ingesta Forense de Archivos de Logs Multiformato")
    add_p("Permite procesar archivos de registro reales (.log, .txt, .csv, .json) de hasta 100 MB. El motor de análisis reconoce automáticamente:")
    add_bullet("FortiGate UTM y FortiOS Traffic/UTM Logs.")
    add_bullet("Linux Syslog (RFC 3164 / RFC 5424 / auth.log / secure).")
    add_bullet("Servidores Web Nginx, Apache HTTP y WAF.")
    add_bullet("Windows Security Event Logs (Event IDs 4624, 4625, 4720, 7045).")
    add_bullet("Correlación automática con más de 1.350 IoCs en memoria.")

    add_h2("3.3 Copiloto Heurístico y Generativo con Inteligencia Artificial")
    add_p("Al presionar el botón 'IA' de cualquier incidente, el copiloto analiza el contexto de la amenaza y entrega en 4 pestañas interactivas:")
    add_bullet("Dictamen ejecutivo del ataque, vector de entrada, nivel de riesgo e impacto institucional.", bold_prefix="Análisis Forense:")
    add_bullet("Mapeo exacto con los controles aplicables de la norma ISO/IEC 27001:2022 (A.5.24, A.8.7, A.8.15, A.8.16).", bold_prefix="Controles ISO 27001:")
    add_bullet("Generación automática de reglas de contención listas para copiar para Fortinet FortiOS, Cisco Firepower/ASA, Linux nftables/iptables y Windows Defender Firewall.", bold_prefix="Reglas de Firewall:")
    add_bullet("Generación del borrador formal de notificación de incidentes para el CSIRT de Gobierno y la Agencia Nacional de Ciberseguridad (ANCI) en estricto apego al Art. 7 de la Ley N° 21.663.", bold_prefix="Notificación CSIRT / ANCI:")

    add_h2("3.4 Integración con OneFirewall CTI (World Crime Feeds)")
    add_p("Conexión bidireccional con el appliance OneFirewall (POV local en 10.100.1.33 o Cloud API):")
    add_bullet("Escala oficial OFA de 0 a 1.000 puntos: Limpio (0-79), Watchlist (80-139), Bloqueo Recomendado (140-249), Bloqueo Inmediato (250-1000).", bold_prefix="Crime Score Estandarizado:")
    add_bullet("Presenta la bandera oficial y el nombre formal del país de procedencia de la IP atacante.", bold_prefix="Banderas de País de Origen:")
    add_bullet("Al consultar una IP maliciosa, permite presionar 'Mitigar y Bloquear IP' para cargar automáticamente el formulario de mitigación del SOC con las reglas perimetrales preconfiguradas.", bold_prefix="Mitigación en 1-Clic:")

    # -------------------------------------------------------------
    # 4. WAZUH SIEM & XDR (/wazuh)
    # -------------------------------------------------------------
    add_h1("4. PLATAFORMA WAZUH OPEN SOURCE SIEM & XDR (/wazuh)")
    add_p("Wazuh proporciona visibilidad integral de seguridad a nivel de endpoints y servidores de la red SERMIG.")
    add_bullet("Supervisión del estado operativo de Wazuh Manager, Indexer y Dashboard.", bold_prefix="Clúster Wazuh:")
    add_bullet("Inventario de agentes instalados en servidores Linux, Windows Server y estaciones de trabajo, reportando IP, versión del agente y estado de sincronización.", bold_prefix="Gestión de Agentes XDR:")
    add_bullet("Clasificación de alertas según severidad (niveles 1 al 15), reglas de decodificación y técnica MITRE.", bold_prefix="Alertas de Seguridad:")
    add_bullet("Monitoreo de Integridad de Archivos (FIM), Auditoría de Configuraciones Seguras (SCA), Detección de Vulnerabilidades CVE y Respuesta Activa (Active Response) para bloqueo automatizado.", bold_prefix="Módulos Forenses:")

    # -------------------------------------------------------------
    # 5. GOBERNANZA ISO 27001, RIESGOS Y ACTIVOS
    # -------------------------------------------------------------
    add_h1("5. GOBERNANZA DE SEGURIDAD (SOA, RIESGOS Y ACTIVOS)")
    
    add_h2("5.1 Declaración de Aplicabilidad SoA ISO/IEC 27001:2022 (/soa)")
    add_p("Administra los 93 controles de seguridad distribuidos en las 4 cláusulas del Anexo A:")
    add_bullet("37 controles de gobernanza, políticas, relaciones con proveedores y gestión de incidentes.", bold_prefix="A.5 Organizacionales:")
    add_bullet("8 controles de seguridad en recursos humanos, teletrabajo y acuerdos de confidencialidad.", bold_prefix="A.6 De Personas:")
    add_bullet("14 controles de perímetro físico, centros de datos y seguridad en oficinas.", bold_prefix="A.7 Físicos:")
    add_bullet("34 controles de criptografía, autenticación, control de accesos, desarrollo seguro y respaldo.", bold_prefix="A.8 Tecnológicos:")
    add_p("Permite actualizar el porcentaje de madurez (0% a 100%), justificación de inclusión/exclusión y adjuntar evidencias documentales auditables.")

    add_h2("5.2 Gestión de Riesgos ISO 27005:2022 (/riesgos)")
    add_p("Matriz de evaluación matricial cuantitativa y cualitativa (Probabilidad x Impacto):")
    add_bullet("Evaluación del escenario de amenaza sin controles aplicados.", bold_prefix="Riesgo Inherente:")
    add_bullet("Ponderación de la eficacia de los controles ISO aplicados (0% a 100%).", bold_prefix="Eficacia de Controles:")
    add_bullet("Cálculo del riesgo remanente. Define si el riesgo es Aceptable, Medio, Alto o Crítico.", bold_prefix="Riesgo Residual:")
    add_bullet("Asignación de responsables CISO, controles mitigantes y fecha límite de tratamiento.", bold_prefix="Plan de Tratamiento:")

    add_h2("5.3 Inventario de Activos de Información y CMDB (/activos)")
    add_p("Registro estructurado de hardware, software, bases de datos y servicios en la nube bajo el marco de la Ley N° 21.663:")
    add_bullet("Evaluación de Confidencialidad, Integridad y Disponibilidad (escala 1 a 5).", bold_prefix="Criticidad CID:")
    add_bullet("Asignación formal del custodio y dueño del activo dentro de SERMIG.", bold_prefix="Propietario / Custodio:")
    add_bullet("Sincronización automatizada con Lansweeper ITAM y GLPI / GLPI-Agent para descubrimiento continuo de red e inventario automatizado.", bold_prefix="Integración ITAM & CMDB:")

    # -------------------------------------------------------------
    # 6. AUDITORÍA, REPORTES Y AJUSTES
    # -------------------------------------------------------------
    add_h1("6. AUDITORÍA, REPORTES, INTEGRACIONES Y AJUSTES")
    
    add_h2("6.1 Pistas de Auditoría Forense (/auditoria)")
    add_p("Registro cronológico e inmutable de todas las acciones ejecutadas en el sistema:")
    add_bullet("Inicios y cierres de sesión exitosos o fallidos con registro de IP de origen.")
    add_bullet("Modificaciones a controles SoA, reevaluaciones de riesgos y cambios de configuración.")
    add_bullet("Mitigaciones de incidentes ejecutadas desde el SOC y exportaciones de datos.")

    add_h2("6.2 Informes Ejecutivos y Exportación de Datos (/reportes)")
    add_p("Módulo de reportería diseñado para comités de auditoría y dirección:")
    add_bullet("Generación de informes ejecutivos en formato Word, PDF y Excel.")
    add_bullet("Descarga directa de matrices completas de Incidentes SOC, SoA y Riesgos en CSV estructurado.")
    add_bullet("Impresión formal del Libro de Guardia Digital del SOC para respaldo legal.")

    add_h2("6.3 Conectores e Integraciones Modulares (/settings)")
    add_p("Panel de administración central para personalización e interoperabilidad externa modular (activable/desactivable según disponibilidad):")
    add_bullet("Nombre de la institución, autoridades CISO, correos de contacto y logotipo corporativo.", bold_prefix="Identidad SERMIG:")
    add_bullet("Endpoint de Ollama (ej: http://localhost:11434) y selección del modelo (llama3.2, mistral, etc.).", bold_prefix="Inteligencia Artificial Local:")
    add_bullet("Interruptor ON/OFF, URL del servidor POV (10.100.1.33), Token JWT y calibración del umbral de Crime Score (0-1000).", bold_prefix="OneFirewall CTI (Control A.5.7):")
    add_bullet("Interruptor ON/OFF, Token GraphQL Cloud / On-Premise, Site ID y botón de sincronización de activos.", bold_prefix="Lansweeper ITAM (Control A.5.9):")
    add_bullet("Interruptor ON/OFF, Endpoint API REST (apirest.php), App-Token y User-Token personal para sincronización bidireccional con agentes GLPI-Agent desplegados en la red institucional.", bold_prefix="GLPI & GLPI-Agent (Control A.5.9):")
    add_bullet("Creación de usuarios, asignación de roles RBAC y rotación segura de contraseñas.", bold_prefix="Usuarios del Sistema:")

    add_h2("6.4 Guía Operativa de Integración con GLPI & GLPI-Agent")
    add_p("GLPI es la solución de código abierto (Open Source) para la gestión integral de activos de TI (ITAM) y CMDB. Proporciona una alternativa libre y altamente personalizable a Lansweeper:")
    add_bullet("En GLPI (Configuración > General > API), habilite el servicio API REST y genere un App-Token institucional.", bold_prefix="Paso 1 - Habilitar API en GLPI:")
    add_bullet("En el perfil de usuario administrador de GLPI (Preferencias > Claves de API), genere un User-Token de acceso.", bold_prefix="Paso 2 - Generar User-Token:")
    add_bullet("En el SGSI (/settings), ingrese la URL del endpoint (ej: http://10.100.1.33/glpi/apirest.php), el App-Token y el User-Token. Presione 'Probar Conexión' para validar latencia y estado.", bold_prefix="Paso 3 - Configurar y Validar en SGSI:")
    add_bullet("Presione 'Sincronizar Activos' en Ajustes o directamente en el botón 'GLPI-Agent' del Inventario de Activos (/activos) para incorporar automáticamente computadores, servidores y switches descubiertos por los agentes hacia el SGSI.", bold_prefix="Paso 4 - Sincronizar:")

    # -------------------------------------------------------------
    # GUARDAR DOCUMENTO
    # -------------------------------------------------------------
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Manual generado exitosamente en: {output_path}")

if __name__ == "__main__":
    out_brain = r"C:\Users\RICARDO.ALFARO\.gemini\antigravity\brain\4c33e225-b58b-48d8-9c4f-e301dc72d475\MANUAL_DE_USUARIO_SGSI_SOC_SERMIG_2026.docx"
    create_manual(out_brain)

    out_ent = r"C:\Users\RICARDO.ALFARO\Documents\sgsi_soc_enterprise\MANUAL_DE_USUARIO_SGSI_SOC_SERMIG_2026.docx"
    create_manual(out_ent)

    out_web = r"C:\Users\RICARDO.ALFARO\Documents\sgsi_soc_web\MANUAL_DE_USUARIO_SGSI_SOC_SERMIG_2026.docx"
    create_manual(out_web)
