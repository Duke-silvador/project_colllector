"""Builds three_statement_model.xlsx: a linked income statement, balance sheet and
cash flow forecast for an illustrative company. Run: python build_model.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

wb = Workbook()
F = "Arial"
BLUE = Font(name=F, color="0000FF"); BLACK = Font(name=F); GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True); TITLE = Font(name=F, bold=True, size=14)
HDR = Font(name=F, bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="1F4E78")
YEL = PatternFill("solid", fgColor="FFFF00")
TOP = Border(top=Side(style="thin")); DBL = Border(top=Side(style="thin"), bottom=Side(style="double"))
NUM = '#,##0;(#,##0);"-"'; PCT = '0.0%;(0.0%);"-"'; DAYS = '0" days"'
YEARS = ["2024A","2025A","2026E","2027E","2028E","2029E","2030E"]
COLS = ["C","D","E","F","G","H","I"]; HIST = COLS[:2]; PROJ = COLS[2:]

def prev(c): return COLS[COLS.index(c)-1]

def setup(ws, title, unit="($ thousands)"):
    ws["B1"] = title; ws["B1"].font = TITLE
    ws["B2"] = f"Illustrative Manufacturing Co. {unit}"; ws["B2"].font = Font(name=F, italic=True, color="595959")
    ws.column_dimensions["A"].width = 2; ws.column_dimensions["B"].width = 38
    for c in COLS: ws.column_dimensions[c].width = 12
    for c, y in zip(COLS, YEARS):
        ws[f"{c}4"] = y; ws[f"{c}4"].font = HDR; ws[f"{c}4"].fill = HFILL
        ws[f"{c}4"].alignment = Alignment(horizontal="right")
    ws["B4"].fill = HFILL; ws.freeze_panes = "C5"

def row(ws, r, label, vals, fmt=NUM, font=None, bold=False, border=None):
    ws[f"B{r}"] = label; ws[f"B{r}"].font = BOLD if bold else BLACK
    for c, v in vals.items():
        cell = ws[f"{c}{r}"]; cell.value = v; cell.number_format = fmt
        if font: f = font
        elif isinstance(v, str) and v.startswith("=") and "!" in v: f = GREEN
        elif isinstance(v, str) and v.startswith("="): f = BLACK
        else: f = BLUE
        cell.font = Font(name=F, color=f.color, bold=bold)
        if border: cell.border = border

# ---------------- historical data (hardcoded, illustrative) ----------------
rev = {"C":10800, "D":12000}; cogs_p = {"C":0.61, "D":0.60}; sga_p = {"C":0.21, "D":0.205}
da = {"C":540, "D":600}; intr = {"C":215, "D":200}; tax_r = 0.25; payout = 0.30
cogs = {c: round(rev[c]*cogs_p[c]) for c in HIST}; sga = {c: round(rev[c]*sga_p[c]) for c in HIST}
ni = {c: (rev[c]-cogs[c]-sga[c]-da[c]-intr[c])*(1-tax_r) for c in HIST}
ar = {"C":1450, "D":1644}; inv = {"C":1080, "D":1184}; ppe = {"C":3700, "D":4000}
ap = {"C":790, "D":888}; debt = {"C":3300, "D":3000}; cs = {"C":1000, "D":1000}
cash24 = 900
re24 = cash24+ar["C"]+inv["C"]+ppe["C"]-ap["C"]-debt["C"]-cs["C"]
re25 = round(re24 + ni["D"]*(1-payout))
cash25 = ap["D"]+debt["D"]+cs["D"]+re25-ar["D"]-inv["D"]-ppe["D"]
cash = {"C":cash24, "D":cash25}; re = {"C":re24, "D":re25}
assert cash25 > 0, cash25

# ---------------- Cover ----------------
cv = wb.active; cv.title = "Cover"
cv.column_dimensions["A"].width = 2; cv.column_dimensions["B"].width = 26; cv.column_dimensions["C"].width = 90
cv["B2"] = "Three-Statement Financial Model"; cv["B2"].font = Font(name=F, bold=True, size=18)
cv["B3"] = "Illustrative Manufacturing Co. | 2024A-2025A historical, 2026E-2030E forecast | $ thousands"
cv["B3"].font = Font(name=F, italic=True, color="595959")
lines = [
 ("Purpose", "Forecasts the income statement, balance sheet and cash flow statement, fully linked, so a change to any assumption flows through all three."),
 ("Data", "Illustrative figures for a fictional company. Not real company data."),
 ("How to use", "Change only the blue cells and yellow cells on the Assumptions sheet. Every other number is a formula."),
 ("Scenario", "Assumptions!C5 switches the revenue case: 1 = Base, 2 = Upside, 3 = Downside."),
 ("Sheets", "Assumptions > Income Statement > Balance Sheet > Cash Flow > Checks"),
 ("Balancing", "Cash on the balance sheet comes from the cash flow statement. The Checks sheet confirms assets = liabilities + equity every year."),
 ("Colour key", "Blue = hardcoded input | Black = formula | Green = link to another sheet | Yellow fill = key assumption to change"),
 ("Author", "Muhammed Anshaj K"),
]
for i, (k, v) in enumerate(lines):
    cv[f"B{5+i}"] = k; cv[f"B{5+i}"].font = BOLD
    cv[f"C{5+i}"] = v; cv[f"C{5+i}"].font = BLACK; cv[f"C{5+i}"].alignment = Alignment(wrap_text=True, vertical="top")
    cv[f"B{5+i}"].alignment = Alignment(vertical="top")

# ---------------- Assumptions ----------------
a = wb.create_sheet("Assumptions"); setup(a, "Assumptions")
a["B5"] = "Scenario selector"; a["B5"].font = BOLD
a["E5"] = "Type 1 = Base, 2 = Upside, 3 = Downside"; a["E5"].font = Font(name=F, italic=True, color="595959")
a["C5"] = 1; a["C5"].font = BLUE; a["C5"].fill = YEL
a["D5"] = '=CHOOSE($C$5,"Base","Upside","Downside")'; a["D5"].font = BLACK
a["B7"] = "Revenue growth"; a["B7"].font = BOLD
g = {"Base":[0.10,0.09,0.08,0.07,0.06], "Upside":[0.14,0.12,0.11,0.10,0.09], "Downside":[0.04,0.03,0.03,0.02,0.02]}
for i, k in enumerate(g):
    row(a, 8+i, f"  {k} case", {c: v for c, v in zip(PROJ, g[k])}, PCT)
row(a, 11, "  Revenue growth used", {c: f"=CHOOSE($C$5,{c}8,{c}9,{c}10)" for c in PROJ}, PCT, bold=True)


# historical growth for reference
a["B13"] = "Operating assumptions"; a["B13"].font = BOLD
proj = lambda vals: {c: v for c, v in zip(PROJ, vals)}
row(a, 14, "  COGS (% of revenue)", proj([0.60,0.595,0.59,0.59,0.59]), PCT)
row(a, 15, "  SG&A (% of revenue)", proj([0.20,0.20,0.195,0.195,0.19]), PCT)
row(a, 16, "  D&A (% of revenue)", proj([0.05]*5), PCT)
row(a, 17, "  Capex (% of revenue)", proj([0.065,0.06,0.06,0.055,0.055]), PCT)
row(a, 18, "  Tax rate", proj([0.25]*5), PCT)
a["B20"] = "Working capital"; a["B20"].font = BOLD
row(a, 21, "  Receivable days (DSO)", proj([50]*5), DAYS)
row(a, 22, "  Inventory days (DIO, on COGS)", proj([60]*5), DAYS)
row(a, 23, "  Payable days (DPO, on COGS)", proj([45]*5), DAYS)
a["B25"] = "Financing"; a["B25"].font = BOLD
row(a, 26, "  Interest rate (on opening debt)", proj([0.07]*5), PCT)
row(a, 27, "  Debt repayment", proj([300]*5), NUM)
row(a, 28, "  Dividend payout (% of net income)", proj([0.30]*5), PCT)
for r in (8,9,10,14,15,17,21,22,23):
    for c in PROJ: a[f"{c}{r}"].fill = YEL
a["B30"] = "Notes: interest is charged on opening debt to avoid a circular reference. All inputs are illustrative."
a["B30"].font = Font(name=F, italic=True, color="595959")
a["C5"].comment = Comment("Type 1, 2 or 3 to switch the revenue scenario.", "Model")

# ---------------- Income Statement ----------------
s = wb.create_sheet("Income Statement"); setup(s, "Income Statement")
A = "Assumptions"
row(s, 5, "Revenue", {**{c: rev[c] for c in HIST}, **{c: f"={prev(c)}5*(1+{A}!{c}11)" for c in PROJ}}, bold=True)
row(s, 6, "  Growth", {"D": "=D5/C5-1", **{c: f"={c}5/{prev(c)}5-1" for c in PROJ}}, PCT)
row(s, 7, "Cost of goods sold", {**{c: -cogs[c] for c in HIST}, **{c: f"=-{c}5*{A}!{c}14" for c in PROJ}})
row(s, 8, "Gross profit", {c: f"={c}5+{c}7" for c in COLS}, bold=True, border=TOP)
row(s, 9, "  Gross margin", {c: f"=IF({c}5=0,0,{c}8/{c}5)" for c in COLS}, PCT)
row(s, 10, "SG&A", {**{c: -sga[c] for c in HIST}, **{c: f"=-{c}5*{A}!{c}15" for c in PROJ}})
row(s, 11, "EBITDA", {c: f"={c}8+{c}10" for c in COLS}, bold=True, border=TOP)
row(s, 12, "  EBITDA margin", {c: f"=IF({c}5=0,0,{c}11/{c}5)" for c in COLS}, PCT)
row(s, 13, "Depreciation & amortisation", {**{c: -da[c] for c in HIST}, **{c: f"=-{c}5*{A}!{c}16" for c in PROJ}})
row(s, 14, "EBIT", {c: f"={c}11+{c}13" for c in COLS}, bold=True, border=TOP)
row(s, 15, "Interest expense", {**{c: -intr[c] for c in HIST}, **{c: f"=-'Balance Sheet'!{prev(c)}17*{A}!{c}26" for c in PROJ}})
row(s, 16, "Profit before tax", {c: f"={c}14+{c}15" for c in COLS}, bold=True, border=TOP)
row(s, 17, "Income tax", {**{c: f"=-MAX(0,{c}16*0.25)" for c in HIST}, **{c: f"=-MAX(0,{c}16*{A}!{c}18)" for c in PROJ}})
row(s, 18, "Net income", {c: f"={c}16+{c}17" for c in COLS}, bold=True, border=DBL)
row(s, 19, "  Net margin", {c: f"=IF({c}5=0,0,{c}18/{c}5)" for c in COLS}, PCT)
s["C17"].comment = Comment("Historical tax rate of 25% (illustrative).", "Model")

# ---------------- Balance Sheet ----------------
b = wb.create_sheet("Balance Sheet"); setup(b, "Balance Sheet")
IS = "'Income Statement'"; CF = "'Cash Flow'"
b["B5"] = "Assets"; b["B5"].font = BOLD
row(b, 6, "  Cash", {**{c: cash[c] for c in HIST}, **{c: f"={CF}!{c}20" for c in PROJ}})
row(b, 7, "  Accounts receivable", {**{c: ar[c] for c in HIST}, **{c: f"={IS}!{c}5*{A}!{c}21/365" for c in PROJ}})
row(b, 8, "  Inventory", {**{c: inv[c] for c in HIST}, **{c: f"=-{IS}!{c}7*{A}!{c}22/365" for c in PROJ}})
row(b, 9, "Total current assets", {c: f"=SUM({c}6:{c}8)" for c in COLS}, bold=True, border=TOP)
row(b, 10, "  Property, plant & equipment (net)", {**{c: ppe[c] for c in HIST}, **{c: f"={prev(c)}10-{CF}!{c}13-{CF}!{c}7" for c in PROJ}})
row(b, 11, "Total assets", {c: f"={c}9+{c}10" for c in COLS}, bold=True, border=DBL)
b["B13"] = "Liabilities & equity"; b["B13"].font = BOLD
row(b, 14, "  Accounts payable", {**{c: ap[c] for c in HIST}, **{c: f"=-{IS}!{c}7*{A}!{c}23/365" for c in PROJ}})
row(b, 15, "Total current liabilities", {c: f"={c}14" for c in COLS}, bold=True, border=TOP)
row(b, 17, "  Long-term debt", {**{c: debt[c] for c in HIST}, **{c: f"=MAX(0,{prev(c)}17-{A}!{c}27)" for c in PROJ}})
row(b, 18, "Total liabilities", {c: f"={c}15+{c}17" for c in COLS}, bold=True, border=TOP)
row(b, 20, "  Common stock", {**{c: cs[c] for c in HIST}, **{c: f"={prev(c)}20" for c in PROJ}})
row(b, 21, "  Retained earnings", {**{c: re[c] for c in HIST}, **{c: f"={prev(c)}21+{IS}!{c}18+{CF}!{c}17" for c in PROJ}})
row(b, 22, "Total equity", {c: f"={c}20+{c}21" for c in COLS}, bold=True, border=TOP)
row(b, 23, "Total liabilities & equity", {c: f"={c}18+{c}22" for c in COLS}, bold=True, border=DBL)
b["C21"].comment = Comment("2024A retained earnings is the balancing figure for the opening balance sheet (illustrative data).", "Model")
b["D6"].comment = Comment("2025A cash is consistent with 2025 net income less a 30% dividend.", "Model")

# ---------------- Cash Flow ----------------
f = wb.create_sheet("Cash Flow"); setup(f, "Cash Flow Statement")
BS = "'Balance Sheet'"
P = PROJ
def frow(r, label, fn, **kw): row(f, r, label, {c: fn(c) for c in P}, **kw)
f["B5"] = "Operating activities"; f["B5"].font = BOLD
frow(6, "  Net income", lambda c: f"={IS}!{c}18")
frow(7, "  Depreciation & amortisation", lambda c: f"=-{IS}!{c}13")
frow(8, "  (Increase) / decrease in receivables", lambda c: f"={BS}!{prev(c)}7-{BS}!{c}7")
frow(9, "  (Increase) / decrease in inventory", lambda c: f"={BS}!{prev(c)}8-{BS}!{c}8")
frow(10, "  Increase / (decrease) in payables", lambda c: f"={BS}!{c}14-{BS}!{prev(c)}14")
frow(11, "Cash from operating activities", lambda c: f"=SUM({c}6:{c}10)", bold=True, border=TOP)
f["B12"] = "Investing activities"; f["B12"].font = BOLD
frow(13, "  Capital expenditure", lambda c: f"=-{IS}!{c}5*{A}!{c}17")
frow(14, "Cash from investing activities", lambda c: f"={c}13", bold=True, border=TOP)
f["B15"] = "Financing activities"; f["B15"].font = BOLD
frow(16, "  Debt repayment", lambda c: f"={BS}!{c}17-{BS}!{prev(c)}17")
frow(17, "  Dividends paid", lambda c: f"=-MAX(0,{IS}!{c}18*{A}!{c}28)")
frow(18, "Cash from financing activities", lambda c: f"={c}16+{c}17", bold=True, border=TOP)
frow(19, "Net change in cash", lambda c: f"={c}11+{c}14+{c}18", bold=True)
frow(20, "Ending cash", lambda c: f"={BS}!{prev(c)}6+{c}19", bold=True, border=DBL)
f.column_dimensions["C"].hidden = True; f.column_dimensions["D"].hidden = True

# ---------------- Checks ----------------
k = wb.create_sheet("Checks"); setup(k, "Model Checks", "(difference should be zero)")
row(k, 5, "Balance sheet: assets less liabilities & equity", {c: f"=ROUND({BS}!{c}11-{BS}!{c}23,2)" for c in COLS})
row(k, 6, "Cash: balance sheet less cash flow ending cash", {c: f"=ROUND({BS}!{c}6-{CF}!{c}20,2)" for c in P})
row(k, 8, "All checks pass?", {"C": f'=IF(SUMPRODUCT(ABS(C5:I6))=0,"OK","ERROR")'}, fmt="General", bold=True)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
wb.save("three_statement_model.xlsx")
print("Saved three_statement_model.xlsx. Open it in Excel (or run a recalculation) to see values.")
