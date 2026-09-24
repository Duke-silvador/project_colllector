"""Consulta masiva de pólizas y exportación, sin escrituras en MySQL."""
from io import BytesIO
import re

from openpyxl import Workbook, load_workbook
from requests import RequestException

import compass
from database import conectar
from franquicias import _formatear_codigo_oficina


def oficinas_para_busqueda():
    connection = conectar()
    try:
        cursor = connection.cursor(dictionary=True)
        try:
            cursor.execute("SELECT office_id, office_number, office_name FROM staging_hub.offices")
            return cursor.fetchall()
        finally:
            cursor.close()
    finally:
        connection.close()


def leer_polizas(data, hoja, encabezado, columna):
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = book[hoja]
        resultado = []
        for numero, row in enumerate(sheet.iter_rows(min_row=encabezado + 1), encabezado + 1):
            if all(c.value is None for c in row):
                continue
            cell = row[columna]
            value = cell.value
            if isinstance(value, (int, float)) and not isinstance(value, bool) and value == int(value):
                value = str(int(value))
                if re.fullmatch(r"0+", cell.number_format):
                    value = value.zfill(len(cell.number_format))
            resultado.append((numero, str(value).strip() if value is not None else ""))
        return resultado
    finally:
        book.close()


def consultar_polizas(filas, oficinas, token, progreso=None):
    mapa = {str(o["office_id"]).strip(): o for o in oficinas}
    cache = {}
    agentes_por_oficina = {}
    resultados = []
    for i, (fila, numero) in enumerate(filas, 1):
        row = {"Fila Excel": fila, "Póliza": numero, "Estado póliza": "",
               "Office ID": "", "Número oficina": "", "Oficina": "",
               "Vigencia desde": "", "Fecha expiración": "", "Resultado": "",
               "Nombre franquicia Compass": "", "ID agente": "", "Nombre agente": "",
               "Correo agente": "", "Teléfono agente": "", "Estado agente": "",
               "Licencia 220 agente": "", "NPN agente": "", "Resultado agente": ""}
        if not numero:
            row["Resultado"] = "Sin número de póliza"
        else:
            if numero not in cache:
                try:
                    cache[numero] = (compass.buscar_poliza(token, numero), None)
                except (RequestException, ValueError, TypeError, KeyError):
                    cache[numero] = (None, "Error al consultar Compass; vuelve a intentar")
            poliza, error = cache[numero]
            if error:
                row["Resultado"] = error
            elif poliza is None:
                row["Resultado"] = "Póliza no encontrada"
            else:
                office_id = str(poliza.get("office_id") or "").strip()
                oficina = mapa.get(office_id)
                row.update({"Estado póliza": str(poliza.get("status_id") or "Sin estado"),
                            "Office ID": office_id,
                            "Vigencia desde": str(poliza.get("effective_date") or ""),
                            "Fecha expiración": _campo(poliza, "expirationdate", "expirydate", "expiration", "expiry"),
                            "Número oficina": (_formatear_codigo_oficina(oficina.get("office_number")) or "") if oficina else "",
                            "Oficina": str(oficina.get("office_name") or "") if oficina else "",
                            "Resultado": "Encontrada" if oficina else "Oficina sin coincidencia"})
                row["Nombre franquicia Compass"] = row["Oficina"]
                agent_id = _campo(poliza, "agentid", "userid")
                row["ID agente"] = agent_id
                if not agent_id:
                    row["Resultado agente"] = "Póliza sin ID de agente"
                elif not office_id:
                    row["Resultado agente"] = "Póliza sin Office ID"
                else:
                    if office_id not in agentes_por_oficina:
                        try:
                            agentes_por_oficina[office_id] = (compass.obtener_agentes_oficina(token, office_id), None)
                        except (RequestException, ValueError, TypeError, KeyError):
                            agentes_por_oficina[office_id] = ([], "Error al consultar agentes de Compass; vuelve a intentar")
                    agentes, error_agente = agentes_por_oficina[office_id]
                    agente = next((a for a in agentes if agent_id in
                                   [_campo(a, clave) for clave in ("userid", "agentid", "id")]), None)
                    if error_agente:
                        row["Resultado agente"] = error_agente
                    elif agente is None:
                        row["Resultado agente"] = "Agente no encontrado en la oficina"
                    else:
                        row["Nombre agente"] = _campo(agente, "agentname", "name", "fullname") or " ".join(
                            filter(None, [_campo(agente, "firstname"), _campo(agente, "lastname")]))
                        row.update({
                            "Correo agente": _campo(agente, "email", "emailaddress", "agentemail"),
                            "Teléfono agente": _campo(agente, "phone", "phonenumber", "telephone", "mobile", "cellphone"),
                            "Estado agente": _campo(agente, "status", "statusid", "agentstatus"),
                            "Licencia 220 agente": _campo(agente, "agentlicence220number", "agentlicense220number", "license220number"),
                            "NPN agente": _campo(agente, "agent220npn", "agent220npnnumber", "npn"),
                        })
                        row.update(_columnas_agente(agente))
                        row["Resultado agente"] = "Encontrado" if row["Nombre agente"] else "Agente sin nombre en Compass"
        resultados.append(row)
        if progreso:
            progreso(i, len(filas))
    # Todas las filas comparten columnas, incluso si la primera póliza no tiene agente.
    columnas = list(dict.fromkeys(k for row in resultados for k in row))
    return [{k: row.get(k, "") for k in columnas} for row in resultados]


def _columnas_agente(datos, ruta="Agente"):
    """Expone cada dato de Compass en una columna, incluidos campos anidados."""
    representados = {
        "id", "userid", "agentid", "officeid",
        "agentname", "name", "fullname", "firstname", "lastname",
        "email", "emailaddress", "agentemail",
        "phone", "phonenumber", "telephone", "mobile", "cellphone",
        "status", "statusid", "agentstatus",
        "agentlicence220number", "agentlicense220number", "license220number",
        "agent220npn", "agent220npnnumber", "npn",
    }
    columnas = {}
    for clave, valor in datos.items():
        if ruta == "Agente" and re.sub(r"[^a-z0-9]", "", str(clave).casefold()) in representados:
            continue
        nombre = f"{ruta} · {clave}"
        if isinstance(valor, dict):
            columnas.update(_columnas_agente(valor, nombre))
        elif isinstance(valor, list):
            for i, elemento in enumerate(valor, 1):
                columnas.update(_columnas_agente({str(i): elemento}, nombre))
        else:
            columnas[nombre] = "" if valor is None else str(valor)
    return columnas


def _campo(datos, *nombres):
    """Reconoce las variantes de nombres usadas por Compass, preservando IDs."""
    campos = {re.sub(r"[^a-z0-9]", "", str(k).casefold()): v for k, v in datos.items()}
    return next((str(campos[n]).strip() for n in nombres
                 if campos.get(n) is not None and str(campos[n]).strip()), "")


def exportar_excel(resultados, columnas=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Pólizas Compass"
    headers = list(columnas) if columnas is not None else list(dict.fromkeys(key for row in resultados for key in row))
    if not headers:
        raise ValueError("Selecciona al menos una columna para exportar.")
    sheet.append(headers)
    for row in resultados:
        sheet.append([row.get(key, "") for key in headers])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(55, max(18, max(len(str(c.value or "")) for c in column) + 2))
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
