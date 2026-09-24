import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from reportes_validacion import ReporteValidacion
from validacion_polizas import validar_polizas


class ReportesTests(unittest.TestCase):
    def test_busqueda_y_estados_combinados(self):
        from streamlit.testing.v1 import AppTest

        libro = Workbook()
        for fila in [("Policy Number", "Cliente"), ("ABC123", "A"), ("XYZ456", "B"), (None, "C")]:
            libro.active.append(fila)
        libro.create_sheet("June").append(["Policy"])
        libro["June"].append(["ABC123"])
        libro.create_sheet("July").append(["Policy"])
        archivo = BytesIO()
        libro.save(archivo)
        archivo.name = "filtros.xlsx"
        archivo.size = len(archivo.getvalue())
        with tempfile.TemporaryDirectory() as carpeta:
            # La página recarga sus módulos locales en cada ejecución (para reflejar
            # cambios de código sin reiniciar Streamlit); eso pisaría este parche si
            # no se neutraliza aquí, ya que reload() vuelve a definir la clase real.
            with patch("streamlit.file_uploader", return_value=archivo), patch(
                "reportes_validacion.ReporteValidacion",
                side_effect=lambda n, t, e: ReporteValidacion(n, t, e, carpeta)
            ), patch("importlib.reload"):
                app = AppTest.from_file(str(Path(__file__).parent / "pages" / "validar_polizas.py"), default_timeout=20).run()
                app.button[0].click().run()
                for estado, cantidad in [("No encontradas", 1), ("Encontradas", 1), ("Todas las filas en rojo", 2)]:
                    app.radio[0].set_value(estado).run()
                    self.assertFalse(app.exception)
                    self.assertEqual(len(app.dataframe[-1].value), cantidad)
                app.text_input[0].input("xyz").run()
                self.assertEqual(app.dataframe[-1].value.iloc[0]["Policy Number"], "XYZ456")
                app.radio[0].set_value("Encontradas").run()
                self.assertEqual(len(app.dataframe[-1].value), 0)

    def test_reporte_con_hojas_y_resultados(self):
        libro = Workbook()
        libro.active.title = "Compass"
        libro.active.append(["Policy Number"])
        libro.active.append(["A"])
        libro.active.append(["B"])
        for nombre in ("June", "Julio"):
            hoja = libro.create_sheet(nombre)
            hoja.append(["Policy"])
            hoja.append(["A"])
        archivo = BytesIO()
        libro.save(archivo)
        with tempfile.TemporaryDirectory() as carpeta:
            reporte = ReporteValidacion("prueba.xlsx", len(archivo.getvalue()), (1, 1, 1), carpeta)
            reporte.leer_hojas(archivo.getvalue())
            _, detalle = validar_polizas(archivo.getvalue(), progreso=reporte.progreso)
            reporte.completar(detalle)
            reporte.guardar()
            guardado = json.loads(reporte.ruta.read_text(encoding="utf-8"))
            self.assertEqual(guardado["archivo"], "prueba.xlsx")
            self.assertEqual(guardado["cantidad_hojas"], 3)
            self.assertEqual([h["nombre"] for h in guardado["hojas"]], ["Compass", "June", "Julio"])
            self.assertEqual(guardado["resumen"]["filas_rojas"], 1)
            self.assertEqual(guardado["estado"], "Completado")
            self.assertTrue(guardado["eventos"])
            self.assertEqual(guardado["detalle"], detalle)

    def test_errores_y_nombres_unicos(self):
        with tempfile.TemporaryDirectory() as carpeta:
            reporte = ReporteValidacion("prueba.xlsx", 0, (1, 1, 1), carpeta)
            otro = ReporteValidacion("prueba.xlsx", 0, (1, 1, 1), carpeta)
            self.assertNotEqual(reporte.ruta, otro.ruta)
            reporte.fallar(ValueError("Falta Policy"))
            reporte.guardar()
            guardado = json.loads(reporte.ruta.read_text(encoding="utf-8"))
            self.assertEqual(guardado["estado"], "Error")
            self.assertEqual(guardado["error"]["mensaje"], "Falta Policy")

    def test_guardar_reintenta_bloqueo_transitorio(self):
        # En Windows, un antivirus/EDR puede negar el acceso (WinError 5) justo al
        # renombrar el .tmp; un reintento corto debe bastar para completarlo.
        with tempfile.TemporaryDirectory() as carpeta:
            reporte = ReporteValidacion("prueba.xlsx", 0, (1,), carpeta)
            original_replace = Path.replace
            llamadas = {"n": 0}

            def replace_falla_una_vez(self, destino):
                llamadas["n"] += 1
                if llamadas["n"] == 1:
                    raise PermissionError(5, "Access is denied")
                return original_replace(self, destino)

            with patch.object(Path, "replace", replace_falla_una_vez), \
                 patch("reportes_validacion.time.sleep"):
                reporte.guardar()
            self.assertEqual(llamadas["n"], 2)
            self.assertTrue(reporte.ruta.exists())

    def test_guardar_agota_reintentos_y_relanza(self):
        with tempfile.TemporaryDirectory() as carpeta:
            reporte = ReporteValidacion("prueba.xlsx", 0, (1,), carpeta)
            with patch.object(Path, "replace", side_effect=PermissionError(5, "Access is denied")), \
                 patch("reportes_validacion.time.sleep"):
                with self.assertRaises(PermissionError):
                    reporte.guardar()

    def test_progreso_no_aborta_la_validacion_si_falla_el_guardado(self):
        with tempfile.TemporaryDirectory() as carpeta:
            reporte = ReporteValidacion("prueba.xlsx", 0, (1,), carpeta)
            with patch.object(ReporteValidacion, "guardar", side_effect=OSError("bloqueado")):
                reporte.progreso("Validando…", 0.5)  # no debe lanzar
            self.assertEqual(reporte.datos["eventos"][-1]["etapa"], "Validando…")

    def test_pagina_guarda_y_ofrece_reporte_de_error(self):
        from streamlit.testing.v1 import AppTest

        archivo = BytesIO(b"contenido invalido")
        archivo.name = "invalido.xlsx"
        archivo.size = len(archivo.getvalue())
        with tempfile.TemporaryDirectory() as carpeta:
            def crear(nombre, tamano, encabezados):
                return ReporteValidacion(nombre, tamano, encabezados, carpeta)

            with patch("streamlit.file_uploader", return_value=archivo), \
                 patch("streamlit.button", return_value=True), \
                 patch("reportes_validacion.ReporteValidacion", side_effect=crear), \
                 patch("importlib.reload"):
                app = AppTest.from_file(str(Path(__file__).parent / "pages" / "validar_polizas.py")).run()
            self.assertFalse(app.exception)
            self.assertTrue(app.error)
            reportes = list(Path(carpeta).glob("*.json"))
            self.assertEqual(len(reportes), 1)
            guardado = json.loads(reportes[0].read_text(encoding="utf-8"))
            self.assertEqual(guardado["estado"], "Error")
            self.assertEqual(guardado["error"]["tipo"], "BadZipFile")
            self.assertIn("validacion_polizas_reporte", app.session_state)
