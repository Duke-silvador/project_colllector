"""Lectura inicial; el mapeo a columnas depende de la tabla de destino."""

from openpyxl import load_workbook
from pypdf import PdfReader


def leer_excel(ruta, hoja=None):
    """Devuelve filas de valores de un archivo .xlsx o .xlsm."""
    workbook = load_workbook(ruta, read_only=True, data_only=True)
    try:
        sheet = workbook[hoja] if hoja else workbook.active
        return list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def leer_pdf(ruta):
    """Devuelve texto por pagina. Los PDF escaneados requieren OCR adicional."""
    return [page.extract_text(extraction_mode="layout") or "" for page in PdfReader(ruta).pages]
