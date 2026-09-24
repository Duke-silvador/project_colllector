import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from fechas_statements import archivos_excel, preparar_fecha


class DateTests(unittest.TestCase):
    def test_only_target_sheet_member_changes_and_other_cells_are_kept(self):
        for anterior in ('08/15/2026', date(2026, 8, 15)):
            with self.subTest(anterior=anterior), TemporaryDirectory() as folder:
                path = Path(folder) / 'report.xlsx'
                book = Workbook()
                book.active['A2'] = 'Statement Date:'
                book.active['C2'] = anterior
                book.active['D4'] = '=SUM(1,2)'
                book.create_sheet('Other')['A1'] = 'unchanged'
                book.save(path)
                blob, celda, _ = preparar_fecha(path, date(2026, 9, 15))
                self.assertEqual(celda, 'C2')
                with ZipFile(path) as before, ZipFile(BytesIO(blob)) as after:
                    changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
                    self.assertEqual(changed, ['xl/worksheets/sheet1.xml'])
                result = load_workbook(BytesIO(blob))
                self.assertEqual(result.active['D4'].value, '=SUM(1,2)')
                self.assertEqual(result['Other']['A1'].value, 'unchanged')
                result.close()

    def test_files_are_only_first_level_and_exclude_excel_lock_files(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'report.xlsx').touch()
            (root / '~$report.xlsx').touch()
            (root / 'nested').mkdir()
            (root / 'nested' / 'other.xlsx').touch()
            self.assertEqual([p.name for p in archivos_excel(folder)], ['report.xlsx'])
