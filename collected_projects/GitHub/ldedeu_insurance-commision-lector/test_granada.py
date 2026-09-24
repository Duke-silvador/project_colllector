import unittest
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch
from openpyxl import Workbook, load_workbook
from granada import HEADERS, preparar_statement, completar_franquicias, guardar_statement, cargar_comisiones_excel, interpretar_fecha
from openpyxl.utils.datetime import WINDOWS_EPOCH, MAC_EPOCH, to_excel


class GranadaTests(unittest.TestCase):
    def test_history_tries_original_then_edition_variants(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.side_effect = [
            [{'code': '0012IA', 'franchise': 'DT120', 'is_master_code': 1}], []]
        rows = preparar_statement(self.data(edition='9'), '2026-09')
        rows[0]['policy_number'] = '0185FL00099940'
        with patch('granada.compass.obtener_token', return_value='token'), \
             patch('granada.compass.buscar_poliza', return_value=None), \
             patch('granada.buscar_franquicia_historica', side_effect=[None, None, 'DTF0073']) as history:
            completar_franquicias(connection, rows)
        self.assertEqual([call.args[1] for call in history.call_args_list],
                         ['0185FL00099940', '0185FL00099940-9', '0185FL00099940 - 9'])
        self.assertEqual(rows[0]['franchise_number'], 'DTF0073')
        self.assertEqual(rows[0]['franchise_number_source'], 'historico')
        self.assertIsNone(rows[0]['code_lookup_alert'])

    def test_master_tries_hyphen_spacing_variants_before_history(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.side_effect = [
            [{'code': '0012IA', 'franchise': 'DT120', 'is_master_code': 1}],
            [{'office_id': 'real', 'office_number': 73, 'office_name': 'Office', 'state': 'FL'}]]
        rows = preparar_statement(self.data(edition='9'), '2026-09')
        rows[0]['policy_number'] = '0185FL00099940'
        with patch('granada.compass.obtener_token', return_value='token'), \
             patch('granada.compass.buscar_poliza', side_effect=[None, {'office_id': 'real'}]) as lookup, \
             patch('granada.buscar_franquicia_historica') as history:
            completar_franquicias(connection, rows)
        self.assertEqual([call.args[1] for call in lookup.call_args_list],
                         ['0185FL00099940-9', '0185FL00099940 - 9'])
        history.assert_not_called()
        self.assertEqual(rows[0]['franchise_number'], 'DTF0073')

    def test_compass_key_cleans_both_fields_and_cache_distinguishes_editions(self):
        from granada import numero_poliza_compass
        self.assertEqual(numero_poliza_compass(' ab-123 / ', ' 0.2 '), 'AB123-02')
        connection = MagicMock()
        connection.cursor.return_value.fetchall.side_effect = [
            [{'code': '0012IA', 'franchise': 'DT120', 'is_master_code': 1}],
            [{'office_id': 'real', 'office_number': 73, 'office_name': 'Office', 'state': 'FL'}]]
        rows = preparar_statement(self.data(), '2026-09')
        rows[0]['policy_number'] = ' p-001 '
        rows.append(dict(rows[0], edition='3'))
        with patch('granada.compass.obtener_token', return_value='token'), \
             patch('granada.compass.buscar_poliza', return_value={'office_id': 'real'}) as lookup:
            completar_franquicias(connection, rows)
        self.assertEqual([call.args[1] for call in lookup.call_args_list], ['P001-2', 'P001-3'])
        self.assertEqual(rows[0]['policy_number'], ' p-001 ')
        self.assertEqual(rows[0]['edition'], '2')

    def test_master_uses_policy_office_and_preserves_statement_agent(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.side_effect = [
            [{'code': '0012IA', 'franchise': 'DT120', 'is_master_code': 1}],
            [{'office_id': 'real', 'office_number': 73, 'office_name': 'Real Office', 'state': 'FL'}]]
        rows = preparar_statement(self.data(), '2026-09')
        rows.append(dict(rows[0]))
        with patch('granada.compass.obtener_token', return_value='token') as login, \
             patch('granada.compass.buscar_poliza', return_value={'office_id': 'real', 'agent_name': 'Wrong Agent'}) as lookup:
            completar_franquicias(connection, rows)
        login.assert_called_once()
        lookup.assert_called_once_with('token', 'P001-2')
        self.assertTrue(all(r['franchise_number'] == 'DTF0073' for r in rows))
        self.assertEqual(rows[0]['franchise_number_source'], 'compass')
        self.assertEqual(rows[0]['agent_name'], 'Statement Agent')

    def test_master_missing_policy_keeps_franchise_blank_and_alerts(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.side_effect = [
            [{'code': '0012IA', 'franchise': 'DT120', 'is_master_code': 1}], []]
        rows = preparar_statement(self.data(), '2026-09')
        rows[0]['franchise_number'] = 'DT120'
        with patch('granada.compass.obtener_token', return_value='token'), \
             patch('granada.compass.buscar_poliza', return_value=None), \
             patch('granada.buscar_franquicia_historica', return_value=None), \
             patch('granada_historico.buscar_alternativas', return_value=None):
            completar_franquicias(connection, rows)
        self.assertIsNone(rows[0]['franchise_number'])
        self.assertTrue(rows[0]['code_lookup_alert'])

    def test_missing_master_policy_resolves_history_once_and_keeps_statement_agent(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.side_effect = [
            [{'code': '0012IA', 'franchise': 'DT120', 'is_master_code': 1}],
            [{'office_id': 'historical-office', 'office_number': 73, 'office_name': 'Office', 'state': 'FL'}]]
        rows = preparar_statement(self.data(), '2026-09')
        rows.append(dict(rows[0]))
        with patch('granada.compass.obtener_token', return_value='token'), \
             patch('granada.compass.buscar_poliza', return_value=None) as compass_lookup, \
             patch('granada.buscar_franquicia_historica', return_value='73') as history:
            completar_franquicias(connection, rows)
        self.assertEqual(compass_lookup.call_count, 4)
        history.assert_called_once_with(connection, 'P001')
        self.assertTrue(all(r['franchise_number'] == 'DTF0073' for r in rows))
        self.assertEqual(rows[0]['franchise_number_source'], 'historico')
        self.assertIsNone(rows[0]['code_lookup_alert'])
        self.assertEqual(rows[0]['agent_name'], 'Statement Agent')

    def test_carrier_registered_for_raw_commission_loading(self):
        from carriers import CARRIERS, cargar_carriers
        self.assertIn('GRANADA', cargar_carriers())
        self.assertEqual(CARRIERS['GRANADA']['table'], 'staging_hub.st_granda_raw')
        self.assertEqual(CARRIERS['GRANADA'].get('commission_business_line'), '')

    def test_english_abbreviated_month_from_actual_statement(self):
        for value in ('15-Aug-26', '15-AUG-26', '15-aug-2026', '15-Aug-26 00:00:00'):
            self.assertEqual(interpretar_fecha(value, WINDOWS_EPOCH), '2026-08-15')
        with self.assertRaises(ValueError):
            interpretar_fecha('31-Feb-26', WINDOWS_EPOCH)

    def test_dates_with_single_digits_and_time_and_excel_serial(self):
        for value in ('8/1/2026', '8/1/2026 12:00:00 AM', '08/01/2026 00:00:00',
                      '2026-08-01T00:00:00', datetime(2026, 8, 1)):
            self.assertEqual(interpretar_fecha(value, WINDOWS_EPOCH), '2026-08-01')
        for epoch in (WINDOWS_EPOCH, MAC_EPOCH):
            self.assertEqual(interpretar_fecha(to_excel(datetime(2026, 8, 1), epoch), epoch), '2026-08-01')
        with self.assertRaises(ValueError):
            interpretar_fecha('invalid date', WINDOWS_EPOCH)

    def test_franchise_receives_statement_rate_without_querying_commission_table(self):
        for premium, amount in (('124.78', '14.97'), ('-100', '-12'), ('0', '0')):
            connection = MagicMock()
            connection.cursor.return_value.fetchall.return_value = [dict(
                policy_number='P1', premium_amount=premium, commission_amount=amount,
                transaction_type='CANCEL', agent_name='Statement Agent', franchise_number='DTF0073')]
            with patch('granada.completar_franquicias'):
                row = cargar_comisiones_excel(connection, {'file_id': 'report.xlsx',
                    'rows': [{'accounting_month': '2026-09-01'}]})[0]
            self.assertEqual(row['franchise_percent'], row['del_toro_percent'])
            self.assertEqual(row['franchise_commission'], Decimal(amount))
            self.assertIsNone(row['commission_alert'])
            self.assertFalse(any('commission_rates' in c.args[0]
                                 for c in connection.cursor.return_value.execute.call_args_list))

    def data(self, tipo='Renewal', edition='2', premium=100, commission=12):
        book = Workbook()
        book.active.title = 'Any sheet name'
        book.active.append(['Statement heading'])
        book.active.append(list(HEADERS))
        book.active.append(['P001', edition, 'Insured', '0012IA', 'Statement Agent', tipo,
                            datetime(2026, 9, 1), premium, commission])
        book.create_sheet('Ignored').append(['Invalid'])
        blob = BytesIO()
        book.save(blob)
        book.close()
        return blob.getvalue()

    def test_types_dates_signed_amounts_and_statement_agent(self):
        for value, expected in [('Renewal', 'RENEWAL'), ('Cancellation\u00a0', 'CANCEL'),
                                ('New Business', 'NEW_BUSINESS'), ('Endorsement', 'ENDORSE'),
                                ('Reinstatement', 'REINSTATE')]:
            row = preparar_statement(self.data(value, premium=-100, commission=-12), '2026-09')[0]
            self.assertEqual(row['transaction_type'], expected)
            self.assertEqual(row['agent_name'], 'Statement Agent')
            self.assertEqual(row['effective_date'], '2026-09-01')
            self.assertEqual(Decimal(row['premium_amount']), Decimal('-100'))
            self.assertEqual(row['source_row'], 3)

    def test_edition_is_kept_without_inferring_transaction(self):
        for edition in ('0', '7'):
            row = preparar_statement(self.data('', edition), '2026-09')[0]
            self.assertIsNone(row['transaction_type'])
            self.assertEqual(row['edition'], edition)

    def test_code_lookup_is_case_insensitive_and_preserves_agent(self):
        rows = preparar_statement(self.data(), '2026-09')
        connection = MagicMock()
        connection.cursor.return_value.fetchall.side_effect = [
            [{'code': '0012ia', 'franchise': 'DTF0073', 'is_master_code': 0}],
            [{'office_id': 'office', 'office_number': 73, 'office_name': 'Office', 'state': 'FL'}]]
        completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['franchise_number'], 'DTF0073')
        self.assertEqual(rows[0]['agent_name'], 'Statement Agent')

    def test_insert_stores_edition_and_agent_and_repeat_is_idempotent(self):
        rows = preparar_statement(self.data(), '2026-09')
        connection = MagicMock()
        connection.cursor.return_value.fetchone.return_value = {'acquired': 1}
        connection.cursor.return_value.fetchall.return_value = []
        with patch('granada.conectar', return_value=connection), patch('granada.completar_franquicias'):
            self.assertEqual(guardar_statement(rows, 'granada.xlsx', b'contenido-granada'), 1)
            connection.cursor.return_value.fetchall.return_value = [
                {k: rows[0][k] for k in ('source_sheet', 'source_row', 'source_data')}]
            self.assertEqual(guardar_statement(rows, 'granada.xlsx', b'contenido-granada'), 0)
        insert = next(c for c in connection.cursor.return_value.execute.call_args_list if c.args[0].startswith('INSERT'))
        self.assertEqual(len(insert.args[1]), 27)
        self.assertEqual(insert.args[1][-2:], ('2', 'Statement Agent'))

    def test_save_uses_content_hash_as_file_id_not_the_file_name(self):
        """Regresion: el mismo statement subido con otro nombre de archivo debe reconocerse
        como el mismo import (mismo file_id, calculado del contenido), no como uno nuevo."""
        rows = preparar_statement(self.data(), '2026-09')
        connection = MagicMock()
        connection.cursor.return_value.fetchone.return_value = {'acquired': 1}
        connection.cursor.return_value.fetchall.return_value = []
        with patch('granada.conectar', return_value=connection), patch('granada.completar_franquicias'):
            guardar_statement(rows, 'nombre-original.xlsx', b'mismo-contenido')
            insert = next(c for c in connection.cursor.return_value.execute.call_args_list if c.args[0].startswith('INSERT'))
            file_id_guardado = insert.args[1][0]

            connection.cursor.return_value.fetchall.return_value = [
                {k: rows[0][k] for k in ('source_sheet', 'source_row', 'source_data')}]
            self.assertEqual(guardar_statement(rows, 'nombre-reenviado (1).xlsx', b'mismo-contenido'), 0)
            select = next(c for c in connection.cursor.return_value.execute.call_args_list if c.args[0].startswith('SELECT source_sheet'))
            self.assertEqual(select.args[1][0], file_id_guardado)
