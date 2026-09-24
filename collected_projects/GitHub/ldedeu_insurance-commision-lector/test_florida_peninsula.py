import unittest
from datetime import date
from io import BytesIO
from unittest.mock import MagicMock, patch
from openpyxl import Workbook
from florida_peninsula import HEADERS, leer_statement, guardar_statement, cargar_comisiones_excel, normalizar_agency_id
from decimal import Decimal
from openpyxl import load_workbook
from conciliacion import generar_excel, LDA_COLUMNS_AMWINS
from openpyxl.utils import get_column_letter


def statement():
    book = Workbook()
    book.remove(book.active)
    for sheet in ('FPI', 'EDI', 'OVH', 'Resumen'):
        ws = book.create_sheet(sheet)
        ws.append(list(HEADERS) + ['Premium', 'Commission'])
        ws.append(['Agency', '000123', 'Insured', date(2026, 8, 1), .12058,
                   100, 12, 'Policy Agency', '0012', 'NEW_BUSINESS',
                   'Unrelated company label', 100, 12])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


class FloridaStatementTests(unittest.TestCase):
    def test_statement_transaction_type_aliases_select_correct_rate(self):
        from reparto_comisiones import calcular_comisiones_raw
        rates = [dict(carrier='FLORIDA PENINSULA', state='FL', transaction_type=kind,
                      business_line='', franchise_percent=percent)
                 for kind, percent in [('NEW_BUSINESS', '.08'), ('RENEWAL', '.06')]]
        for label, expected in [('Renew', '.06'), ('RN', '.06'), ('R', '.06'), ('Renewal', '.06'),
                                ('New', '.08'), ('New Business', '.08'), ('New Bussines', '.08'),
                                ('NEW_BUSINESS', '.08'), ('  new   business ', '.08')]:
            with self.subTest(label=label):
                row = dict(leer_statement(statement(), '2026-08')[0], transaction_type=label, state='FL')
                result = calcular_comisiones_raw([row], rates, 'FLORIDA PENINSULA')[0]
                self.assertIsNone(result['commission_alert'])
                self.assertEqual(result['franchise_percent'], Decimal(expected))
                self.assertEqual(result['transaction_type'], label)

    def test_compass_dates_are_converted_for_sql_without_changing_day(self):
        from florida_peninsula import _db_fecha
        from datetime import datetime
        for value in ('2026-08-10T00:00:00.000Z', '2026-08-10',
                      '2026-08-10T23:30:00-05:00', datetime(2026, 8, 10), date(2026, 8, 10)):
            self.assertEqual(_db_fecha(value), date(2026, 8, 10))
        self.assertIsNone(_db_fecha(None))
        self.assertIsNone(_db_fecha(''))
        with self.assertRaises(ValueError):
            _db_fecha('invalid')

    def test_agency_id_suffixes_and_leading_zeros_are_removed(self):
        for value in ('0045595_FPI', '0045595_EDI', '0045595_OVH', '0045595', '45595'):
            with self.subTest(value=value):
                self.assertEqual(normalizar_agency_id(value), '45595')

    def test_normalized_producer_code_preserves_original_source(self):
        row = leer_statement(statement(), '2026-08')[0]
        self.assertEqual(row['producer_code'], '12')
        self.assertEqual(row['source_data']['Agencyid'], '0012')

    def test_loader_uses_statement_type_and_rate_without_calling_type_bot(self):
        original = dict(leer_statement(statement(), '2026-08')[0], id=1, state='FL', office_number='82')
        original['transaction_type'] = 'N'
        c = MagicMock()
        history = [dict(policy_number='000123', carrier='EDISON', transaction_type='RENEWAL',
                        franchise_percent='.99')]
        rates = [dict(carrier='FLORIDA PENINSULA', state='FL', transaction_type='NEW_BUSINESS',
                      business_line='', franchise_percent='.070123')]
        c.cursor.return_value.fetchall.side_effect = [[original], rates]
        with patch('florida_peninsula.obtener_tipos_lote') as bot, patch('florida_peninsula.completar_franquicias'):
            rows = cargar_comisiones_excel(c, dict(file_id='test.xlsx', rows=[dict(accounting_month='2026-08-01')]),
                                          consultar_tipos=True)
        bot.assert_not_called()
        self.assertEqual(rows[0]['transaction_type'], 'N')
        self.assertEqual(rows[0]['del_toro_percent'], Decimal('.12058'))
        self.assertEqual(rows[0]['del_toro_commission'], Decimal('12'))
        self.assertEqual(rows[0]['franchise_commission'], Decimal('7.0123'))
        self.assertFalse(any('historic_data_commissions' in call.args[0] for call in c.cursor.return_value.execute.call_args_list))

    def test_all_three_carriers_use_florida_commission_rates(self):
        originals = [dict(r, id=i, state='FL') for i, r in enumerate(leer_statement(statement(), '2026-08'), 1)]
        rates = [dict(carrier='FLORIDA PENINSULA', state='FL', transaction_type='NEW_BUSINESS', business_line='', franchise_percent='.08')]
        c = MagicMock()
        c.cursor.return_value.fetchall.side_effect = [originals, rates]
        with patch('florida_peninsula.completar_franquicias'):
            rows = cargar_comisiones_excel(c, dict(file_id='test.xlsx', rows=[dict(accounting_month='2026-08-01')]))
        self.assertEqual([r['carrier'] for r in rows], ['FLORIDA PENINSULA', 'EDISON', 'OVATION HOME'])
        self.assertTrue(all(r['franchise_percent'] == Decimal('.08') and r['franchise_commission'] == Decimal('8') for r in rows))

    def test_registered_franchise_code_avoids_compass_and_history(self):
        from florida_peninsula import completar_franquicias
        rows = leer_statement(statement(), '2026-08')
        c = MagicMock()
        c.cursor.return_value.fetchall.return_value = [dict(code='12', franchise='DTF0082', state_code='FL', is_master_code=0)]
        with patch('florida_peninsula.cargar_mapa_office_numbers', return_value={'uid': '82'}), patch('florida_peninsula.compass.obtener_token') as login, patch('florida_peninsula.buscar_franquicia_historica') as historic:
            completar_franquicias(c, rows)
        login.assert_not_called()
        historic.assert_not_called()
        self.assertTrue(all(r['franchise_number'] == 'DTF0082' and r['franchise_number_source'] == 'codigos' for r in rows))

    def test_export_preserves_statement_rate_even_when_ratio_differs(self):
        rows = leer_statement(statement(), '2026-08')
        bank = Workbook()
        bank.active.title = 'Statement'
        bank.active.append(['Description', 'Amount'])
        bank.active.append(['FLORIDA PENINSULA', 36])
        buf = BytesIO()
        bank.save(buf)
        result = generar_excel(rows, Decimal('36'), buf.getvalue(), 'Statement', 1, 2, 'FLORIDA PENINSULA', {})
        formulas = load_workbook(BytesIO(result))
        cached = load_workbook(BytesIO(result), data_only=True)
        def col(field):
            return get_column_letter(next(i for i, (_, f) in enumerate(LDA_COLUMNS_AMWINS, 1) if f == field))
        toro_pct, toro_com, empresa = col('del_toro_percent'), col('del_toro_commission'), col('carrier')
        try:
            self.assertEqual(Decimal(str(formulas['Data']['N2'].value)), Decimal('.12058'))
            self.assertEqual(formulas['Report_LDA'][toro_pct + '2'].value, '=Data!N2')
            self.assertEqual(cached['Report_LDA'][toro_pct + '2'].value, .12058)
            self.assertEqual(cached['Report_LDA'][toro_com + '2'].value, 12)
            self.assertEqual(cached['Report_LDA'][empresa + '3'].value, 'EDISON')
            totals = {cached['Pivot'].cell(i, 1).value: cached['Pivot'].cell(i, 2).value for i in range(1, cached['Pivot'].max_row + 1)}
            for name in ('FLORIDA PENINSULA', 'OVATION HOME', 'EDISON'):
                self.assertEqual(totals[f'Total {name}'], 12)
            self.assertEqual(totals['Grand Total'], 36)
            self.assertTrue(any(str(cell.value).startswith('=SUMIFS(') for row in formulas['Pivot'] for cell in row))
        finally:
            formulas.close()
            cached.close()

    def test_import_saves_all_three_carriers_atomically(self):
        rows = leer_statement(statement(), '2026-08')
        c = MagicMock()
        c.cursor.return_value.fetchone.return_value = {'acquired': 1}
        c.cursor.return_value.fetchall.return_value = []
        with patch('florida_peninsula.conectar', return_value=c), patch('florida_peninsula.completar_franquicias'):
            self.assertEqual(guardar_statement(rows, 'statement.xlsx', b'contenido-fpi'), 3)
        inserts = [call for call in c.cursor.return_value.execute.call_args_list
                   if call.args[0].startswith('INSERT')]
        self.assertEqual([call.args[1][3] for call in inserts], ['FLORIDA PENINSULA', 'EDISON', 'OVATION HOME'])
        self.assertTrue(all(call.args[0].count('%s') == len(call.args[1]) for call in inserts))
        c.commit.assert_called_once()

    def test_identical_retry_does_not_duplicate_and_changed_source_rolls_back(self):
        rows = leer_statement(statement(), '2026-08')
        c = MagicMock()
        c.cursor.return_value.fetchone.return_value = {'acquired': 1}
        existing = [{k: r[k] for k in ('source_sheet', 'source_row', 'source_data')} for r in rows]
        c.cursor.return_value.fetchall.return_value = existing
        with patch('florida_peninsula.conectar', return_value=c), patch('florida_peninsula.completar_franquicias'):
            self.assertEqual(guardar_statement(rows, 'statement.xlsx', b'contenido-fpi'), 0)
            existing[0] = {**existing[0], 'source_data': {'Policy': 'changed'}}
            with self.assertRaisesRegex(ValueError, 'difiere'):
                guardar_statement(rows, 'statement.xlsx', b'contenido-fpi')
        c.rollback.assert_called_once()

    def test_save_uses_content_hash_as_file_id_not_the_file_name(self):
        """Regresion: el mismo statement subido con otro nombre de archivo debe reconocerse
        como el mismo import (mismo file_id, calculado del contenido), no como uno nuevo."""
        rows = leer_statement(statement(), '2026-08')
        c = MagicMock()
        c.cursor.return_value.fetchone.return_value = {'acquired': 1}
        c.cursor.return_value.fetchall.return_value = []
        with patch('florida_peninsula.conectar', return_value=c), patch('florida_peninsula.completar_franquicias'):
            guardar_statement(rows, 'nombre-original.xlsx', b'mismo-contenido')
            insert = next(call for call in c.cursor.return_value.execute.call_args_list
                         if call.args[0].startswith('INSERT'))
            file_id_guardado = insert.args[1][0]

            existing = [{k: r[k] for k in ('source_sheet', 'source_row', 'source_data')} for r in rows]
            c.cursor.return_value.fetchall.return_value = existing
            self.assertEqual(guardar_statement(rows, 'nombre-reenviado (1).xlsx', b'mismo-contenido'), 0)
            select = next(call for call in c.cursor.return_value.execute.call_args_list
                         if call.args[0].startswith('SELECT source_sheet'))
            self.assertEqual(select.args[1][0], file_id_guardado)

    def test_three_sheets_determine_carrier_and_preserve_statement_values(self):
        rows = leer_statement(statement(), '2026-08', premium_column='Premium', commission_column='Commission')
        self.assertEqual([r['carrier'] for r in rows], ['FLORIDA PENINSULA', 'EDISON', 'OVATION HOME'])
        self.assertEqual(rows[0]['del_toro_percent'], '0.12058')
        self.assertEqual(rows[0]['transaction_type_source'], 'statement')
        self.assertEqual(rows[0]['policy_number'], '000123')
        self.assertEqual(rows[0]['source_data']['sumtier'], '100')
        self.assertEqual(rows[0]['premium_amount'], '100')

    def test_unknown_amount_columns_are_never_guessed(self):
        row = leer_statement(statement(), '2026-08', premium_column=None, commission_column=None)[0]
        self.assertIsNone(row['premium_amount'])
        self.assertIsNone(row['commission_amount'])

    def test_missing_explicit_column_blocks_instead_of_using_other_value(self):
        with self.assertRaisesRegex(ValueError, 'faltan columnas'):
            leer_statement(statement(), '2026-08', premium_column='Not present')
