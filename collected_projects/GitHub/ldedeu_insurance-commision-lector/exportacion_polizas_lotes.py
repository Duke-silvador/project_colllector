"""Excel de resultados completos y hojas por lote."""
from io import BytesIO
from openpyxl import load_workbook
from busqueda_polizas import exportar_excel


def exportar_excel_lotes(resultados, tamano_lote=10):
    if tamano_lote < 1:
        raise ValueError('El tamaño del lote debe ser mayor que cero.')
    book = load_workbook(BytesIO(exportar_excel(resultados)))
    completa = book.active
    completa.title = 'All policies'
    headers = [c.value for c in completa[1]]
    for inicio in range(0, len(resultados), tamano_lote):
        sheet = book.create_sheet(f'Batch {inicio // tamano_lote + 1:04d}')
        sheet.append(headers)
        for row in resultados[inicio:inicio + tamano_lote]:
            sheet.append([row.get(k, '') for k in headers])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for key, dimension in completa.column_dimensions.items():
            sheet.column_dimensions[key].width = dimension.width
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
