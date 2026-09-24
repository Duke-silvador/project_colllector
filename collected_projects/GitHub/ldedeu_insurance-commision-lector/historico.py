"""Carga historica en segundo plano, lectura incremental y transaccion atomica."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from threading import Thread, Lock
from uuid import uuid4
import json
import sqlite3
import re

from openpyxl import load_workbook
from database import conectar

FIELDS = {
    'Date': ('date', 'date', None), 'Insured Name': ('insured_name', 'text', 250),
    'Agency Code': ('agency_code', 'text', 150), 'Transaction': ('transaction_type', 'text', 250),
    'Office': ('franchise', 'text', 150), 'Premium': ('premium', 'decimal', None),
    'Del Toro %': ('del_toro_percentage', 'optional_decimal', 10), 'Del Toro $': ('del_toro_commission', 'decimal', None),
    'Policy': ('policy_number', 'text', 200), 'Company': ('carrier', 'required', 200),
    'Franchise %': ('franchise_percentage', 'optional_decimal', 10), 'Franchise $': ('franchise_commission', 'decimal', None),
    'Month': ('statement_month', 'date', None), 'Report Month': ('report_month', 'required_date', None),
    'LOB': ('lob', 'text', 150), 'Agent': ('agent_name', 'text', 250),
    'Term Length': ('term_length', 'text', 50), 'Agency Code + Office': ('agency_code_office', 'text', 200),
    'State': ('state', 'text', 50),
}
# Conservar el historial al recargar el modulo cuando no hay trabajos activos.
JOBS = globals().get('JOBS', {})
GUARD = globals().get('GUARD', Lock())


def clave_registro(values):
    normalized = []
    for value in values:
        if isinstance(value, Decimal):
            normalized.append(format(value.normalize(), 'f') if value else '0')
        elif isinstance(value, (date, datetime)):
            normalized.append(value.isoformat())
        else:
            normalized.append(value)
    return json.dumps(normalized, ensure_ascii=False, separators=(',', ':'))


def convertir(value, kind, limit):
    empty = value is None or isinstance(value, str) and not value.strip()
    if empty:
        return None
    if kind in ('decimal', 'optional_decimal'):
        precision = limit or 18
        maximum = Decimal(10) ** (precision - 5)
        try:
            raw = str(value).strip()
            parentheses = raw.startswith('(') and raw.endswith(')')
            if parentheses:
                raw = raw[1:-1].strip()
            raw = raw.replace('\u2212', '-')
            if '$' in raw or ',' in raw:
                # Formato monetario estadounidense; no adivinar comas decimales.
                match = re.fullmatch(r'([+-]?)\s*\$?\s*([+-]?)(\d{1,3}(?:,\d{3})+|\d+)(\.\d+)?', raw)
                if not match or match[1] and match[2]:
                    raise InvalidOperation
                raw = (match[1] or match[2]) + match[3].replace(',', '') + (match[4] or '')
            result = Decimal(raw)
            if parentheses:
                result = -abs(result)
            if not result.is_finite() or abs(result) >= maximum:
                raise ValueError(f'fuera del rango DECIMAL({precision},5)')
            result = result.quantize(Decimal('.00001'), rounding=ROUND_HALF_UP)
            if abs(result) >= maximum:
                raise ValueError(f'fuera del rango DECIMAL({precision},5) después de redondear')
            return result
        except InvalidOperation:
            raise ValueError(f'valor numérico no reconocido: {str(value)[:120]!r}. Se admite 1234.56, $1,234.56 y (1,234.56).') from None
    if kind in ('date', 'required_date'):
        if isinstance(value, datetime):
            return value.date() if value.year >= 1000 else None
        if isinstance(value, date):
            return value if value.year >= 1000 else None
        for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%Y-%m'):
            try:
                parsed = datetime.strptime(str(value).strip(), fmt).date()
                return parsed if parsed.year >= 1000 else None
            except ValueError:
                pass
        return None  # Las columnas DATE no admiten el texto original invalido.
    result = str(value)
    if len(result) > limit:
        raise ValueError(f'excede {limit} caracteres')
    return result


def importar_historico(path, sheet, header=1, batch_size=1000, progress=lambda **kw: None):
    book = load_workbook(path, read_only=True, data_only=True)
    connection = None
    cursor = None
    locked = False
    count = 0
    processed = 0
    skipped = 0
    seen = sqlite3.connect('')  # indice temporal en disco, no una copia en memoria del historico
    seen.execute('CREATE TABLE seen (value TEXT PRIMARY KEY) WITHOUT ROWID')
    try:
        ws = book[sheet]
        iterator = ws.iter_rows(min_row=header, values_only=True)
        labels = [str(value or '').strip().casefold() for value in next(iterator, ())]
        indexes = []
        for label in FIELDS:
            if labels.count(label.casefold()) != 1:
                raise ValueError(f'Falta el encabezado {label} o está duplicado.')
            indexes.append(labels.index(label.casefold()))
        connection = conectar()
        connection.raise_on_warnings = True
        cursor = connection.cursor()
        cursor.execute("SELECT GET_LOCK('staging_hub.historic_data_commissions.import', 0)")
        locked = cursor.fetchone()[0] == 1
        if not locked:
            raise ValueError('Ya hay una importación histórica en curso.')
        cursor.execute("SELECT ENGINE FROM information_schema.TABLES WHERE TABLE_SCHEMA='staging_hub' AND TABLE_NAME='historic_data_commissions'")
        engine = cursor.fetchone()
        if not engine or engine[0].upper() != 'INNODB':
            raise ValueError('La tabla debe existir y usar InnoDB para revertir errores.')
        columns = ', '.join('`' + spec[0] + '`' for spec in FIELDS.values())
        cursor.execute('SELECT ' + columns + ' FROM staging_hub.historic_data_commissions')
        while existing := cursor.fetchmany(batch_size):
            seen.executemany('INSERT OR IGNORE INTO seen VALUES (?)',
                             ((clave_registro(values),) for values in existing))
            progress(rows=0, skipped=0, status='running')
        seen.commit()
        sql = 'INSERT INTO staging_hub.historic_data_commissions (' + ', '.join('`' + spec[0] + '`' for spec in FIELDS.values()) + ') VALUES (' + ', '.join(['%s'] * len(FIELDS)) + ')'
        batch = []
        for row_number, row in enumerate(iterator, header + 1):
            if all(value is None or value == '' for value in row):
                continue
            values = []
            for index, (label, (_, kind, limit)) in zip(indexes, FIELDS.items()):
                try:
                    values.append(convertir(row[index] if index < len(row) else None, kind, limit))
                except ValueError as exc:
                    raise ValueError(f'Fila Excel {row_number}, columna {label}: {exc}') from None
            processed += 1
            added = seen.execute('INSERT OR IGNORE INTO seen VALUES (?)', (clave_registro(values),)).rowcount
            if not added:
                skipped += 1
                if processed % batch_size == 0:
                    progress(rows=count, skipped=skipped, processed=processed, excel_row=row_number, status='running')
                continue
            batch.append(tuple(values))
            if len(batch) == batch_size:
                cursor.executemany(sql, batch)
                count += len(batch)
                batch.clear()
                seen.commit()
                progress(rows=count, skipped=skipped, processed=processed, excel_row=row_number, status='running')
        if batch:
            cursor.executemany(sql, batch)
            count += len(batch)
        if not processed:
            raise ValueError('La hoja no contiene registros.')
        progress(rows=count, skipped=skipped, processed=processed, status='committing')
        connection.commit()
        return count
    except Exception:
        if connection:
            connection.rollback()
        raise
    finally:
        book.close()
        seen.close()
        if cursor:
            try:
                if locked:
                    cursor.execute("SELECT RELEASE_LOCK('staging_hub.historic_data_commissions.import')")
                    cursor.fetchone()
            finally:
                cursor.close()
        if connection:
            connection.close()


def iniciar(path, sheet, header, batch_size):
    job_id = uuid4().hex
    def update(**values):
        with GUARD:
            JOBS[job_id].update(values)
    with GUARD:
        JOBS[job_id] = dict(status='running', rows=0)
    def worker():
        try:
            count = importar_historico(path, sheet, header, batch_size, update)
            update(status='complete', rows=count)
        except Exception as exc:
            update(status='error', error=str(exc))
        finally:
            Path(path).unlink(missing_ok=True)
    Thread(target=worker, daemon=True).start()
    return job_id


def estado(job_id):
    with GUARD:
        return dict(JOBS.get(job_id, {'status': 'missing'}))
