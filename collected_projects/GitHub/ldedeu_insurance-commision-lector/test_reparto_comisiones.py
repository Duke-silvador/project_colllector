import unittest
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch
from openpyxl import load_workbook
from reparto_comisiones import calcular_comisiones_raw, cargar_comisiones_excel, completar_datos_historicos
from conciliacion import generar_excel
from test_conciliacion import bank_bytes


TARIFAS = [dict(state='TX', carrier='COMMONWEALTH', transaction_type=tipo, business_line='',
                del_toro_percent=Decimal(toro), franchise_percent=Decimal(franchise))
           for tipo, toro, franchise in [('NEW_BUSINESS', '.12', '.07'), ('RENEWAL', '.10', '.05')]]


class RepartoTests(unittest.TestCase):
    def test_all_transactions_uniform_output_accepts_new_statement_percent(self):
        rates = [dict(state='TX', carrier='COMMONWEALTH', transaction_type='ALL', business_line='',
                      del_toro_percent=Decimal(toro), franchise_percent=Decimal('.08'))
                 for toro in ('.10', '.12')]
        for tipo in ('NEW_BUSINESS', 'RENEWAL', 'CANCEL', 'ENDORSE', 'REINSTATE'):
            original = dict(self.row(tipo), commission_amount='150')
            row = calcular_comisiones_raw([original], rates, 'COMMONWEALTH', 'TX', '')[0]
            self.assertEqual(row['del_toro_percent'], Decimal('.15'))
            self.assertEqual(row['franchise_commission'], Decimal('80'))
            self.assertIsNone(row['commission_alert'])

    def test_variants_choose_matching_received_rate_or_report_ambiguity(self):
        rates = [dict(state='TX', carrier='COMMONWEALTH', transaction_type='ALL', business_line='',
                      del_toro_percent=Decimal(toro), franchise_percent=Decimal(franchise))
                 for toro, franchise in (('.10', '.08'), ('.12', '.09'))]
        for amount, expected in (('100', '80'), ('120', '90')):
            row = calcular_comisiones_raw([dict(self.row('CANCEL'), commission_amount=amount)],
                                         rates, 'COMMONWEALTH', 'TX', '')[0]
            self.assertEqual(row['franchise_commission'], Decimal(expected))
        row = calcular_comisiones_raw([dict(self.row('CANCEL'), commission_amount='150')],
                                     rates, 'COMMONWEALTH', 'TX', '')[0]
        self.assertTrue(row['commission_alert'])
        self.assertIsNone(row['franchise_commission'])

    def test_statement_ratio_is_not_quantized_and_franchise_does_not_require_saved_toro(self):
        original = dict(self.row(premium='124.78'), commission_amount='14.97')
        history = [dict(policy_number='P1', carrier='COMMONWEALTH',
                        del_toro_percent=None, franchise_percent='0.08')]
        row = calcular_comisiones_raw([original], [], 'COMMONWEALTH', historico=history)[0]
        self.assertEqual(row['del_toro_percent'], Decimal('14.97') / Decimal('124.78'))
        self.assertEqual(row['del_toro_commission'], Decimal('14.97'))
        self.assertEqual(row['franchise_percent'], Decimal('0.08'))
        self.assertEqual(row['franchise_commission'], Decimal('9.9824'))
        self.assertIsNone(row['commission_alert'])

    def test_zero_premium_has_no_invented_rate(self):
        row = calcular_comisiones_raw([dict(self.row(premium='0'), commission_amount='0')],
                                      TARIFAS, 'COMMONWEALTH', 'TX')[0]
        self.assertEqual(row['del_toro_percent'], Decimal(0))
        with self.assertRaises(ValueError):
            calcular_comisiones_raw([self.row(premium='0')], TARIFAS, 'COMMONWEALTH', 'TX')

    def test_del_toro_uses_statement_instead_of_configured_rate(self):
        tarifas = [dict(TARIFAS[0], del_toro_percent=Decimal('0.123456'))]
        original = dict(self.row(), comm_percent='0.11', del_toro_percent='0.11')
        rows = calcular_comisiones_raw([original], tarifas, 'COMMONWEALTH', 'TX', '')
        self.assertEqual(rows[0]['del_toro_percent'], Decimal('0.06485'))
        data = generar_excel(rows, Decimal('64.85'), bank_bytes(), 'Statement', 1, 2, 'COMMONWEALTH', {})
        book = load_workbook(BytesIO(data), data_only=True)
        try:
            ws = book['Report_LDA']
            cols = {cell.value: cell.column for cell in ws[1]}
            self.assertEqual(Decimal(str(ws.cell(2, cols['Del Toro %']).value)), Decimal('0.06485'))
        finally:
            book.close()

    def test_history_percentages_keep_all_digits_in_final_excel(self):
        history = [dict(policy_number='P1', carrier='COMMONWEALTH',
                        del_toro_percent=Decimal('0.12058'), franchise_percent=Decimal('0.070123'))]
        rows = calcular_comisiones_raw([self.row()], TARIFAS, 'COMMONWEALTH', 'TX', '', historico=history)
        self.assertEqual(rows[0]['del_toro_percent'], Decimal('0.06485'))
        self.assertEqual(rows[0]['franchise_percent'], Decimal('0.070123'))
        data = generar_excel(rows, Decimal('64.85'), bank_bytes(), 'Statement', 1, 2, 'COMMONWEALTH', {})
        book = load_workbook(BytesIO(data), data_only=True)
        try:
            ws = book['Data']
            cols = {cell.value: cell.column for cell in ws[1]}
            for label, expected in (('Del Toro %', '0.06485'), ('Franchise %', '0.070123')):
                sheet = ws if label == 'Del Toro %' else book['Report_LDA']
                headers = {c.value: c.column for c in sheet[1]}
                cell = sheet.cell(2, headers[label])
                self.assertEqual(Decimal(str(cell.value)), Decimal(expected))
                self.assertEqual(cell.number_format, '0.###' if label == 'Del Toro %' else '0.###')
            lda = book['Report_LDA']
            cols = {cell.value: cell.column for cell in lda[1]}
            cell = lda.cell(2, cols['Del Toro %'])
            self.assertEqual(Decimal(str(cell.value)), Decimal('0.06485'))
            self.assertEqual(cell.number_format, '0.###')
        finally:
            book.close()

    def test_loader_recovers_type_after_compass_miss_before_preview(self):
        c = MagicMock()
        history = [dict(policy_number='P1', carrier='COMMONWEALTH', transaction_type='RENEWAL',
                        franchise='DTF0134', state='TX', del_toro_percent='.09', franchise_percent='.04')]
        c.cursor.return_value.fetchall.side_effect = [[dict(self.row(), id=1, office_number='134')], history]
        with patch('reparto_comisiones.obtener_tipos_lote', return_value=[
                dict(request_id='1', transaction_type=None, error='policy_not_found')]):
            rows = cargar_comisiones_excel(c, {'file_id': 'a.pdf', 'rows': [{'accounting_month': '2026-09-01'}]},
                                          'COMMONWEALTH', consultar_tipos=True)
        self.assertEqual(rows[0]['transaction_type'], 'RENEWAL')
        self.assertEqual(rows[0]['transaction_type_source'], 'historico')
        self.assertEqual(rows[0]['compass_lookup_error'], 'policy_not_found')
        self.assertIsNone(rows[0]['commission_alert'])

    def test_history_type_franchise_and_office_reach_excel(self):
        originals = [dict(self.row(None), franchise_number=None, type_lookup_error='policy_not_found')]
        history = [dict(policy_number='P1', carrier='OTHER', transaction_type='Renewal',
                        franchise='DTF0134', state='TX', del_toro_percent='.99', franchise_percent='.99')]
        completar_datos_historicos(originals, history)
        rows = calcular_comisiones_raw(originals, TARIFAS, 'COMMONWEALTH', historico=history)
        self.assertEqual(rows[0]['transaction_type'], 'RENEWAL')
        self.assertEqual(rows[0]['office_number'], 'DTF0134')
        self.assertEqual(rows[0]['franchise_number_source'], 'historico')
        self.assertEqual(rows[0]['del_toro_commission'], Decimal('64.85'))
        data = generar_excel(rows, Decimal('64.85'), bank_bytes(), 'Statement', 1, 2, 'COMMONWEALTH', {})
        book = load_workbook(BytesIO(data), data_only=True)
        ws = book['Data']
        cols = {cell.value: cell.column for cell in ws[1]}
        self.assertEqual(ws.cell(2, cols['Transaction Type']).value, 'RENEWAL')
        book.close()

    def test_history_does_not_replace_valid_compass_data(self):
        rows = [dict(self.row(), office_id='office', office_number='82', franchise_number='DTF0082')]
        history = [dict(policy_number='P1', carrier='COMMONWEALTH', transaction_type='RENEWAL', franchise='DTF0134')]
        completar_datos_historicos(rows, history)
        self.assertEqual(rows[0]['transaction_type'], 'NEW_BUSINESS')
        self.assertEqual(rows[0]['franchise_number'], 'DTF0082')
        self.assertEqual(rows[0]['office_number'], '82')
        self.assertNotIn('transaction_type_source', rows[0])

    def test_persisted_history_franchise_is_labeled(self):
        rows = [dict(self.row(), office_id=None, franchise_number='DTF0134')]
        completar_datos_historicos(rows, [dict(policy_number='P1', franchise='DTF0134')])
        self.assertEqual(rows[0]['franchise_number_source'], 'historico')

    def row(self, tipo='NEW_BUSINESS', premium='1000'):
        return dict(policy_number='P1', transaction_type=tipo, premium_amount=premium,
                    commission_amount='64.85', producer_code='12IA', franchise_number='DTF0082')

    def test_types_choose_configured_rate_and_preserve_statement(self):
        rows = calcular_comisiones_raw([self.row(), self.row('RENEWAL')], TARIFAS, 'COMMONWEALTH', 'TX', '')
        self.assertEqual([r['del_toro_commission'] for r in rows], [Decimal('64.85'), Decimal('64.85')])
        self.assertEqual([r['franchise_commission'] for r in rows], [Decimal('70'), Decimal('50')])
        self.assertEqual(rows[0]['commission_amount'], '64.85')

    def test_missing_type_state_or_rate_is_not_inferred(self):
        for tipo, estado, tarifas in [('active', 'TX', TARIFAS), ('NEW_BUSINESS', None, TARIFAS),
                                      ('NEW_BUSINESS', 'FL', TARIFAS), ('NEW_BUSINESS', 'TX', TARIFAS * 2)]:
            row = calcular_comisiones_raw([self.row(tipo)], tarifas, 'COMMONWEALTH', estado, '')[0]
            self.assertTrue(row['commission_alert'])
            self.assertEqual(row['del_toro_commission'], Decimal('64.85'))

    def test_signed_rounding_and_chargebacks(self):
        rows = calcular_comisiones_raw([self.row(premium='-10.50'),
            dict(policy_number='', transaction_type='Report fee', commission_amount='-5.50', comm_percent='1')],
            TARIFAS, 'COMMONWEALTH', 'TX', '')
        self.assertEqual(rows[0]['franchise_commission'], Decimal('-.735'))
        self.assertIsNone(rows[1]['franchise_commission'])
        self.assertIsNone(rows[1]['commission_alert'])
        self.assertEqual(rows[1]['comm_percent'], '1')

    def test_excel_stores_fraction_and_formats_as_decimal(self):
        rows = calcular_comisiones_raw([self.row()], TARIFAS, 'COMMONWEALTH', 'TX', '')
        data = generar_excel(rows, Decimal('64.85'), bank_bytes(), 'Statement', 1, 2, 'COMMONWEALTH', {})
        book = load_workbook(BytesIO(data), data_only=True)
        ws = book['Data']
        cols = {cell.value: cell.column for cell in ws[1]}
        self.assertEqual(ws.cell(2, cols['Del Toro %']).value, .06485)
        self.assertEqual(ws.cell(2, cols['Del Toro %']).number_format, '0.###')
        lda = book['Report_LDA']
        cols = {cell.value: cell.column for cell in lda[1]}
        self.assertEqual(lda.cell(2, cols['Franchise %']).number_format, '0.###')
        self.assertEqual(lda.cell(2, cols['Franchise $']).value, 70)
        self.assertEqual(book['Pivot']['B7'].value, 0)
        book.close()

    def test_office_state_selects_rate_per_policy(self):
        tarifas = TARIFAS + [dict(TARIFAS[0], state='FL', del_toro_percent=Decimal('.15'), franchise_percent=Decimal('.08'))]
        c = MagicMock()
        c.cursor.return_value.fetchall.side_effect = [[dict(self.row(), office_id='tx-office', state='Texas'),
            dict(self.row(), office_id='fl-office', state='FL'), dict(self.row(), office_id='unknown', state=None)], [], tarifas]
        rows = cargar_comisiones_excel(c, {'file_id': 'a.pdf', 'rows': [{'accounting_month': '2026-09-01'}]}, 'COMMONWEALTH')
        self.assertEqual(rows[0]['del_toro_commission'], Decimal('64.85'))
        self.assertEqual(rows[1]['del_toro_commission'], Decimal('64.85'))
        self.assertEqual(rows[2]['del_toro_commission'], Decimal('64.85'))
        self.assertTrue(rows[2]['commission_alert'])
        sql = c.cursor.return_value.execute.call_args_list[0].args[0]
        self.assertIn('o.office_id = r.office_id', sql)

    def test_commonwealth_chargebacks_get_100_percent_franchise_passthrough_regardless_of_code_resolution(self):
        """Los chargebacks de COMMONWEALTH son 100% a cargo de la franquicia, igual que los MVR
        de THE GENERAL: Del Toro % debe salir limpio en 1 (Decimal exacto, no un valor impreciso
        que se muestre como '1.') y Franchise %/$ deben llenarse con el mismo importe del
        chargeback, se haya podido resolver la franquicia por código (completar_chargebacks) o
        no -esa resolución no cambia."""
        c = MagicMock()
        resuelto = dict(policy_number='', transaction_type='Motor Vehicle Report x 3 for tx12ia1',
                        commission_amount='-16.50', comm_percent='1', producer_code='', producer_name='')
        no_resuelto = dict(policy_number='', transaction_type='Motor Vehicle Report x 1 for tx12ia5',
                           commission_amount='-5.50', comm_percent='1', producer_code='', producer_name='')
        c.cursor.return_value.fetchall.side_effect = [
            [resuelto, no_resuelto],  # rows
            [],  # historico
            [{'code': '12ia1', 'franchise': 'DTF0073', 'is_master_code': 0}],  # codigos (completar_chargebacks)
        ]
        rows = cargar_comisiones_excel(c, {'file_id': 'a.pdf', 'rows': [{'accounting_month': '2026-09-01'}]}, 'COMMONWEALTH')
        con_franquicia, sin_franquicia = rows
        self.assertEqual(con_franquicia['franchise_number'], 'DTF0073')
        self.assertIsNone(sin_franquicia['franchise_number'])
        for row, monto in ((con_franquicia, Decimal('-16.50')), (sin_franquicia, Decimal('-5.50'))):
            self.assertEqual(row['comm_percent'], Decimal('1'))
            self.assertEqual(row['franchise_percent'], Decimal('1'))
            self.assertEqual(row['franchise_commission'], monto)
            self.assertEqual(row['del_toro_commission'], monto)
            self.assertEqual(row['commission_difference'], Decimal('0'))

    def test_commonwealth_chargeback_report_shows_clean_percent_and_filled_franchise_columns(self):
        """Regresion de pantalla: el Report_LDA debe mostrar Del Toro % en formato entero limpio
        (sin punto colgante) y Franchise %/$ /Difference con valores, no vacios, para chargebacks
        de COMMONWEALTH."""
        rows = [
            dict(policy_number='', transaction_type='Motor Vehicle Report x 1 for tx12ia5',
                commission_amount='-5.50', comm_percent='1', producer_code='tx12ia5', producer_name='',
                franchise_number=None),
            dict(policy_number='', transaction_type='Motor Vehicle Report x 3 for tx12ia1',
                commission_amount='-16.50', comm_percent='1', producer_code='tx12ia1', producer_name='',
                franchise_number='DTF0073'),
        ]
        for row in rows:
            commission = Decimal(str(row['commission_amount']))
            row['comm_percent'] = Decimal('1')
            row['del_toro_commission'] = commission
            row['franchise_percent'] = Decimal('1')
            row['franchise_commission'] = commission
            row['commission_difference'] = Decimal('0')
        data = generar_excel(rows, Decimal('-22'), None, 'Sheet', 1, None, 'COMMONWEALTH', {})
        book = load_workbook(BytesIO(data), data_only=True)
        lda = book['Report_LDA']
        cols = {cell.value: cell.column for cell in lda[1]}
        for fila in (2, 3):
            self.assertEqual(lda.cell(fila, cols['Del Toro %']).value, 1.0)
            self.assertEqual(lda.cell(fila, cols['Del Toro %']).number_format, '0')
            self.assertEqual(lda.cell(fila, cols['Franchise %']).number_format, '0')
            self.assertIsNotNone(lda.cell(fila, cols['Franchise $']).value)
            self.assertIsNotNone(lda.cell(fila, cols['Difference']).value)
        book.close()

    def test_loads_persisted_raw_instead_of_stale_snapshot(self):
        c = MagicMock()
        c.cursor.return_value.fetchall.side_effect = [[self.row('RENEWAL')], [], TARIFAS]
        rows = cargar_comisiones_excel(c, {'file_id': 'a.pdf', 'rows': [{'accounting_month': '2026-09-01'}]},
                                      'COMMONWEALTH', 'TX', '')
        self.assertEqual(rows[0]['del_toro_commission'], Decimal('64.85'))

    def test_history_takes_priority_without_type_or_state(self):
        history = [dict(policy_number='P1', carrier='COMMONWEALTH',
                        del_toro_percent='.09', franchise_percent='.04'),
                   dict(policy_number='P1', carrier='COMMONWEALTH',
                        del_toro_percent='.08', franchise_percent='.03')]
        original = dict(self.row(None), type_lookup_error='Compass no disponible')
        row = calcular_comisiones_raw([original], TARIFAS, 'COMMONWEALTH', historico=history)[0]
        self.assertEqual(row['del_toro_commission'], Decimal('64.85'))
        self.assertEqual(row['franchise_commission'], Decimal('40'))
        self.assertIsNone(row['commission_alert'])
        self.assertEqual(row['commission_amount'], original['commission_amount'])

    def test_missing_or_incomplete_or_other_carrier_history_uses_rates(self):
        for history in ([], [dict(policy_number='P2', carrier='COMMONWEALTH',
                                  del_toro_percent='.09', franchise_percent='.04')],
                        [dict(policy_number='P1', carrier='COMMONWEALTH',
                              del_toro_percent='.09', franchise_percent=None)],
                        [dict(policy_number='P1', carrier='OTHER',
                              del_toro_percent='.09', franchise_percent='.04')]):
            with self.subTest(history=history):
                row = calcular_comisiones_raw([self.row()], TARIFAS, 'COMMONWEALTH', 'TX', '',
                                              historico=history)[0]
                self.assertEqual(row['del_toro_commission'], Decimal('64.85'))

    def test_loader_queries_history_first_and_skips_rates_when_found(self):
        c = MagicMock()
        history = [dict(policy_number='P1', carrier='COMMONWEALTH',
                        del_toro_percent='.09', franchise_percent='.04')]
        c.cursor.return_value.fetchall.side_effect = [[self.row()], history]
        rows = cargar_comisiones_excel(c, {'file_id': 'a.pdf', 'rows': [{'accounting_month': '2026-09-01'}]},
                                      'COMMONWEALTH')
        self.assertEqual(rows[0]['del_toro_commission'], Decimal('64.85'))
        calls = c.cursor.return_value.execute.call_args_list
        self.assertEqual(len(calls), 2)
        self.assertIn('historic_data_commissions', calls[1].args[0])
        self.assertIn('ORDER BY h.report_month DESC, h.date DESC', calls[1].args[0])
        self.assertEqual(calls[1].args[1], ('2026-09-01', 'a.pdf', '2026-09-01'))

    def test_loader_falls_back_only_for_policies_missing_from_history(self):
        c = MagicMock()
        history = [dict(policy_number='P1', carrier='COMMONWEALTH',
                        del_toro_percent='.09', franchise_percent='.04')]
        c.cursor.return_value.fetchall.side_effect = [[self.row(), dict(self.row(), policy_number='P2')],
                                                     history, TARIFAS]
        rows = cargar_comisiones_excel(c, {'file_id': 'a.pdf', 'rows': [{'accounting_month': '2026-09-01'}]},
                                      'COMMONWEALTH', 'TX', '')
        self.assertEqual([r['del_toro_commission'] for r in rows], [Decimal('64.85'), Decimal('64.85')])
        calls = c.cursor.return_value.execute.call_args_list
        self.assertIn('historic_data_commissions', calls[1].args[0])
        self.assertIn('commission_rates', calls[2].args[0])
