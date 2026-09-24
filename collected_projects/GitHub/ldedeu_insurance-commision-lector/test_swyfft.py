import unittest
from datetime import date, datetime
from io import BytesIO
from openpyxl import Workbook
from swyfft import HEADERS, leer_statement, SHEET_CARRIERS
from swyfft import (preparar_statement, guardar_statement, cargar_comisiones_excel, completar_franquicias,
                    _es_formato_cheque, extraer_filas_pagina_cheque, preparar_statement_cheque)
from unittest.mock import MagicMock, patch
from decimal import Decimal


def _bloque_cheque(texto, x, y, height=0.015):
    return {'text': texto, 'x': x, 'y': y, 'height': height}


def _bloques_pagina_cheque_real():
    """Bloques reales (posiciones OCR) de SWYFFT-refund-separado.pdf: un talon de cheque de
    SWYFFT con una mini tabla propia debajo (Producer Location Name/Loc.#/Insured Name/Policy
    Number/Account Number/Description/Gross Premium/Net Due/Paid Now), seguido en la misma
    pagina de una copia completa del cheque (para archivo del banco) sin esa tabla."""
    return [
        _bloque_cheque('SWIFFI,LLC', 0.0671, 0.0052),
        _bloque_cheque('CANOPIUS REFUND-5001', 0.1106, 0.0177),
        _bloque_cheque('9534', 0.9379, 0.0191),
        _bloque_cheque('CHASE BANK', 0.6740, 0.0200),
        _bloque_cheque('350 MOUNT KEMBLE AVENUE', 0.1136, 0.0324),
        _bloque_cheque('225 SOUTH STREET', 0.6731, 0.0340),
        _bloque_cheque('MORRISTOWN, NJ 07960', 0.0976, 0.0450),
        _bloque_cheque('MORRISTOWN, NJ 07960', 0.6723, 0.0466),
        _bloque_cheque('DATE', 0.7606, 0.0736),
        _bloque_cheque('09/11/2026', 0.7593, 0.0936),
        _bloque_cheque('PAY TO THE', 0.0498, 0.0995),
        _bloque_cheque('Del Toro Franchising Corp', 0.1955, 0.1014),
        _bloque_cheque('ORDER OF', 0.0473, 0.1093),
        _bloque_cheque('$', 0.7965, 0.1188),
        _bloque_cheque('483.50', 0.9396, 0.1188),
        _bloque_cheque('DOLLARS', 0.9244, 0.1538),
        _bloque_cheque('VOID AFTER 90 DAYS', 0.7796, 0.1723),
        _bloque_cheque('Del Toro Franchising Corp', 0.2086, 0.1823),
        _bloque_cheque('42 NW 27TH AVENUE', 0.1968, 0.1950),
        _bloque_cheque('MIAMI, FL 33125', 0.1782, 0.2073),
        _bloque_cheque('MEMO', 0.0431, 0.2379),
        _bloque_cheque('Commission', 0.0621, 0.2572),
        _bloque_cheque('AUTHORIZED SIGNATURE', 0.8940, 0.2575),
        _bloque_cheque('Producer Location Name Loc. # Insured Name', 0.1398, 0.3550),
        _bloque_cheque('Policy Number', 0.3623, 0.3580),
        _bloque_cheque('Account Number Description', 0.5283, 0.3603),
        _bloque_cheque('Gross Premium', 0.7154, 0.3626),
        _bloque_cheque('Net Due', 0.8294, 0.3632),
        _bloque_cheque('Paid Now', 0.9096, 0.3649),
        _bloque_cheque('De La Cerda & Associates, L 13234 MIRACLE ESTATES L CA92-000831-00', 0.2061, 0.3668),
        _bloque_cheque('New', 0.5617, 0.3721),
        _bloque_cheque('$4,025.68', 0.7323, 0.3737),
        _bloque_cheque('$0.00', 0.8361, 0.3747),
        _bloque_cheque('$483.50', 0.9143, 0.3753),
        # Copia del cheque para el archivo del banco (misma pagina, sin tabla propia).
        _bloque_cheque('SWYFFT, LLC', 0.0541, 0.6882),
        _bloque_cheque('9534', 0.9248, 0.7009),
        _bloque_cheque('CANOPIUS REFUND-5001', 0.0967, 0.7017),
        _bloque_cheque('CHASE BANK', 0.6601, 0.7021),
        _bloque_cheque('225 SOUTH STREET', 0.6601, 0.7163),
        _bloque_cheque('350 MOUNT KEMBLE AVENUE', 0.0997, 0.7166),
        _bloque_cheque('MORRISTOWN, NJ 07960', 0.0845, 0.7281),
        _bloque_cheque('MORRISTOWN, NJ 07960', 0.6592, 0.7281),
        _bloque_cheque('DATE', 0.7483, 0.7510),
        _bloque_cheque('09/11/2026', 0.7470, 0.7723),
        _bloque_cheque('PAY TO THE', 0.0384, 0.7817),
        _bloque_cheque('ORDER OF', 0.0355, 0.7920),
        _bloque_cheque('Del Toro Franchising Corp', 0.1837, 0.7916),
        _bloque_cheque('483.50', 0.9278, 0.7958),
        _bloque_cheque('$', 0.7846, 0.7981),
        _bloque_cheque('Non-Transferable.This Is Not A Check.', 0.4540, 0.9221),
        _bloque_cheque('SEP 21 2026', 0.8695, 0.9699),
    ]


def _pagina_cheque_real():
    bloques = _bloques_pagina_cheque_real()
    return {'blocks': bloques, 'text': '\n'.join(b['text'] for b in bloques), 'image': None}


class SwyfftTests(unittest.TestCase):
    def test_summary_rows_are_not_imported_but_incomplete_policies_still_fail(self):
        for total_label in ('Grand Total', ''):
            book, sheet = self.workbook()
            sheet.append(['Agency', '12', '', '', 'INV', datetime(2026, 9, 20),
                          'Client', '00123', 'Home', datetime(2026, 9, 1), 'New', '', '', 100, 12])
            sheet.append([total_label] + [None] * 12 + [100, 12])
            rows = preparar_statement(self.data(book), '2026-09')
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]['commission_amount'], '12')
        book, sheet = self.workbook()
        sheet.append(['Agency', '12', '', '', 'INV', '', 'Client', '', 'Home',
                      datetime(2026, 9, 1), 'New', '', '', 100, 12])
        with self.assertRaisesRegex(ValueError, 'No se reconoció como total'):
            preparar_statement(self.data(book), '2026-09')

    def test_detects_header_at_different_rows_and_preserves_source_row(self):
        for header in (1, 5, 35):
            with self.subTest(header=header):
                book = Workbook()
                sheet = book.active
                for _ in range(header - 1):
                    sheet.append(['SWYFFT report', 'Producer Name'])
                sheet.append(['  PRODUCER   NAME  '] + list(HEADERS[1:]))
                sheet.append([None] * len(HEADERS))
                sheet.append(['Agency', '0012', 'Sub', 'Contact', 'INV', datetime(2026, 9, 20),
                              'Client', '00123', 'Home', datetime(2026, 9, 1), 'New',
                              datetime(2026, 9, 15), datetime(2026, 9, 18), 100, 12])
                rows = preparar_statement(self.data(book), '2026-09')
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0]['source_row'], header + 2)
                self.assertEqual(rows[0]['policy_number'], '00123')

    def prepared(self, tipo='New', premium=100, commission=12):
        book, sheet = self.workbook()
        sheet.append(['Agency', '0012', 'Sub', 'Contact', 'INV', datetime(2026, 9, 20),
                      'Client', '00123', 'Home', datetime(2026, 9, 1), tipo,
                      datetime(2026, 9, 15), datetime(2026, 9, 18), premium, commission])
        return preparar_statement(self.data(book), '2026-09')

    def test_transaction_mapping_signed_amounts_and_policy_date(self):
        for label, expected in [('New', 'NEW_BUSINESS'), ('Cancellation', 'CANCEL'),
                                ('Reinstatement', 'REINSTATE'), ('Endorsement', 'ENDORSE')]:
            row = self.prepared(label, -100, -12)[0]
            self.assertEqual(row['transaction_type'], expected)
            self.assertEqual(row['effective_date'], '2026-09-01')
            self.assertEqual(row['producer_code'], '12')
            self.assertEqual(row['source_data']['Producer Location ID'], '0012')
            self.assertEqual(row['commission_amount'], '-12')
            self.assertEqual(Decimal(row['del_toro_percent']), Decimal('.12'))
        with self.assertRaises(ValueError):
            self.prepared('Unknown')
        with self.assertRaises(ValueError):
            self.prepared(premium=0, commission=12)
        self.assertEqual(Decimal(self.prepared(premium=0, commission=0)[0]['del_toro_percent']), 0)

    def test_sub_location_contains_registered_name_without_compass(self):
        rows = self.prepared()
        rows[0]['source_data']['Sub Location'] = 'SWYFFT - My   Franchise LLC - Florida'
        c = MagicMock()
        c.cursor.return_value.fetchall.return_value = [dict(office_id='uid', office_number='82', office_name='my franchise', state='FL')]
        with patch('swyfft.compass.obtener_token') as login:
            completar_franquicias(c, rows)
        login.assert_not_called()
        self.assertEqual(rows[0]['franchise_number'], 'DTF0082')
        self.assertEqual(rows[0]['franchise_number_source'], 'sub_location')
        self.assertEqual(rows[0]['office_id'], 'uid')
        self.assertEqual(rows[0]['state'], 'FL')
        self.assertNotIn('franchises_carrier_codes', c.cursor.return_value.execute.call_args.args[0])

    def test_ambiguous_sub_location_does_not_choose_first_matching_office(self):
        import json
        rows = self.prepared()
        rows[0]['source_data'] = json.dumps({'Sub Location': 'North Agency South Agency'})
        c = MagicMock()
        c.cursor.return_value.fetchall.return_value = [
            dict(office_id='north', office_number='82', office_name='North Agency', state='FL'),
            dict(office_id='south', office_number='83', office_name='South Agency', state='FL')]
        with patch('swyfft.compass.obtener_token', return_value='token'), \
             patch('swyfft.compass.buscar_poliza', return_value=None), \
             patch('swyfft.resolver_franquicia', return_value=None):
            completar_franquicias(c, rows)
        self.assertIsNone(rows[0]['franchise_number'])
        self.assertIn('ambiguo', rows[0]['code_lookup_alert'])

    def test_calculates_each_type_using_swyfft_rates_and_signed_premium(self):
        for tipo in ('New', 'Cancellation', 'Reinstatement', 'Endorsement'):
            row = dict(self.prepared(tipo, -100, -12)[0], id=1, state='FL')
            rates = [dict(carrier='SWYFFT', state='FL', transaction_type=row['transaction_type'],
                          business_line='', franchise_percent='.08')]
            c = MagicMock()
            c.cursor.return_value.fetchall.side_effect = [[row], rates]
            with patch('swyfft.completar_franquicias'):
                result = cargar_comisiones_excel(c, {'file_id': 'test.xlsx', 'rows': [{'accounting_month': '2026-09-01'}]}, consultar_tipos=True)[0]
            self.assertEqual(result['del_toro_percent'], Decimal('.12'))
            self.assertEqual(result['del_toro_commission'], Decimal('-12'))
            self.assertEqual(result['franchise_commission'], Decimal('-8'))
            self.assertIsNone(result['commission_alert'])

    def test_commission_split_without_file_id_combines_every_file_for_the_month(self):
        """SWYFFT puede llegar en varios ficheros del mismo mes (el Excel de transferencia y uno
        o más talones de cheque, cada uno su propio file_id): sin file_id en el snapshot, se
        concilia y exporta como un único statement combinado, no uno por archivo."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        row_excel = dict(id=1, policy_number='P1', producer_code='13234', insured_name='A',
                         transaction_type='NEW_BUSINESS', franchise_number='DTF0092', state='FL',
                         premium_amount=Decimal('100'), commission_amount=Decimal('12'),
                         statement_rate=Decimal('0.12'), source_sheet='SWYFFT',
                         accounting_month=date(2026, 9, 1))
        row_cheque = dict(id=2, policy_number='CA92-000831-00', producer_code='13234', insured_name='B',
                          transaction_type='NEW_BUSINESS', franchise_number='DTF0092', state='FL',
                          premium_amount=Decimal('4025.68'), commission_amount=Decimal('483.50'),
                          statement_rate=Decimal('0.12'), source_sheet='SWYFFT_CHEQUE',
                          accounting_month=date(2026, 9, 1))
        cursor.fetchall.side_effect = [[row_excel, row_cheque], []]
        with patch('swyfft.completar_franquicias'):
            snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 9, 1)}]}
            result = cargar_comisiones_excel(connection, snapshot)
        select_sql = cursor.execute.call_args_list[0].args[0]
        self.assertNotIn('file_id', select_sql)
        self.assertEqual(len(result), 2)
        self.assertEqual({r['policy_number'] for r in result}, {'P1', 'CA92-000831-00'})
        self.assertEqual({r['source_sheet'] for r in result}, {'SWYFFT', 'SWYFFT_CHEQUE'})

    def test_atomic_import_retry_and_modified_source(self):
        rows = self.prepared()
        c = MagicMock()
        c.cursor.return_value.fetchone.return_value = {'acquired': 1}
        c.cursor.return_value.fetchall.return_value = []
        with patch('swyfft.conectar', return_value=c), patch('swyfft.completar_franquicias'):
            self.assertEqual(guardar_statement(rows, 'test.xlsx', b'contenido-swyfft'), 1)
            existing = [{k: rows[0][k] for k in ('source_sheet', 'source_row', 'source_data')}]
            c.cursor.return_value.fetchall.return_value = existing
            self.assertEqual(guardar_statement(rows, 'test.xlsx', b'contenido-swyfft'), 0)
            existing[0] = {**existing[0], 'source_data': {'changed': True}}
            with self.assertRaises(ValueError):
                guardar_statement(rows, 'test.xlsx', b'contenido-swyfft')
        inserts = [call for call in c.cursor.return_value.execute.call_args_list if call.args[0].startswith('INSERT')]
        self.assertEqual(len(inserts), 1)
        self.assertEqual(inserts[0].args[1][3], 'SWYFFT')
        c.rollback.assert_called_once()

    def test_save_uses_content_hash_as_file_id_not_the_file_name(self):
        """Regresion: el mismo statement subido con otro nombre de archivo debe reconocerse
        como el mismo import (mismo file_id, calculado del contenido), no como uno nuevo."""
        rows = self.prepared()
        c = MagicMock()
        c.cursor.return_value.fetchone.return_value = {'acquired': 1}
        c.cursor.return_value.fetchall.return_value = []
        with patch('swyfft.conectar', return_value=c), patch('swyfft.completar_franquicias'):
            guardar_statement(rows, 'nombre-original.xlsx', b'mismo-contenido')
            insert = next(call for call in c.cursor.return_value.execute.call_args_list
                         if call.args[0].startswith('INSERT'))
            file_id_guardado = insert.args[1][0]

            existing = [{k: rows[0][k] for k in ('source_sheet', 'source_row', 'source_data')}]
            c.cursor.return_value.fetchall.return_value = existing
            self.assertEqual(guardar_statement(rows, 'nombre-reenviado (1).xlsx', b'mismo-contenido'), 0)
            select = next(call for call in c.cursor.return_value.execute.call_args_list
                         if call.args[0].startswith('SELECT source_sheet'))
            self.assertEqual(select.args[1][0], file_id_guardado)

    def test_export_reconciles_swyfft_and_preserves_calculated_commissions(self):
        from conciliacion import generar_excel
        from openpyxl import load_workbook
        from reparto_comisiones import calcular_comisiones_raw
        row = dict(self.prepared()[0], state='FL', franchise_number='DTF0082')
        rows = calcular_comisiones_raw([row], [dict(carrier='SWYFFT', state='FL',
            transaction_type='NEW_BUSINESS', business_line='', franchise_percent='.08')], 'SWYFFT')
        rows[0]['commission_alert'] = 'Aviso anterior de una búsqueda ya resuelta'
        bank = Workbook()
        bank.active.append(['Description', 'Amount'])
        bank.active.append(['SWYFFT PAYMENT', 12])
        data = generar_excel(rows, Decimal('12'), self.data(bank), 'Sheet', 1, 2, 'SWYFFT', {})
        book = load_workbook(BytesIO(data), data_only=True)
        lda_cols = {c.value: c.column for c in book['Report_LDA'][1]}
        self.assertEqual(book['Data']['A2'].value, 'DTF0082')
        self.assertEqual(book['Data']['N2'].value, .12)
        self.assertEqual(book['Data']['O2'].value, 12)
        self.assertEqual(book['Report_LDA'].cell(2, lda_cols['Franchise $']).value, 8)
        self.assertEqual(book['Report_LDA'].cell(2, lda_cols['Difference']).value, 4)
        self.assertNotEqual(book['Data']['A2'].fill.fgColor.rgb, '00FFC7CE')
        columnas = {cell.value: cell.column for cell in book['Data'][1]}
        self.assertEqual(book['Data'].cell(2, columnas['Commission Alert']).value,
                         'Aviso anterior de una búsqueda ya resuelta')
        self.assertIn('Report_LDA', book.sheetnames)
        book.close()

    def workbook(self, labels=None):
        book = Workbook()
        first = book.active
        first.title = 'Arbitrary name'
        first.append(labels or list(HEADERS))
        return book, first

    def data(self, book):
        output = BytesIO()
        book.save(output)
        book.close()
        return output.getvalue()

    def test_first_sheet_whitespace_dates_and_leading_zero_ids(self):
        labels = list(HEADERS)
        labels[0] = 'Producer   Name'
        book, sheet = self.workbook(labels)
        values = ['Agency', 7, 'Sub', 'Contact', 'INV-001', datetime(2026, 9, 17),
                  'Client', '00123', 'Home', datetime(2026, 9, 1), 'New Business',
                  datetime(2026, 9, 1), datetime(2026, 9, 17), -1250.50, -125.05]
        sheet.append(values)
        sheet['B2'].number_format = '0000'
        book.create_sheet('Ignore').append(['Invalid headers'])
        rows = leer_statement(self.data(book))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['Producer Location ID'], '0007')
        self.assertEqual(rows[0]['Policy Number'], '00123')
        self.assertEqual(rows[0]['Policy Effective Date'], '2026-09-01')
        self.assertEqual(rows[0]['Premium Collected'], '-1250.5')
        self.assertEqual(rows[0]['Commissions Paid'], '-125.05')

    def test_missing_or_duplicate_headers_rejected(self):
        for labels in (list(HEADERS[:-1]), list(HEADERS) + ['Producer   Name']):
            book, sheet = self.workbook(labels)
            with self.assertRaises(ValueError):
                leer_statement(self.data(book))

    def test_sheet_carriers_accepts_cheque_source_for_the_same_carrier(self):
        """SWYFFT puede pagar parte por transferencia (Excel) y parte por cheque (PDF/foto): el
        origen 'cheque' es un source_sheet distinto pero cuenta como el mismo carrier SWYFFT."""
        self.assertEqual(SHEET_CARRIERS['SWYFFT_CHEQUE'], 'SWYFFT')

    def test_es_formato_cheque_detects_the_check_stub_mini_table(self):
        self.assertTrue(_es_formato_cheque([_pagina_cheque_real()]))
        book, sheet = self.workbook()
        sheet.append(['Agency', '12', '', '', 'INV', datetime(2026, 9, 20),
                      'Client', '00123', 'Home', datetime(2026, 9, 1), 'New', '', '', 100, 12])
        pagina_excel = {'blocks': [], 'text': 'Producer Name Policy Number Premium Collected', 'image': None}
        self.assertFalse(_es_formato_cheque([pagina_excel]))

    def test_extraer_filas_pagina_cheque_ignores_the_duplicate_check_copy_on_the_same_page(self):
        """Regresion: el talon repite el cheque completo (sin la tabla) mas abajo en la misma
        pagina, para el archivo del banco. Sin el corte por salto vertical, esa copia se leia
        como filas de datos y fallaba al no encontrar Producer/Loc.#/Insured/Policy Number."""
        filas = extraer_filas_pagina_cheque(_pagina_cheque_real(), 1)
        self.assertEqual(len(filas), 1)
        fila = filas[0]
        self.assertEqual(fila['Producer Name'], 'De La Cerda & Associates, L')
        self.assertEqual(fila['Producer Location ID'], '13234')
        self.assertEqual(fila['Insured'], 'MIRACLE ESTATES L')
        self.assertEqual(fila['Policy Number'], 'CA92-000831-00')
        self.assertEqual(fila['Transaction Type'], 'New')
        self.assertEqual(fila['Premium Collected'], '4025.68')
        self.assertEqual(fila['Commissions Paid'], '483.50')
        self.assertEqual(fila['Sub Location'], '')
        self.assertEqual(fila['Policy Effective Date'], '')

    def test_preparar_statement_cheque_builds_the_same_row_shape_as_the_excel_path(self):
        rows = preparar_statement_cheque([_pagina_cheque_real()], '2026-09')
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row['carrier'], 'SWYFFT')
        self.assertEqual(row['source_sheet'], 'SWYFFT_CHEQUE')
        self.assertEqual(row['policy_number'], 'CA92-000831-00')
        self.assertEqual(row['insured_name'], 'MIRACLE ESTATES L')
        self.assertEqual(row['producer_name'], 'De La Cerda & Associates, L')
        self.assertEqual(row['producer_code'], '13234')
        self.assertEqual(row['transaction_type'], 'NEW_BUSINESS')
        self.assertEqual(row['premium_amount'], '4025.68')
        self.assertEqual(row['commission_amount'], '483.50')
        # No hay Policy Effective Date en el talon: se usa la fecha del propio cheque (emision).
        self.assertEqual(row['effective_date'], '2026-09-11')

    def test_preparar_statement_cheque_handles_more_than_one_row_close_together(self):
        """No hay un ejemplo real con dos filas en el mismo cheque; se construye sinteticamente
        con las mismas anclas de columna para probar que el corte por salto vertical no rompe
        con varias filas reales pegadas entre si."""
        base = _bloques_pagina_cheque_real()
        segunda_fila = [
            _bloque_cheque('Otra Agencia LLC 55555 OTRO ASEGURADO SIC1111111-00', 0.2061, 0.4000),
            _bloque_cheque('Renewal', 0.5617, 0.4050),
            _bloque_cheque('$1,000.00', 0.7323, 0.4060),
            _bloque_cheque('$0.00', 0.8361, 0.4065),
            _bloque_cheque('$100.00', 0.9143, 0.4070),
        ]
        pagina = {'blocks': base + segunda_fila, 'text': '', 'image': None}
        filas = extraer_filas_pagina_cheque(pagina, 1)
        self.assertEqual(len(filas), 2)
        self.assertEqual(filas[0]['Policy Number'], 'CA92-000831-00')
        self.assertEqual(filas[1]['Policy Number'], 'SIC1111111-00')
        self.assertEqual(filas[1]['Transaction Type'], 'Renewal')
        self.assertEqual(filas[1]['Commissions Paid'], '100.00')

    def test_guardar_statement_accepts_cheque_rows(self):
        rows = preparar_statement_cheque([_pagina_cheque_real()], '2026-09')
        c = MagicMock()
        c.cursor.return_value.fetchone.return_value = {'acquired': 1}
        c.cursor.return_value.fetchall.return_value = []
        with patch('swyfft.conectar', return_value=c), patch('swyfft.completar_franquicias'):
            self.assertEqual(guardar_statement(rows, 'SWYFFT-refund-separado.pdf', b'contenido-cheque'), 1)
            insert = next(call for call in c.cursor.return_value.execute.call_args_list
                         if call.args[0].startswith('INSERT'))
            self.assertEqual(insert.args[1][4], 'SWYFFT_CHEQUE')
