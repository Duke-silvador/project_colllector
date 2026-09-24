import unittest
from io import BytesIO
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from requests import Timeout

from busqueda_polizas import consultar_polizas, exportar_excel, leer_polizas


class BusquedaTests(unittest.TestCase):
    def test_resolves_agent_by_id_and_caches_office(self):
        offices = [{'office_id': 'a', 'office_number': '007', 'office_name': 'Miami'}]
        agents = [{'id': 'other', 'firstname': 'Wrong'},
                  {'user_ID': '42', 'firstname': 'Ana', 'lastname': 'Perez',
                   'email': 'ana@example.com', 'npn': '123'}]
        with patch('busqueda_polizas.compass.buscar_poliza', return_value={
            'office_id': 'a', 'agent_id': 42, 'status_id': 'active',
            'expiration_date': '2027-09-17T00:00:00Z'
        }) as policies, patch('busqueda_polizas.compass.obtener_agentes_oficina', return_value=agents) as query:
            rows = consultar_polizas([(2, '001'), (3, '002'), (4, '001')], offices, 'token')
        query.assert_called_once_with('token', 'a')
        self.assertEqual(policies.call_count, 2)
        for row in rows:
            self.assertEqual(row['Nombre franquicia Compass'], 'Miami')
            self.assertEqual(row['ID agente'], '42')
            self.assertEqual(row['Fecha expiración'], '2027-09-17T00:00:00Z')
            self.assertEqual(row['Nombre agente'], 'Ana Perez')
            self.assertEqual(row['Correo agente'], 'ana@example.com')
            self.assertEqual(row['NPN agente'], '123')
            for duplicado in ('firstname', 'lastname', 'user_ID', 'email', 'npn'):
                self.assertNotIn(f'Agente · {duplicado}', row)
            self.assertNotIn('Datos agente', row)
            self.assertEqual(row['Resultado agente'], 'Encontrado')
        book = load_workbook(BytesIO(exportar_excel(rows)))
        headers = [c.value for c in book.active[1]]
        self.assertEqual(book.active.cell(2, headers.index('Nombre agente') + 1).value, 'Ana Perez')
        self.assertEqual(book.active.cell(2, headers.index('Fecha expiración') + 1).value, '2027-09-17T00:00:00Z')
        book.close()

    def test_agent_columns_survive_missing_first_row_and_nested_fields(self):
        with patch('busqueda_polizas.compass.buscar_poliza', side_effect=[None, {
            'office_id': 'a', 'agent_id': '42'
        }]), patch('busqueda_polizas.compass.obtener_agentes_oficina', return_value=[{
            'id': '42', 'name': 'Ana', 'address': {'city': 'Miami'},
            'licenses': [{'number': '00123'}], 'active': False
        }]):
            rows = consultar_polizas([(2, 'missing'), (3, 'found')], [], 'token')
        self.assertEqual(list(rows[0]), list(rows[1]))
        self.assertEqual(rows[0]['Agente · address · city'], '')
        self.assertEqual(rows[1]['Agente · address · city'], 'Miami')
        self.assertEqual(rows[1]['Agente · licenses · 1 · number'], '00123')
        self.assertEqual(rows[1]['Agente · active'], 'False')
        book = load_workbook(BytesIO(exportar_excel(rows)))
        headers = [c.value for c in book.active[1]]
        self.assertEqual(book.active.cell(3, headers.index('Agente · licenses · 1 · number') + 1).value, '00123')
        book.close()

    def test_agent_failure_does_not_discard_policy(self):
        for agents, status in [([], 'Agente no encontrado'), (Timeout(), 'Error al consultar')]:
            with self.subTest(status=status), patch('busqueda_polizas.compass.buscar_poliza', return_value={
                'office_id': 'a', 'agent_id': 'missing', 'status_id': 'active'
            }), patch('busqueda_polizas.compass.obtener_agentes_oficina',
                      **({'side_effect': agents} if isinstance(agents, Exception) else {'return_value': agents})) as query:
                rows = consultar_polizas([(2, '001'), (3, '002')], [], 'token')
            query.assert_called_once_with('token', 'a')
            self.assertEqual(rows[0]['Estado póliza'], 'active')
            self.assertEqual(rows[0]['ID agente'], 'missing')
            self.assertEqual(rows[0]['Nombre agente'], '')
            self.assertIn(status, rows[0]['Resultado agente'])

    def test_missing_agent_id_skips_query(self):
        with patch('busqueda_polizas.compass.buscar_poliza', return_value={'office_id': 'a'}), \
             patch('busqueda_polizas.compass.obtener_agentes_oficina') as query:
            rows = consultar_polizas([(None, '001')], [], 'token')
        query.assert_not_called()
        self.assertEqual(rows[0]['Resultado agente'], 'Póliza sin ID de agente')

    def test_match_duplicates_missing_and_error(self):
        filas = [(2, '001'), (3, '001'), (4, '002'), (5, '003'), (6, ''), (7, '004')]
        oficinas = [{'office_id': 'a', 'office_number': '007', 'office_name': 'Miami'}]
        with patch('busqueda_polizas.compass.buscar_poliza', side_effect=[
            {'office_id': 'a', 'status_id': 'active'}, None, Timeout(),
            {'office_id': 'b', 'status_id': 'cancelled'}
        ]) as buscar:
            rows = consultar_polizas(filas, oficinas, 'token')
        self.assertEqual(buscar.call_count, 4)
        self.assertEqual(len(rows), 6)
        self.assertEqual(rows[0]['Número oficina'], 'DTF0007')
        self.assertEqual(rows[1]['Oficina'], 'Miami')
        self.assertEqual(rows[2]['Resultado'], 'Póliza no encontrada')
        self.assertIn('Error', rows[3]['Resultado'])
        self.assertEqual(rows[4]['Resultado'], 'Sin número de póliza')
        self.assertEqual(rows[5]['Estado póliza'], 'cancelled')
        self.assertEqual(rows[5]['Resultado'], 'Oficina sin coincidencia')

    def test_excel_preserves_numbers_and_source_rows(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['Póliza', 'Otro'])
        sheet.append(['00123', None])
        sheet.append([45, None])
        sheet['A3'].number_format = '00000'
        sheet.append([None, None])
        sheet.append([None, 'dato'])
        data = BytesIO()
        book.save(data)
        self.assertEqual(leer_polizas(data.getvalue(), sheet.title, 1, 0),
                         [(2, '00123'), (3, '00045'), (5, '')])

    def test_export_keeps_literal_text(self):
        data = exportar_excel([{'Póliza': '00123', 'Oficina': '=1+1'}])
        book = load_workbook(BytesIO(data))
        self.assertEqual(book.active['A2'].value, '00123')
        self.assertEqual(book.active['B2'].data_type, 's')
        book.close()


if __name__ == '__main__':
    unittest.main()
