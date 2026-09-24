"""Carga THE GENERAL. El statement siempre llega en Excel (formato fijo, con o sin encabezado
agrupado en varias filas); no hay version PDF/imagen para este carrier. Si el archivo trae una
segunda hoja llamada MVR, sus filas con "Chargeback Applies"='Yes' se importan junto con las
demas, como chargebacks (100% a cargo de la franquicia, igual que COMMONWEALTH). Puedes subir uno
o varios documentos a la vez; se importan por separado pero se concilian y exportan juntos, mas
abajo."""
import hashlib
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from idiomas import interfaz as st, texto
from bass import importe
from the_general import detectar_encabezado, preparar_statement, guardar
from database import conectar
from importacion import calcular_file_id, buscar_posibles_duplicados
from conciliacion_ui import mostrar_conciliacion
from tablas_statements import mostrar_tabla


def _revisar_duplicados(records, month, file_id):
    """Un solo aviso resumido (no uno por poliza: con cientos de filas repetidas se vuelve una
    pared de cajas de advertencia). guardar() ya protege de verdad contra duplicar el mismo
    contenido con otro archivo (solo marca updated_at); esto es solo informativo, antes de
    importar."""
    try:
        accounting_month = datetime.strptime(month, '%Y-%m').date()
        pares = [(r['policy_prefix_plus_policy_number'], importe(r['total_commission_amount']))
                 for r in records if r.get('policy_prefix_plus_policy_number')]
        connection = conectar()
        try:
            duplicados = buscar_posibles_duplicados(connection, 'st_the_general_raw',
                                                     'policy_prefix_plus_policy_number', 'total_commission_amount',
                                                     accounting_month, file_id, pares)
        finally:
            connection.close()
        if not duplicados:
            return
        archivos = sorted({archivo for d in duplicados for archivo in d['archivos']})
        st.info(texto(
            f'{len(duplicados)} de estas filas ya están guardadas este mes en {", ".join(archivos)} con el mismo '
            'importe. Si es el mismo statement con otro nombre de archivo, se puede importar igual: no se '
            'duplican, solo se actualiza la fecha de la fila ya guardada.',
            f'{len(duplicados)} of these rows are already saved this month under {", ".join(archivos)} with the '
            'same amount. If this is the same statement under a different file name, it can be imported anyway: '
            "it won't be duplicated, only the already-saved row's date gets updated."))
    except (ValueError, KeyError):
        pass  # Los importes invalidos ya se senalan por campo arriba; no bloquea el aviso.


def _bloque_excel(archivo, month):
    contenido = archivo.getvalue()
    try:
        book = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
        try:
            header = detectar_encabezado(book.worksheets[0])
            tiene_mvr = 'MVR' in book.sheetnames
        finally:
            book.close()
        rows = preparar_statement(contenido, month, header=header)
    except Exception as exc:
        st.error(texto(f'No se pudo procesar el statement: {type(exc).__name__}: {exc}',
                       f'Could not process the statement: {type(exc).__name__}: {exc}'))
        return
    chargebacks = sum(1 for r in rows if r['is_chargeback'])
    st.caption(texto(
        f'Encabezado detectado en la fila {header}. {len(rows) - chargebacks} registros'
        + (f', {chargebacks} chargebacks (hoja MVR).' if tiene_mvr else '.'),
        f'Header detected on row {header}. {len(rows) - chargebacks} records'
        + (f', {chargebacks} chargebacks (MVR sheet).' if tiene_mvr else '.')))
    key = hashlib.sha256(archivo.name.encode() + contenido + month.encode()).hexdigest()
    mostrar_tabla([{k: v for k, v in r.items() if k != 'source_data'} for r in rows], 'general_preview_' + key, st=st)
    total = sum((importe(r['total_commission_amount']) for r in rows), start=importe('0'))
    st.write(texto(f'Comisión total del statement: {total:.2f}', f'Total statement commission: {total:.2f}'))
    file_id = calcular_file_id(contenido)
    _revisar_duplicados(rows, month, file_id)
    confirmed = st.checkbox(texto('Revisé los registros y confirmo la importación.',
                                  'I reviewed the records and confirm the import.'), key='general_confirm_' + key)
    if st.button(texto('Importar faltantes y continuar', 'Import missing and continue'), disabled=not confirmed,
                 key='general_save_' + key, type='primary'):
        try:
            with st.spinner(texto('Guardando THE GENERAL…', 'Saving THE GENERAL…')):
                inserted, updated, duplicated = guardar(rows, month, archivo.name, contenido)
            mensaje = f'{inserted} filas nuevas, {updated} corregidas.'
            mensaje_en = f'{inserted} new rows, {updated} corrected.'
            if duplicated:
                mensaje += f' {duplicated} ya existían con el mismo contenido este mes (otro archivo); no se duplicaron.'
                mensaje_en += f' {duplicated} already existed with the same content this month (a different file); not duplicated.'
            st.success(texto(mensaje, mensaje_en))
        except Exception as exc:
            st.error(str(exc))


def mostrar(month):
    st.title('THE GENERAL')
    st.caption(texto(
        'El statement siempre llega en Excel (columnas fijas). Si el archivo trae una segunda hoja llamada MVR, '
        'sus filas con Chargeback Applies = Yes se importan junto con las demás, como chargebacks. Puedes subir '
        'uno o varios documentos a la vez; se importan por separado pero se concilian y exportan juntos, más '
        'abajo.', 'The statement always arrives as Excel (fixed columns). If the file has a second sheet named '
        'MVR, its rows with Chargeback Applies = Yes are imported along with the rest, as chargebacks. You can '
        'upload one or more documents at once; they are imported separately but reconciled and exported '
        'together, further down.'))
    st.caption(texto(
        'Franquicia por código de agente (Agent #): directo en la tabla de códigos, salvo que sea el código '
        'master, en cuyo caso se busca la póliza en Compass (con el prefijo de estado y, si no aparece, sin él), '
        'luego en el histórico, luego por el mismo código ya visto, y por último por nombre del cliente en '
        'Compass (búsqueda visual). Las filas de chargeback (hoja MVR) solo resuelven franquicia por su propio '
        'código, sin Compass ni histórico, y son 100% a cargo de la franquicia (igual que COMMONWEALTH).',
        'Franchise by agent code (Agent #): direct in the code table, unless it is the master code, in which '
        'case the policy is looked up in Compass (with the state prefix and, if not found, without it), then in '
        "history, then by the same code already seen, and finally by the client's name in Compass (visual "
        'search). Chargeback rows (MVR sheet) only resolve the franchise by their own code, without Compass or '
        'history, and the franchise is charged 100% of it (same as COMMONWEALTH).'))
    archivos = st.file_uploader(texto('Statements THE GENERAL (Excel)', 'THE GENERAL statements (Excel)'),
                                type=['xlsx', 'xlsm'], accept_multiple_files=True, key='general_files')
    if not archivos:
        return
    if any(f.size > 25 * 1024 * 1024 for f in archivos):
        st.error(texto('Máximo 25 MB por documento.', 'Maximum 25 MB per document.'))
        return
    for archivo in archivos:
        with st.expander(archivo.name, expanded=True):
            _bloque_excel(archivo, month)

    st.divider()
    st.subheader(texto('Conciliación y Excel final del mes', "Month's reconciliation and final Excel"))
    st.caption(texto(
        'Combina todos los statements de THE GENERAL ya guardados para este mes contable, sin importar de qué '
        'archivo vinieron.', "Combines every THE GENERAL statement already saved for this accounting month, "
        'regardless of which file it came from.'))
    accounting_month = datetime.strptime(month, '%Y-%m').date()
    connection = conectar()
    try:
        cursor = connection.cursor()
        cursor.execute('SELECT SUM(total_commission_amount) FROM staging_hub.st_the_general_raw WHERE accounting_month=%s',
                       (accounting_month,))
        total = cursor.fetchone()[0]
        cursor.close()
    finally:
        connection.close()
    if total is None:
        st.info(texto('Importa al menos un statement para calcular las comisiones y conciliar con el banco.',
                      'Import at least one statement to calculate commissions and reconcile with the bank.'))
        return
    snapshot = {'file_id': None, 'rows': [{'accounting_month': accounting_month}], 'grand_total': total}
    mostrar_conciliacion(snapshot, 'general_month_' + month, 'THE GENERAL')
