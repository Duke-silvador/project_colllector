"""Importación revisable de códigos; nombres Soffront desde offices."""
from io import BytesIO
import re
from openpyxl import load_workbook
from database import conectar
from excel_codigos import normalizar_agency_id, texto_celda as _text
import compass


def normalizar_codigo(value, carrier):
    if carrier == 'FLORIDA PENINSULA':
        return normalizar_agency_id(value)
    code = str(value if value is not None else '').strip()
    if not code:
        raise ValueError('Falta el código del carrier.')
    return code


def leer_codigos(data, hoja, header=1, carrier='FLORIDA PENINSULA'):
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        ws = book[hoja]
        cells = next(ws.iter_rows(min_row=header, max_row=header))
        names = [re.sub(r'\s+', '', str(c.value or '')).casefold() for c in cells]
        names = ['dtfanchise' if n == 'dtfranchise' else n for n in names]
        required = ('dtfanchise', 'franchisename', 'code')
        for name in required:
            if names.count(name) != 1:
                raise ValueError(f'Falta o está duplicada la columna {name}.')
        rows = []
        for number, cells in enumerate(ws.iter_rows(min_row=header + 1), header + 1):
            if not any(c.value is not None for c in cells):
                continue
            def get(name):
                return _text(cells[names.index(name)]) if name in names else ''
            franchise = get('dtfanchise').upper()
            if not re.fullmatch(r'DTF0\d+(?:-\d+)?|DT120|DTF1001', franchise):
                raise ValueError(f'Fila {number}: código de franquicia inválido {franchise!r}.')
            name = get('franchisename')
            if not name:
                raise ValueError(f'Fila {number}: falta Franchise Name.')
            rows.append(dict(source_row=number, franchise=franchise, franchise_name=name, carrier=carrier,
                             code=normalizar_codigo(get('code'), carrier), is_master_code=False,
                             office_id='', franchise_soffront_name='', state_code='', user_ID='', agent_name='',
                             agent_licence_220_number='', agent_220_npn=''))
        if not rows:
            raise ValueError('No hay códigos para importar.')
        return rows
    finally:
        book.close()


def cargar_oficinas():
    c = conectar()
    try:
        q = c.cursor(dictionary=True)
        try:
            q.execute('SELECT office_id, office_number, office_name, state FROM staging_hub.offices')
            return q.fetchall()
        finally:
            q.close()
    finally:
        c.close()


def completar_oficinas(rows, offices):
    for row in rows:
        number = re.sub(r'^DTF?', '', row['franchise']).lstrip('0')
        matches = [o for o in offices if str(o.get('office_number') or '').lstrip('0') == number]
        if row.get('office_id'):
            matches = [o for o in offices if str(o['office_id']) == str(row['office_id'])]
        row['office_id'] = str(matches[0]['office_id']) if len(matches) == 1 else ''
        row['franchise_soffront_name'] = matches[0]['office_name'] if len(matches) == 1 else ''
        row['state_code'] = matches[0].get('state') or '' if len(matches) == 1 else ''
    return rows


def completar_agentes(rows):
    token = None
    result = []
    for row in rows:
        uid = row.get('office_id')
        if not uid:
            result.append({**row, 'user_ID': '', 'agent_name': '', 'agent_office_id': '',
                           'agent_status': 'Sin Office ID: no se consultó Compass', 'importar': False,
                           'agent_licence_220_number': '', 'agent_220_npn': ''})
            continue
        if token is None:
            token = compass.obtener_token()
        agents = compass.obtener_agentes_oficina(token, uid)
        if not agents:
            result.append({**row, 'user_ID': '', 'agent_name': '', 'agent_office_id': uid,
                           'agent_status': 'Sin agentes activos', 'importar': row.get('importar', True),
                           'agent_licence_220_number': '', 'agent_220_npn': ''})
            continue
        for agent in agents:
            fields = {re.sub(r'[^a-z0-9]', '', str(k).casefold()): v for k, v in agent.items()}
            def get(*names):
                return next((str(fields[n]).strip() for n in names if fields.get(n) is not None and str(fields[n]).strip()), '')
            user = get('userid', 'agentid', 'id')
            name = get('agentname', 'name', 'fullname') or ' '.join(filter(None, [get('firstname'), get('lastname')]))
            if not user or not name:
                raise ValueError('Compass devolvió un agente sin identificador o nombre reconocible; revisa el formato de la respuesta.')
            result.append({**row, 'user_ID': user, 'agent_name': name, 'agent_office_id': uid,
                'agent_status': 'Agente obtenido de Compass', 'importar': row.get('importar', True),
                'agent_licence_220_number': get('agentlicence220number', 'agentlicense220number', 'license220number'),
                'agent_220_npn': get('agent220npn', 'agent220npnnumber', 'npn')})
    return result


def actualizar_agentes(rows):
    """Actualiza agentes sin volver a enriquecer ni sustituir datos de origen."""
    def identity(row):
        return (row.get('source_row'), row['franchise'], row['code'], row.get('franchise_name'), row.get('office_id'))
    groups = {}
    for row in rows:
        groups.setdefault(identity(row), []).append(row)
    result = []
    for group in groups.values():
        refreshed = completar_agentes([{f: v for f, v in group[0].items() if f != '_agents'}])
        row = {**group[0], **{f: refreshed[0].get(f) for f in
               ('agent_office_id', 'agent_status', 'agent_licence_220_number', 'agent_220_npn')}}
        row['_agents'] = [{f: r.get(f) for f in ('user_ID', 'agent_name', 'agent_office_id',
                          'agent_licence_220_number', 'agent_220_npn')} for r in refreshed if r.get('user_ID')]
        row['agent_name'] = ', '.join(r['agent_name'] for r in row['_agents'])
        row['user_ID'] = ', '.join(r['user_ID'] for r in row['_agents'])
        row['importar'] = group[0].get('importar', True) if row.get('office_id') else False
        result.append(row)
    return result


def importar_codigos(rows, file_name, *, conflictos=None):
    rows = [dict(r) for r in rows]
    for row in rows:
        agents = row.get('_agents')
        if agents:
            if any(a.get('agent_office_id') != row.get('office_id') for a in agents):
                raise ValueError('Cambió la oficina; vuelve a consultar sus agentes en Compass.')
            agents = sorted(agents, key=lambda a: a['user_ID'])
            for field in ('user_ID', 'agent_name', 'agent_licence_220_number', 'agent_220_npn'):
                values = [str(a.get(field) or '') for a in agents]
                row[field] = ', '.join(values) if any(values) else ''
    if not rows:
        raise ValueError('No hay códigos para importar.')
    if any(not r.get('office_id') or ((not r.get('agent_name') or not r.get('user_ID'))
           and r.get('agent_status') != 'Sin agentes activos') for r in rows):
        raise ValueError('Completa la oficina, el User ID y el agente de cada registro antes de importar.')
    if any(r.get('agent_office_id') != r['office_id'] for r in rows):
        raise ValueError('Cambió la oficina; vuelve a consultar sus agentes en Compass antes de importar.')
    if any(not isinstance(r.get('is_master_code'), bool) for r in rows):
        raise ValueError('La selección master debe ser una casilla marcada o desmarcada.')
    c = conectar()
    q = c.cursor(dictionary=True)
    locked = False
    try:
        q.execute("SELECT GET_LOCK('staging_hub.franchises_carrier_codes.import',10) AS acquired")
        locked = q.fetchone()['acquired'] == 1
        if not locked:
            raise ValueError('Otra importación de códigos está en curso.')
        q.execute('SELECT office_id, office_number, office_name, state FROM staging_hub.offices')
        completar_oficinas(rows, q.fetchall())
        if any(not r['office_id'] for r in rows):
            raise ValueError('Hay una oficina que no existe o es ambigua.')
        count = 0
        for row in rows:
            carrier = row.get('carrier', 'FLORIDA PENINSULA')
            code = normalizar_codigo(row['code'], carrier)
            q.execute('SELECT codeID, is_master_code, franchise_name, franchise_soffront_name, agent_name, user_ID,'
                      ' agent_licence_220_number, agent_220_npn FROM staging_hub.franchises_carrier_codes'
                      ' WHERE carrier=%s AND code=%s AND franchise=%s FOR UPDATE',
                      (carrier, code, row['franchise']))
            existing = q.fetchall()
            if existing:
                differences = []
                if len(existing) != 1:
                    differences.append(f'{len(existing)} registros existentes para el mismo código')
                else:
                    if bool(existing[0]['is_master_code']) != row['is_master_code']:
                        differences.append('Master')
                    differences.extend(f for f in ('franchise_name', 'franchise_soffront_name', 'user_ID',
                        'agent_name', 'agent_licence_220_number', 'agent_220_npn')
                        if str(existing[0].get(f) or '') != str(row.get(f) or ''))
                if differences:
                    if conflictos is None:
                        raise ValueError(f'Franquicia {row["franchise"]}, código {code}: diferencias en {", ".join(differences)}.')
                    conflictos.append({'Franquicia': row['franchise'], 'Código': code,
                                       'Nombre del Excel': row['franchise_name'],
                                       'Diferencias': ', '.join(differences)})
                continue
            q.execute('INSERT INTO staging_hub.franchises_carrier_codes'
                      ' (file_name,state_code,franchise,franchise_name,franchise_soffront_name,carrier,code,'
                      ' user_ID,agent_name,agent_licence_220_number,agent_220_npn,master_code_override)'
                      ' VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)',
                      (file_name,row['state_code'],row['franchise'],row['franchise_name'],row['franchise_soffront_name'],
                       carrier,code,row['user_ID'],row['agent_name'],row['agent_licence_220_number'] or None,
                       row['agent_220_npn'] or None,int(row['is_master_code'])))
            count += 1
        c.commit()
        return count
    except Exception:
        c.rollback()
        raise
    finally:
        try:
            if locked:
                q.execute("SELECT RELEASE_LOCK('staging_hub.franchises_carrier_codes.import')")
                q.fetchone()
        finally:
            q.close()
            c.close()
