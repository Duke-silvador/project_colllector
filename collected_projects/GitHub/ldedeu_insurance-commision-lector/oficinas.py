"""Alta, listado e importación masiva de oficinas (staging_hub.offices)."""
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

from database import conectar
from models import Office
from repositories import OFFICE_COLUMNS, OfficeRepository

TABLE = "offices"

_BOOLEANOS_SI = {"SI", "SÍ", "S", "YES", "Y", "TRUE", "1", "X"}
_BOOLEANOS_NO = {"NO", "N", "FALSE", "0", ""}

_CAMPOS_TEXTO = ("office_number", "status", "office_type", "phone_number", "email",
                  "street_address", "city", "state", "zip", "qq_credential")
_CAMPOS_ENTEROS = ("active_users", "total_users")
_CAMPOS_BOOLEANOS = ("is_corporate_hq", "shared_email_account", "sms_upload_enabled")


def _verificar_esquema(connection):
    cursor = connection.cursor(dictionary=True)
    try:
        cursor.execute(
            "SELECT ENGINE FROM information_schema.TABLES"
            " WHERE TABLE_SCHEMA='staging_hub' AND TABLE_NAME=%s",
            (TABLE,),
        )
        row = cursor.fetchone()
        if not row:
            raise ValueError(
                "Falta la tabla staging_hub.offices. Ejecuta sql/crear_oficinas.sql antes de continuar."
            )
        if row["ENGINE"].upper() != "INNODB":
            raise ValueError("staging_hub.offices debe ser InnoDB para poder revertir errores.")
    finally:
        cursor.close()


def _texto(valor):
    texto = str(valor).strip() if valor is not None else ""
    return texto or None


def _booleano(valor):
    if isinstance(valor, bool):
        return valor
    texto = str(valor).strip().upper() if valor is not None else ""
    if texto in _BOOLEANOS_SI:
        return True
    if texto in _BOOLEANOS_NO:
        return False
    raise ValueError(f"'{valor}' no es un valor Sí/No reconocido.")


def _entero(valor):
    if valor is None or str(valor).strip() == "":
        return None
    try:
        return int(float(valor))
    except (TypeError, ValueError):
        raise ValueError(f"'{valor}' no es un número entero válido.") from None


def _decimal(valor):
    if valor is None or str(valor).strip() == "":
        return None
    try:
        numero = Decimal(str(valor))
        if not numero.is_finite():
            raise InvalidOperation
        return numero
    except InvalidOperation:
        raise ValueError(f"'{valor}' no es un número válido.") from None


def _fecha(valor):
    if valor is None or str(valor).strip() == "":
        return None
    if isinstance(valor, datetime):
        return valor
    texto = str(valor).strip()
    for formato in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    raise ValueError(f"'{valor}' no es una fecha reconocida.")


def _construir_registro(datos):
    """datos: dict con las llaves de Office (excepto id/uploaded_at); valores ya crudos
    (texto de formulario o celda de Excel). Normaliza y valida cada campo."""
    office_id = str(datos.get("office_id") or "").strip()
    office_name = str(datos.get("office_name") or "").strip()
    if not office_id:
        raise ValueError("Indica el Office ID.")
    if not office_name:
        raise ValueError("Indica el nombre de la oficina.")
    return Office(
        id=None,
        office_id=office_id,
        office_number=_texto(datos.get("office_number")),
        office_name=office_name,
        status=_texto(datos.get("status")),
        office_type=_texto(datos.get("office_type")),
        is_corporate_hq=_booleano(datos.get("is_corporate_hq", False)),
        phone_number=_texto(datos.get("phone_number")),
        email=_texto(datos.get("email")),
        street_address=_texto(datos.get("street_address")),
        city=_texto(datos.get("city")),
        state=_texto(datos.get("state")),
        zip=_texto(datos.get("zip")),
        monthly_sales_goal=_decimal(datos.get("monthly_sales_goal")),
        active_users=_entero(datos.get("active_users")),
        total_users=_entero(datos.get("total_users")),
        shared_email_account=_booleano(datos.get("shared_email_account", False)),
        sms_upload_enabled=_booleano(datos.get("sms_upload_enabled", False)),
        qq_credential=_texto(datos.get("qq_credential")),
        created_date=_fecha(datos.get("created_date")),
    )


def crear_oficina(datos) -> int:
    registro = _construir_registro(datos)
    connection = conectar()
    try:
        _verificar_esquema(connection)
        repo = OfficeRepository(connection)
        if repo.bloquear_por_office_id(registro.office_id):
            raise ValueError(
                f"Ya existe una oficina con Office ID '{registro.office_id}'. Edítala en vez de crear otra."
            )
        office_row_id = repo.crear(registro)
        connection.commit()
        return office_row_id
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def listar_oficinas():
    connection = conectar()
    try:
        _verificar_esquema(connection)
        return OfficeRepository(connection).listar()
    finally:
        connection.close()


def actualizar_estado_oficina(office_id, status):
    if status not in ('Active', 'Inactive'):
        raise ValueError('Selecciona Activa o Inactiva.')
    connection = conectar()
    cursor = connection.cursor(dictionary=True)
    try:
        cursor.execute('SELECT office_id FROM staging_hub.offices WHERE office_id=%s FOR UPDATE', (office_id,))
        if not cursor.fetchone():
            raise ValueError('La oficina seleccionada ya no existe.')
        cursor.execute('UPDATE staging_hub.offices SET status=%s WHERE office_id=%s', (status, office_id))
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        cursor.close()
        connection.close()


def cargar_alias_franquicias(connection):
    cursor = connection.cursor(dictionary=True)
    try:
        cursor.execute('SELECT office_id, office_number, franchise_alias FROM staging_hub.offices WHERE franchise_alias IS NOT NULL AND TRIM(franchise_alias) <> ""')
        return cursor.fetchall()
    finally:
        cursor.close()


def actualizar_alias_franquicia(office_id, alias):
    alias = str(alias or '').strip() or None
    if alias and len(alias) > 255:
        raise ValueError('El alias admite hasta 255 caracteres.')
    if alias:
        if not re.fullmatch(r'\d+(?:-\d+)*', alias):
            raise ValueError('Escribe el alias sin prefijo, por ejemplo 99 o 144-0024.')
    connection = conectar()
    cursor = connection.cursor()
    try:
        cursor.execute('UPDATE staging_hub.offices SET franchise_alias=%s WHERE office_id=%s', (alias, office_id))
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        cursor.close()
        connection.close()


def cargar_mapa_office_numbers(connection):
    """Devuelve {office_id: office_number} para traducir el office_id que devuelve
    Compass al Office # real de la oficina (ver franquicias.resolver_franquicia)."""
    cursor = connection.cursor(dictionary=True)
    try:
        cursor.execute(
            "SELECT office_id, office_number FROM staging_hub.offices WHERE office_number IS NOT NULL"
        )
        return {row["office_id"]: row["office_number"] for row in cursor.fetchall()}
    finally:
        cursor.close()


_COLUMNAS_EXCEL = {
    "officeid": "office_id",
    "office": "office_number",
    "officenumber": "office_number",
    "officename": "office_name",
    "status": "status",
    "officetype": "office_type",
    "corporatehq": "is_corporate_hq",
    "phonenumber": "phone_number",
    "email": "email",
    "streetaddress": "street_address",
    "city": "city",
    "state": "state",
    "zip": "zip",
    "monthlysalesgoal": "monthly_sales_goal",
    "activeusers": "active_users",
    "totalusers": "total_users",
    "sharedemailaccount": "shared_email_account",
    "smsuploadenabled": "sms_upload_enabled",
    "qqcredential": "qq_credential",
    "createddate": "created_date",
}


def _normalizar_encabezado(valor):
    return re.sub(r"[^a-z0-9]", "", str(valor or "").strip().lower())


def _localizar_columnas(headers, errores):
    normalizados = [_normalizar_encabezado(h) for h in headers]
    indices = {}
    for i, normalizado in enumerate(normalizados):
        campo = _COLUMNAS_EXCEL.get(normalizado)
        if campo and campo not in indices:
            indices[campo] = i
    faltantes = [nombre for nombre in ("office_id", "office_name") if nombre not in indices]
    if faltantes:
        errores.append(f"Faltan columnas para {', '.join(faltantes)}.")
        return None
    return indices


def parsear_excel_oficinas(data):
    """Valida y normaliza un libro de oficinas; no toca MySQL.

    Se toma la primera hoja del libro. La fila de encabezado debe traer, al
    menos, 'Office ID' y 'Office Name'; el resto de columnas (Office #,
    Status, Office Type, Corporate HQ?, Phone Number, Email, Street Address,
    City, State, Zip, Monthly Sales Goal, Active Users, Total Users, Shared
    Email Account, SMS Upload Enabled, QQ Credential, Created Date) son
    opcionales. Los campos Sí/No aceptan Sí/No/Yes/No/True/False/1/0.
    """
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = book[book.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise ValueError("El archivo no contiene filas.")
        errores = []
        indices = _localizar_columnas(rows[0], errores)
        if indices is None:
            raise ValueError(" ".join(dict.fromkeys(errores)))
        filas = []
        vistos = set()
        for row_number, row in enumerate(rows[1:], start=2):
            if all(value is None for value in row):
                continue
            datos = {}
            for campo, indice in indices.items():
                datos[campo] = row[indice] if indice < len(row) else None
            office_id = str(datos.get("office_id") or "").strip()
            if not office_id:
                errores.append(f"Fila {row_number}: falta el Office ID.")
                continue
            if office_id in vistos:
                errores.append(f"Fila {row_number}: el Office ID '{office_id}' ya aparece antes en el archivo.")
                continue
            try:
                registro = _construir_registro(datos)
            except ValueError as exc:
                errores.append(f"Fila {row_number}: {exc}")
                continue
            vistos.add(office_id)
            filas.append(registro)
        if errores:
            raise ValueError(" ".join(dict.fromkeys(errores)))
        if not filas:
            raise ValueError("El archivo no contiene registros para importar.")
        return filas
    finally:
        book.close()


def _mismos_datos(existente, registro):
    return all(existente[campo] == getattr(registro, campo) for campo in OFFICE_COLUMNS)


def importar_excel_oficinas(data):
    filas = parsear_excel_oficinas(data)
    connection = conectar()
    try:
        _verificar_esquema(connection)
        repo = OfficeRepository(connection)
        creadas = actualizadas = sin_cambios = 0
        for registro in filas:
            existente = repo.bloquear_por_office_id(registro.office_id)
            if not existente:
                repo.crear(registro)
                creadas += 1
            elif _mismos_datos(existente, registro):
                sin_cambios += 1
            else:
                repo.actualizar(existente["id"], registro)
                actualizadas += 1
        connection.commit()
        return {"creadas": creadas, "actualizadas": actualizadas, "sin_cambios": sin_cambios, "total": len(filas)}
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
