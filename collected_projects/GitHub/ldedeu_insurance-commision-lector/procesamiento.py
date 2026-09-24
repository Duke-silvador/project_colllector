"""Preparacion y validacion de vistas previas, sin acceso a MySQL."""

from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

DECIMAL_FIELDS = ("premium_amount", "comm_percent", "commission_amount")


def hojas_excel(data):
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def texto(value):
    return "" if value is None else str(value)


def preparar_filas(filas, encabezado):
    """Usa indices para distinguir encabezados vacios o duplicados."""
    if encabezado < 1 or encabezado > len(filas):
        raise ValueError("La fila de encabezados no existe.")
    headers = filas[encabezado - 1]
    labels = [f"{index + 1}: {texto(value) or '(sin nombre)'}" for index, value in enumerate(headers)]
    rows = [row for row in filas[encabezado:] if any(value is not None for value in row)]
    return labels, rows


def mapear_filas(rows, mapping, month):
    return [
        {
            field: (month if field == "accounting_month" else
                    texto(row[index]) if index is not None and index < len(row) else "")
            for field, index in mapping.items()
        }
        for row in rows
    ]


def validar_preview(rows):
    """Valida numeros sin inferir separadores ni transformar porcentajes."""
    errors = []
    if not rows:
        return ["Agrega al menos un registro."]
    for index, row in enumerate(rows, start=1):
        for field in DECIMAL_FIELDS:
            if field == 'premium_amount' and row.get('_chargeback'):
                continue
            try:
                number = Decimal(texto(row.get(field)).strip())
                if not number.is_finite():
                    raise InvalidOperation
            except InvalidOperation:
                errors.append(f"Registro {index}: {field} requiere un numero con punto decimal, sin simbolos ni separadores de miles.")
    return errors
