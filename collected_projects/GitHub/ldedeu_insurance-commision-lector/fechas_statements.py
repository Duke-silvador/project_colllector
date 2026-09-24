"""Edita solamente la celda Statement Date del XML de la primera hoja."""
from io import BytesIO
from datetime import date, datetime
from pathlib import Path
import posixpath
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel


def archivos_excel(carpeta):
    path = Path(carpeta).expanduser().resolve()
    if not path.is_dir():
        raise ValueError('La carpeta no existe.')
    return sorted(p for p in path.iterdir() if p.is_file() and
                  p.suffix.lower() in ('.xlsx', '.xlsm') and not p.name.startswith('~$'))


def preparar_fecha(path, fecha):
    original = Path(path).read_bytes()
    book = load_workbook(BytesIO(original), data_only=False)
    try:
        sheet = book.worksheets[0]
        labels = [cell for row in sheet for cell in row
                  if str(cell.value or '').strip().casefold().rstrip(':').strip() == 'statement date']
        if len(labels) != 1:
            raise ValueError('La primera hoja debe tener un único campo Statement Date.')
        label = labels[0]
        column = label.column + 1
        for merged in sheet.merged_cells.ranges:
            if label.coordinate in merged:
                column = merged.max_col + 1
        target = next((sheet.cell(label.row, c) for c in range(column, sheet.max_column + 1)
                       if sheet.cell(label.row, c).value is not None), None)
        if target is None or target.data_type == 'f':
            raise ValueError('No se encontró una fecha editable junto a Statement Date.')
        if not isinstance(target.value, (date, datetime)):
            try:
                datetime.strptime(str(target.value).strip(), '%m/%d/%Y')
            except ValueError:
                raise ValueError('El valor junto a Statement Date no es una fecha reconocida.') from None
        address, anterior, epoch = target.coordinate, str(target.value), book.epoch
    finally:
        book.close()
    with ZipFile(BytesIO(original)) as source:
        workbook = ET.fromstring(source.read('xl/workbook.xml'))
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        first = workbook.find('s:sheets', ns)[0]
        rid = first.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
        relations = ET.fromstring(source.read('xl/_rels/workbook.xml.rels'))
        relative = next(r.attrib['Target'] for r in relations if r.attrib['Id'] == rid)
        part = relative.lstrip('/') if relative.startswith('/') else posixpath.normpath('xl/' + relative)
        xml = source.read(part).decode('utf-8')
        pattern = r'<c\b[^>]*\br="' + re.escape(address) + r'"[^>]*>.*?</c>'
        matches = list(re.finditer(pattern, xml, re.S))
        if len(matches) != 1:
            raise ValueError('No se pudo identificar de forma única la celda en el Excel.')
        cell = matches[0].group()
        opening = cell[:cell.index('>') + 1]
        tipo = re.search(r'\bt="([^"]+)"', opening)
        kind = tipo.group(1) if tipo else 'n'
        if kind == 'n':
            body = '<v>' + str(to_excel(fecha, epoch)) + '</v>'
        elif kind == 'd':
            body = '<v>' + fecha.isoformat() + '</v>'
        elif kind in ('s', 'inlineStr', 'str'):
            opening = re.sub(r'\s+t="[^"]+"', '', opening)[:-1] + ' t="inlineStr">'
            body = '<is><t>' + fecha.strftime('%m/%d/%Y') + '</t></is>'
        else:
            raise ValueError('Tipo de celda no compatible con una fecha.')
        xml = xml[:matches[0].start()] + opening + body + '</c>' + xml[matches[0].end():]
        result = BytesIO()
        with ZipFile(result, 'w') as output:
            output.comment = source.comment
            for info in source.infolist():
                output.writestr(info, xml.encode('utf-8') if info.filename == part else source.read(info.filename))
    return result.getvalue(), address, anterior
