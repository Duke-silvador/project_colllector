"""Statement THE GENERAL: llega en Excel, con un encabezado que puede ocupar 1 fila (formato
"canonico") o 2-3 filas (grupos como "Add On"/"Product Fee" encima de la fila real de nombres
de columna). En vez de exigir texto exacto, se detecta la fila ancla por 'Year'+'Month' en las
primeras dos columnas y el resto de las columnas se leen por POSICION FIJA (confirmado con dos
statements reales: el orden de columnas es identico en ambos formatos).

st_the_general_raw ya existia con datos reales cargados antes de esta app (columnas crudas:
year, month, agent_master_number, agent_number, agent_name, customer_name,
combined_policy_number, company_number, policy_prefix, policy_number,
policy_prefix_plus_policy_number, state_code, eff_date, exp_date, active_during_month_days,
premium_amount, premium_commission_percent, premium_commission_amount, product_fee_amount,
product_fee_comm_percent, product_fee_comm_amount, total_commission_amount, transaction_desc,
dt_or_dtf); esas columnas y sus nombres NO se tocan. Solo se le agregaron columnas nuevas
(source_row, source_data, is_chargeback, del_toro_percent, franchise_number,
franchise_number_source, office_id, producer_name, term_length, compass_policy_id,
policy_status, lob) via sql/ampliar_the_general_raw.sql.

La poliza canonica para Compass/historico es policy_prefix_plus_policy_number (ej.
'FL8606397'); si Compass no la encuentra, se reintenta con policy_number solo (el numero sin
el prefijo de estado, ej. '8606397'), que ya viene como columna propia del mismo statement.

Premium/comision: si Premium=0 Y Premium Commission %=0, se usa Add On Product Fee en su
lugar (Amount/Comm %) para el calculo de reparto (del_toro_percent, y el premium "usado" que
se computa al leer). Los valores crudos del statement (premium_amount, premium_commission_*,
product_fee_*) nunca se pisan.

Franquicia (codigo = Agent #): codigo directo en franchises_carrier_codes -> si es master
(is_master_code=1, no hardcodeado) o no esta -> Compass por poliza (con/sin prefijo) ->
historico (staging_hub.historic_data_commissions) -> mismo codigo ya visto antes en el lote o
en filas guardadas -> por ultimo, Company Search por nombre del cliente (busqueda visual).
LOB se toma de Compass (line_business_id) si la trae; si no, queda vacia. Reparto: tabla de
tarifas (commission_rates), no pass-through. dt_or_dtf: 'DT' si el codigo es master, 'DTF' en
caso contrario (coincide con la convencion ya usada en los datos existentes).

Hoja MVR (si existe): cada fila con "Chargeback Applies"='Yes' se agrega como una fila mas de
la MISMA tabla, con is_chargeback=1 y sin poliza (como los chargebacks de COMMONWEALTH: 100% a
cargo de la franquicia, sin reparto), y su valor es el MVR Cost en negativo. La franquicia sale
del codigo de la MISMA fila (columna Code), solo por tabla directa -sin Compass, historico ni
busqueda visual.
"""
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO

from openpyxl import load_workbook

import compass
from requests import RequestException
from mysql.connector import Error as MySQLError

from database import conectar
from importacion import nombres_archivo, calcular_file_id
from franquicias import _formatear_codigo_oficina
from calculos import calcular_term_length
from compass_bot_client import buscar_franquicias_por_nombre

CARRIER = 'THE GENERAL'

# Posicion fija (0-based) de cada columna del Excel, confirmada contra dos statements reales
# con encabezados de 1 y de 3 filas: el orden es identico en ambos. La clave es el nombre de
# columna real en st_the_general_raw (sin traducir a otra convencion).
_POSICIONES = {
    'year': 0, 'month': 1, 'agent_master_number': 2, 'agent_number': 3, 'agent_name': 4,
    'customer_name': 5, 'combined_policy_number': 6, 'company_number': 7, 'policy_prefix': 8,
    'policy_number': 9, 'policy_prefix_plus_policy_number': 10, 'state_code': 11,
    'eff_date': 12, 'exp_date': 13, 'active_during_month_days': 14, 'premium_amount': 15,
    'premium_commission_percent': 16, 'premium_commission_amount': 17, 'product_fee_amount': 18,
    'product_fee_comm_percent': 19, 'product_fee_comm_amount': 20, 'total_commission_amount': 21,
}
_ANCHO_MINIMO = max(_POSICIONES.values()) + 1
_COLUMNAS_TEXTO = ('agent_master_number', 'agent_number', 'agent_name', 'customer_name',
                   'combined_policy_number', 'company_number', 'policy_prefix', 'policy_number',
                   'policy_prefix_plus_policy_number', 'state_code')
_COLUMNAS_DECIMAL = ('premium_amount', 'premium_commission_percent', 'premium_commission_amount',
                    'product_fee_amount', 'product_fee_comm_percent', 'product_fee_comm_amount',
                    'total_commission_amount')

FIELDS = ('policy_prefix_plus_policy_number', 'customer_name', 'state_code', 'eff_date', 'exp_date',
          'premium_amount', 'premium_commission_percent', 'total_commission_amount', 'agent_number', 'agent_name')

MVR_FIELDS = ('quote_number', 'insured_name', 'state', 'chargeback_applies', 'mvr_cost', 'code', 'agency_name')


def _normalizar_texto(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def _normalizar_codigo(value):
    """El Code de la hoja MVR trae un espacio no-separable (\xa0) delante, y el Agent #/Agent
    Master Number del statement traen un cero a la izquierda (ej. '090903') que no está en la
    tabla de códigos ni en los datos ya cargados (ej. '90903'): se quita para que coincidan."""
    limpio = re.sub(r'[^A-Za-z0-9]', '', str(value or '')).upper()
    return limpio.lstrip('0') or '0' if limpio.isdigit() else limpio


def _fecha_celda(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f'fecha inválida: {value!r}')


def _decimal_celda(value, campo, fila):
    if value is None or value == '':
        return Decimal('0')
    try:
        numero = Decimal(str(value))
        if not numero.is_finite():
            raise InvalidOperation
        return numero
    except InvalidOperation:
        raise ValueError(f'Fila {fila}: {campo} inválido.') from None


# ---------------------------------------------------------------------------
# Lectura de Excel: hoja principal (posicion fija) y hoja MVR (chargebacks)
# ---------------------------------------------------------------------------

def detectar_encabezado(sheet):
    """La fila real de encabezado varia (1 a 3 filas de grupos encima), pero 'Year' y 'Month'
    siempre estan en las dos primeras columnas de la fila donde arrancan los nombres reales."""
    for index, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        if len(row) >= 2 and str(row[0] or '').strip().casefold() == 'year' and str(row[1] or '').strip().casefold() == 'month':
            return index
    raise ValueError('No se encontró el encabezado (fila con "Year"/"Month") de THE GENERAL.')


def _detectar_encabezado_mvr(sheet):
    for index, row in enumerate(sheet.iter_rows(min_row=1, max_row=5, values_only=True), 1):
        primeros = [str(v or '').strip().casefold() for v in row[:3]]
        if primeros[:1] == ['quote #'] or (len(primeros) > 2 and primeros[2] == 'insured name'):
            return index
    return None


def preparar_statement(data, month, *, header=None):
    """Filas listas para guardar(): claves = columnas reales de st_the_general_raw. Incluye
    tambien las filas de chargeback de la hoja MVR, si existe, con Chargeback Applies='Yes'."""
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', month):
        raise ValueError('Indica el mes contable en formato YYYY-MM.')
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    result = []
    try:
        sheet = book.worksheets[0]
        header = header or detectar_encabezado(sheet)
        numero = header
        for cells in sheet.iter_rows(min_row=header + 1):
            numero += 1
            if all(cell.value is None for cell in cells):
                continue
            if len(cells) < _ANCHO_MINIMO:
                continue
            fila = _fila_principal(cells, numero)
            if fila is not None:
                result.append(fila)
        if 'MVR' in book.sheetnames:
            mvr_sheet = book['MVR']
            mvr_header = _detectar_encabezado_mvr(mvr_sheet)
            if mvr_header is not None:
                # numero+1 en vez de len(result)+1: la hoja principal puede tener filas
                # vacias/descartadas entre medio, por lo que la cantidad de filas validas es
                # menor que el ultimo numero de fila iterado; usar esa cantidad duplicaria un
                # source_row ya usado por la hoja principal.
                result.extend(_filas_mvr(mvr_sheet, mvr_header, numero + 1))
        if not result:
            raise ValueError('No hay filas para importar.')
        return result
    finally:
        book.close()


def _fila_principal(cells, numero):
    valores = [c.value for c in cells]
    row = {campo: valores[pos] for campo, pos in _POSICIONES.items()}
    for campo in _COLUMNAS_TEXTO:
        row[campo] = _normalizar_texto(row[campo])
    row['agent_number'] = _normalizar_codigo(row['agent_number'])
    row['agent_master_number'] = _normalizar_codigo(row['agent_master_number'])
    policy_combinada = row['policy_prefix_plus_policy_number']
    if not policy_combinada or policy_combinada.upper() in ('N/A', '--'):
        return None
    try:
        for campo in _COLUMNAS_DECIMAL:
            row[campo] = _decimal_celda(row[campo], campo, numero)
    except ValueError:
        raise
    if row['premium_amount'] == 0 and row['premium_commission_percent'] == 0:
        premium_usada, tasa_usada = row['product_fee_amount'], row['product_fee_comm_percent']
    else:
        premium_usada, tasa_usada = row['premium_amount'], row['premium_commission_percent']
    try:
        row['eff_date'] = _fecha_celda(row['eff_date'])
        row['exp_date'] = _fecha_celda(row['exp_date'])
    except ValueError as exc:
        raise ValueError(f'Fila {numero}: fecha inválida ({exc}).') from None
    try:
        row['year'] = int(row['year']) if row['year'] not in (None, '') else None
        row['month'] = int(row['month']) if row['month'] not in (None, '') else None
        row['active_during_month_days'] = int(row['active_during_month_days']) if row['active_during_month_days'] not in (None, '') else None
    except (TypeError, ValueError):
        raise ValueError(f'Fila {numero}: Year/Month/# of Days Active inválido.') from None
    row.update(
        source_row=numero,
        source_data={k: (v.isoformat() if isinstance(v, (date, datetime)) else str(v)) for k, v in row.items()},
        del_toro_percent=str(tasa_usada.quantize(Decimal('0.000001'), rounding=ROUND_HALF_UP)),
        term_length=calcular_term_length(row['eff_date'], row['exp_date']),
        producer_code=row['agent_number'], insured_name=row['customer_name'],
        state=row['state_code'] or None,
        is_chargeback=False,
    )
    for campo in _COLUMNAS_DECIMAL:
        row[campo] = str(row[campo])
    return row


def _filas_mvr(sheet, header, numero_inicial):
    labels = ['quote_number', 'policy_number', 'insured_name', 'state', 'chargeback_applies',
              'mvr_cost', 'code', 'agency_name']
    filas = []
    numero = numero_inicial
    for cells in sheet.iter_rows(min_row=header + 1, values_only=True):
        if all(v is None for v in cells):
            continue
        valores = list(cells) + [None] * (len(labels) - len(cells))
        source = dict(zip(labels, valores))
        if _normalizar_texto(source.get('chargeback_applies')).casefold() != 'yes':
            continue
        try:
            costo = _decimal_celda(source.get('mvr_cost'), 'MVR Cost', numero)
        except ValueError:
            raise
        insured = _normalizar_texto(source.get('insured_name'))
        estado = _normalizar_texto(source.get('state')).upper() or None
        filas.append(dict(
            source_row=numero,
            source_data={k: ('' if v is None else str(v)) for k, v in source.items()},
            year=None, month=None, agent_master_number=None,
            agent_number=_normalizar_codigo(source.get('code')), agent_name=_normalizar_texto(source.get('agency_name')) or None,
            customer_name=insured, combined_policy_number=None, company_number=None, policy_prefix=None,
            policy_number=None, policy_prefix_plus_policy_number=None, state_code=estado,
            eff_date=None, exp_date=None, active_during_month_days=None,
            premium_amount='0', premium_commission_percent='1', premium_commission_amount=str(-costo),
            product_fee_amount='0', product_fee_comm_percent='0', product_fee_comm_amount='0',
            total_commission_amount=str(-costo), transaction_desc='MVR',
            del_toro_percent='1', term_length=None,
            producer_code=_normalizar_codigo(source.get('code')), insured_name=insured,
            state=estado, is_chargeback=True,
        ))
        numero += 1
    return filas


def alertas_lectura(row, original=None, threshold=.95):
    alerts = {}
    for field in FIELDS:
        value = str(row.get(field) or '').strip()
        if not value:
            if field == 'agent_name':
                continue
            alerts[field] = 'missing'
            continue
        if original and value == str(original['values'].get(field) or '').strip():
            score = original['confidence'].get(field)
            if score is not None and score < threshold:
                alerts[field] = f'{score:.0%}'
    return alerts


# ---------------------------------------------------------------------------
# Franquicia
# ---------------------------------------------------------------------------

def variantes_poliza_the_general(numero, numero_sin_prefijo):
    """No sabemos si Compass guarda la poliza con el prefijo de estado o sin el; se prueba
    primero completa (como la tenemos, ej. 'FL8606397') y luego solo el numero (columna
    aparte del mismo statement, ej. '8606397')."""
    variantes = [v for v in (str(numero or '').strip(), str(numero_sin_prefijo or '').strip()) if v]
    return list(dict.fromkeys(variantes))


def _franquicias_vistas(connection, rows, producer_code):
    vistas = {str(r.get('franchise_number') or '').strip().upper()
              for r in rows
              if _normalizar_codigo(r.get('producer_code')) == _normalizar_codigo(producer_code) and r.get('franchise_number')}
    cursor = connection.cursor(dictionary=True)
    try:
        cursor.execute('SELECT DISTINCT franchise_number FROM staging_hub.st_the_general_raw'
                       ' WHERE UPPER(agent_number)=UPPER(%s) AND franchise_number IS NOT NULL', (producer_code,))
        vistas |= {str(r['franchise_number']).strip().upper() for r in cursor.fetchall()}
    finally:
        cursor.close()
    vistas.discard('')
    return vistas


def _franquicia_observada(connection, rows, producer_code):
    vistas = _franquicias_vistas(connection, rows, producer_code)
    if len(vistas) == 1:
        return next(iter(vistas)), None
    if len(vistas) > 1:
        return None, sorted(vistas)
    return None, None


def _franquicia_historica(connection, numero):
    cursor = connection.cursor(dictionary=True)
    try:
        cursor.execute(
            "SELECT franchise FROM staging_hub.historic_data_commissions"
            " WHERE carrier=%s AND policy_number=%s AND NULLIF(TRIM(franchise), '') IS NOT NULL"
            ' ORDER BY report_month DESC, date DESC LIMIT 1', (CARRIER, str(numero).strip()))
        registro = cursor.fetchone()
    finally:
        cursor.close()
    return str(registro['franchise']).strip().upper() if registro else None


def completar_franquicias(connection, rows):
    """Filas de poliza real: codigo (Agent #) directo en la tabla -> si es master o no esta,
    Compass por poliza (con y sin prefijo) -> historico -> mismo codigo ya visto -> busqueda
    visual por nombre (el master no equivale a una sola franquicia real: cada poliza puede
    tocar una oficina distinta, por eso SIEMPRE va a Compass en vez de usar la franquicia
    registrada). Filas de MVR (chargeback): solo codigo directo, sin Compass/historico/busqueda
    visual (no tienen poliza que buscar) -> si es el codigo master, usa la franquicia que el
    master tiene registrada en la tabla de codigos (no hay poliza para desambiguar cual de sus
    oficinas reales es), igual que un codigo normal."""
    if not rows:
        return
    cursor = connection.cursor(dictionary=True)
    try:
        cursor.execute('SELECT code, franchise, is_master_code FROM staging_hub.franchises_carrier_codes'
                       ' WHERE UPPER(TRIM(carrier))=%s', (CARRIER,))
        codes = {}
        masters = set()
        franquicia_master = {}
        for entry in cursor.fetchall():
            code = _normalizar_codigo(entry['code'])
            franchise = str(entry.get('franchise') or '').strip().upper()
            if franchise and not franchise.startswith('DT'):
                franchise = _formatear_codigo_oficina(franchise) or franchise
            if entry.get('is_master_code'):
                masters.add(code)
                if franchise:
                    franquicia_master[code] = franchise
                continue
            codes.setdefault(code, set()).add(franchise)
        cursor.execute('SELECT office_id, office_number, office_name, state FROM staging_hub.offices')
        offices = cursor.fetchall()
    finally:
        cursor.close()
    token = None
    polizas = {}

    def oficina_por_numero(numero_oficina):
        return next((o for o in offices if _formatear_codigo_oficina(o['office_number']) == numero_oficina), None)

    def resolver_con_oficina(row, franquicia, source, *, poliza=None):
        office = None
        if poliza is not None:
            office_id = str(poliza.get('office_id') or '').strip()
            office = next((o for o in offices if str(o['office_id']).strip() == office_id), None) if office_id else None
        if office is None:
            office = oficina_por_numero(franquicia)
        row.update(franchise_number=franquicia, franchise_number_source=source,
                   office_id=str(office['office_id']) if office else row.get('office_id'),
                   state=row.get('state') or (office.get('state') if office else None),
                   producer_name=office.get('office_name') if office else None, code_lookup_alert=None)

    def intentar_observado(row, mensaje_no_encontrado):
        franquicia_previa, ambiguas = _franquicia_observada(connection, rows, row['producer_code'])
        if franquicia_previa:
            resolver_con_oficina(row, franquicia_previa, 'codigo_observado')
        elif ambiguas:
            row['code_lookup_alert'] = (f'{row["producer_code"]}: aparece con franquicias distintas en otras filas '
                                        f'({", ".join(ambiguas)}); revisa el código')
        else:
            row['code_lookup_alert'] = mensaje_no_encontrado

    principales = [r for r in rows if not r.get('is_chargeback')]
    chargebacks = [r for r in rows if r.get('is_chargeback')]

    for row in principales:
        code = _normalizar_codigo(row.get('producer_code'))
        row.update(franchise_number=None, franchise_number_source=None, office_id=None,
                   compass_policy_id=None, policy_status=None, producer_name=None, lob=None,
                   dt_or_dtf='DT' if code in masters else 'DTF')
        franquicia_directa = None
        if code not in masters:
            matches = codes.get(code, set())
            franquicia_directa = next(iter(matches)) if len(matches) == 1 and '' not in matches else None
        if franquicia_directa:
            resolver_con_oficina(row, franquicia_directa, 'codigos')
            continue
        numero = row.get('policy_prefix_plus_policy_number')
        if not numero:
            intentar_observado(row, f'{row["producer_code"]}: sin franquicia única por código y sin póliza para consultar Compass')
            continue
        try:
            if token is None:
                token = compass.obtener_token()
            if numero not in polizas:
                encontrada = None
                for variante in variantes_poliza_the_general(numero, row.get('policy_number')):
                    encontrada = compass.buscar_poliza(token, variante)
                    if encontrada is not None:
                        break
                polizas[numero] = encontrada
            poliza = polizas[numero]
            if poliza is None:
                franquicia = _franquicia_historica(connection, numero)
                if franquicia and not franquicia.startswith('DT'):
                    franquicia = _formatear_codigo_oficina(franquicia) or ''
                if franquicia:
                    resolver_con_oficina(row, franquicia, 'historico')
                else:
                    intentar_observado(row, f'{row["producer_code"]}: póliza {numero} sin franquicia en la tabla de códigos, Compass ni histórico')
                continue
            office_id = str(poliza.get('office_id') or '').strip()
            office = next((o for o in offices if str(o['office_id']).strip() == office_id), None)
            franquicia = _formatear_codigo_oficina(office.get('office_number')) if office else None
            row.update(compass_policy_id=poliza.get('policy_id') or poliza.get('id'),
                       policy_status=poliza.get('status_id'), lob=poliza.get('line_business_id'))
            if franquicia:
                resolver_con_oficina(row, franquicia, 'compass', poliza=poliza)
            else:
                intentar_observado(row, f'{row["producer_code"]}: no se pudo determinar la franquicia de la póliza {numero} en Compass')
        except (RequestException, MySQLError, ValueError, TypeError, KeyError):
            intentar_observado(row, f'{row["producer_code"]}: error al consultar Compass/histórico para la póliza {numero}')

    _completar_por_nombre_cliente(principales, offices)

    for row in chargebacks:
        row.update(franchise_number=None, franchise_number_source=None, office_id=None,
                   compass_policy_id=None, policy_status=None, producer_name=None, lob=None,
                   dt_or_dtf='DT' if _normalizar_codigo(row.get('producer_code')) in masters else 'DTF')
        code = _normalizar_codigo(row.get('producer_code'))
        if code in masters:
            franquicia_directa = franquicia_master.get(code)
        else:
            matches = codes.get(code, set())
            franquicia_directa = next(iter(matches)) if len(matches) == 1 and '' not in matches else None
        if franquicia_directa:
            resolver_con_oficina(row, franquicia_directa, 'codigos')
        else:
            row['code_lookup_alert'] = (f'Chargeback: código {row.get("producer_code") or "no identificado"} '
                                        'sin franquicia única registrada para The General')


def _completar_por_nombre_cliente(rows, offices):
    """Ultimo respaldo para filas de poliza (no chargebacks): busqueda visual por nombre del
    cliente en Company Search, cuando ni el codigo ni Compass ni el historico resolvieron
    nada."""
    pendientes = [row for row in rows if not row.get('franchise_number') and str(row.get('insured_name') or '').strip()]
    if not pendientes:
        return
    nombres = list(dict.fromkeys(str(row['insured_name']).strip() for row in pendientes))
    clientes = [{'request_id': str(i), 'client_name': nombre} for i, nombre in enumerate(nombres)]
    try:
        resultados = buscar_franquicias_por_nombre(clientes)
    except ValueError:
        return
    por_nombre = {nombres[int(r['request_id'])]: r for r in resultados}
    for row in pendientes:
        resultado = por_nombre.get(str(row['insured_name']).strip())
        if not resultado or resultado.get('error'):
            continue
        numeros_oficina = {str(o['office_number']).strip() for o in resultado.get('offices', [])}
        if len(numeros_oficina) != 1:
            if len(numeros_oficina) > 1:
                row['code_lookup_alert'] = (f'{row["code_lookup_alert"]} · el cliente aparece en varias franquicias '
                                            f'en Company Search ({", ".join(sorted(numeros_oficina))}); revisa manualmente')
            continue
        franquicia = _formatear_codigo_oficina(next(iter(numeros_oficina)))
        office = next((o for o in offices if _formatear_codigo_oficina(o['office_number']) == franquicia), None)
        if not franquicia or office is None:
            continue
        row.update(franchise_number=franquicia, franchise_number_source='compass_visual',
                   office_id=str(office['office_id']), state=row.get('state') or office.get('state'),
                   producer_name=office.get('office_name'), code_lookup_alert=None)


# ---------------------------------------------------------------------------
# Guardado
# ---------------------------------------------------------------------------

_COLUMNAS_INSERT = (
    'year', 'month', 'dt_or_dtf', 'agent_master_number', 'agent_number', 'agent_name',
    'customer_name', 'combined_policy_number', 'company_number', 'policy_prefix', 'policy_number',
    'policy_prefix_plus_policy_number', 'state_code', 'eff_date', 'exp_date',
    'active_during_month_days', 'premium_amount', 'premium_commission_percent',
    'premium_commission_amount', 'product_fee_amount', 'product_fee_comm_percent',
    'product_fee_comm_amount', 'total_commission_amount', 'transaction_desc', 'is_chargeback',
    'del_toro_percent', 'franchise_number', 'franchise_number_source', 'office_id',
    'producer_name', 'term_length', 'compass_policy_id', 'policy_status', 'lob',
)


def guardar(rows, month, file_name, contenido):
    """Guarda el statement completo (una sola pagina logica; THE GENERAL solo llega en Excel).
    Reintento identico no duplica; una fila corregida se actualiza. file_id se calcula del
    CONTENIDO del archivo (no del nombre). Escribe en las columnas crudas ya existentes de
    st_the_general_raw (con datos reales previos, intactos) y en las derivadas agregadas.

    Si el mismo statement llega de nuevo con OTRO archivo (otro file_id, ej. el usuario lo
    reenvia con otro nombre o una version corregida solo en algunas filas), tampoco se duplica:
    una fila cuyo contenido (poliza+monto, o asegurado+codigo+monto en un MVR) ya existe ese mes
    contable solo se marca con updated_at, sin insertar una copia igual.

    Devuelve (insertadas, corregidas, con_contenido_repetido)."""
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', month):
        raise ValueError('Mes contable invalido; usa YYYY-MM.')
    if not rows:
        raise ValueError('No hay registros para importar.')
    nombres_archivo(file_name)  # valida que el nombre sea utilizable
    file_id = calcular_file_id(contenido)
    source_rows = [row['source_row'] for row in rows]
    if len(set(source_rows)) != len(source_rows):
        raise ValueError('Hay filas con la misma posición de origen.')
    connection = conectar()
    cursor = connection.cursor(dictionary=True)
    locked = False
    try:
        cursor.execute("SELECT GET_LOCK('staging_hub.st_the_general_raw.import',10) AS acquired")
        locked = cursor.fetchone()['acquired'] == 1
        if not locked:
            raise ValueError('Hay otra importación de THE GENERAL en curso.')
        accounting = datetime.strptime(month, '%Y-%m').date()
        payloads = {row['source_row']: json.dumps(row['source_data'], ensure_ascii=False, sort_keys=True, default=str)
                   for row in rows}
        cursor.execute('SELECT id, source_row, source_data FROM staging_hub.st_the_general_raw'
                       ' WHERE file_id=%s AND accounting_month=%s AND source_row IS NOT NULL FOR UPDATE',
                       (file_id, accounting))
        existentes = {r['source_row']: r for r in cursor.fetchall()}
        for numero in existentes:
            if numero not in source_rows:
                raise ValueError('Este archivo tiene menos filas que la versión ya guardada; no se eliminó nada.')

        def sin_cambios(row):
            previo = existentes.get(row['source_row'])
            if previo is None:
                return False
            guardado = json.loads(previo['source_data']) if isinstance(previo['source_data'], str) else previo['source_data']
            return guardado == row['source_data']

        completar_franquicias(connection, [r for r in rows if not sin_cambios(r)])

        # Un mismo statement puede volver a llegar con otro nombre/archivo (otro file_id, aunque
        # el contenido de cada fila sea igual); no se duplica: se busca por contenido entre TODO
        # lo ya guardado ese mes (sin importar el file_id) y, si coincide, solo se marca
        # updated_at en la fila existente en vez de insertar una copia igual.
        cursor.execute('SELECT id, policy_prefix_plus_policy_number, total_commission_amount,'
                       ' customer_name, agent_number, is_chargeback FROM staging_hub.st_the_general_raw'
                       ' WHERE accounting_month=%s', (accounting,))
        por_contenido = {}
        for existente in cursor.fetchall():
            monto = Decimal(str(existente['total_commission_amount'])) if existente['total_commission_amount'] is not None else None
            if existente['is_chargeback']:
                clave = ('chargeback', existente['customer_name'], existente['agent_number'], monto)
            else:
                clave = ('principal', existente['policy_prefix_plus_policy_number'], monto)
            por_contenido.setdefault(clave, existente['id'])

        def clave_contenido(row):
            monto = Decimal(str(row['total_commission_amount'])) if row.get('total_commission_amount') not in (None, '') else None
            if row.get('is_chargeback'):
                return ('chargeback', row.get('insured_name'), row.get('producer_code'), monto)
            return ('principal', row.get('policy_prefix_plus_policy_number'), monto)

        inserted = updated = duplicated = 0
        for row in rows:
            if sin_cambios(row):
                continue
            previo = existentes.get(row['source_row'])
            valores = tuple(row.get(col) for col in _COLUMNAS_INSERT)
            if previo is not None:
                asignaciones = ', '.join(f'{col}=%s' for col in _COLUMNAS_INSERT)
                cursor.execute(
                    f'UPDATE staging_hub.st_the_general_raw SET {asignaciones}, source_data=%s'
                    ' WHERE id=%s', valores + (payloads[row['source_row']], previo['id']))
                updated += 1
                continue
            id_existente = por_contenido.get(clave_contenido(row))
            if id_existente is not None:
                cursor.execute('UPDATE staging_hub.st_the_general_raw SET updated_at=CURRENT_TIMESTAMP WHERE id=%s',
                               (id_existente,))
                duplicated += 1
            else:
                columnas = ', '.join(_COLUMNAS_INSERT)
                marcadores = ', '.join(['%s'] * len(_COLUMNAS_INSERT))
                cursor.execute(
                    f'INSERT INTO staging_hub.st_the_general_raw ({columnas}, source_data, file_id, file_name,'
                    f' accounting_month, source_row) VALUES ({marcadores}, %s, %s, %s, %s, %s)',
                    valores + (payloads[row['source_row']], file_id, file_name, accounting, row['source_row']))
                inserted += 1
        connection.commit()
        return inserted, updated, duplicated
    except Exception:
        connection.rollback()
        raise
    finally:
        try:
            if locked:
                cursor.execute("SELECT RELEASE_LOCK('staging_hub.st_the_general_raw.import')")
                cursor.fetchone()
        finally:
            cursor.close()
            connection.close()


def cargar_comisiones_excel(connection, snapshot, carrier=None, state=None, business_line=None, *, consultar_tipos=False):
    """Sin file_id en el snapshot, combina TODOS los archivos guardados de ese mes contable
    en un solo reparto. Tabla de tarifas (commission_rates), no pass-through. El premium
    "usado" (con el respaldo de Product Fee cuando Premium es cero) se calcula aqui, sin
    tocar los valores crudos guardados."""
    cursor = connection.cursor(dictionary=True)
    try:
        base = (
            'SELECT r.id, r.file_id, r.accounting_month, r.is_chargeback, r.source_row,'
            ' r.policy_prefix_plus_policy_number AS policy_number, r.customer_name AS insured_name,'
            ' r.state_code AS state, r.eff_date AS effective_date, r.exp_date AS expiration_date,'
            ' CASE WHEN r.premium_amount=0 AND r.premium_commission_percent=0'
            '      THEN r.product_fee_amount ELSE r.premium_amount END AS premium_amount,'
            ' r.total_commission_amount AS commission_amount, r.del_toro_percent,'
            ' r.agent_number AS producer_code, r.agent_name, r.dt_or_dtf, r.term_length,'
            ' r.compass_policy_id, r.policy_status, r.lob AS line_business_id,'
            ' r.franchise_number, r.franchise_number_source, r.office_id, r.producer_name'
            ' FROM staging_hub.st_the_general_raw r WHERE r.accounting_month=%s'
        )
        if snapshot.get('file_id'):
            cursor.execute(base + ' AND r.file_id=%s ORDER BY r.id',
                           (snapshot['rows'][0]['accounting_month'], snapshot['file_id']))
        else:
            cursor.execute(base + ' ORDER BY r.id', (snapshot['rows'][0]['accounting_month'],))
        rows = cursor.fetchall()
        if not rows:
            raise ValueError('No hay registros guardados de THE GENERAL para este mes.')
        for row in rows:
            row['carrier'] = CARRIER
            row['transaction_type_source'] = 'statement'
            row['del_toro_percent_source'] = 'statement'
            row['business_line'] = ''
            if row.get('is_chargeback'):
                row['transaction_type'] = 'MVR'
                row['code_lookup_alert'] = None if row.get('franchise_number') else (
                    f'Chargeback: código {row.get("producer_code") or "sin identificar"} sin franquicia única')
            else:
                row['code_lookup_alert'] = None if row.get('franchise_number') else (
                    f"Póliza {row.get('policy_number') or 'sin número'}: sin franquicia en código, Compass, histórico ni búsqueda visual")
        cursor.execute('SELECT state, carrier, transaction_type, business_line, del_toro_percent, franchise_percent'
                       ' FROM staging_hub.commission_rates WHERE carrier=%s', (CARRIER,))
        tarifas = cursor.fetchall()
        from reparto_comisiones import calcular_comisiones_raw
        calculated = calcular_comisiones_raw(rows, tarifas, CARRIER, state, business_line)
        for row in calculated:
            if row.get('is_chargeback'):
                # Los MVR son 100% a cargo de la franquicia (no se reparten): Del Toro y
                # Franquicia reciben exactamente lo que entra, sin diferencia.
                commission = Decimal(str(row['commission_amount']))
                row['comm_percent'] = Decimal('1')
                row['del_toro_commission'] = commission
                row['franchise_percent'] = Decimal('1')
                row['franchise_commission'] = commission
                row['commission_difference'] = Decimal('0')
                row['producer_name'] = row.get('agent_name')
            row['commission_alert'] = '; '.join(filter(None, (row.get('commission_alert'), row.get('code_lookup_alert')))) or None
        return calculated
    finally:
        cursor.close()
