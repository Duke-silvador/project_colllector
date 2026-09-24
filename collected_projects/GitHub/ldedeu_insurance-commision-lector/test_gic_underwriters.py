import unittest
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch
from openpyxl import Workbook
from requests import RequestException
from carriers import CARRIERS
from gic_underwriters import (preparar_statement, detectar_encabezado, porcentaje_fraccion,
                              extraer_tabla_pagina, alertas_lectura, validar, construir_filas,
                              completar_franquicias, guardar, cargar_comisiones_excel, FIELDS)


def libro(policy='04-CIM-000062634', rate=0.1, premium=635.0, commission=63.5,
          insured='LBBROTHERS LLC', eff=datetime(2026, 4, 22), exp=datetime(2027, 4, 22)):
    wb = Workbook()
    ws = wb.active
    ws.append(['Carrier', 'Policy Number', 'Insured Name ', 'Eff Date', 'Exp Date', 'Premium', 'Comm', 'Comm %'])
    ws.append(['Mid-Continent Casualty Company', policy, insured, eff, exp, premium, commission, rate])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def oficina(**over):
    base = dict(office_id='abc', office_number='120', office_name='Del Toro Insurance Agency Inc', state='FL')
    base.update(over)
    return base


class GicExcelTests(unittest.TestCase):
    def test_carrier_registered_for_raw_commission_loading(self):
        self.assertEqual(CARRIERS['GIC Underwriters']['table'], 'staging_hub.st_gic_underwriters_raw')

    def test_header_detection_ignores_the_carrier_column_and_computes_term_length(self):
        """La columna 'Carrier' del Excel (aseguradora real) se ignora: nuestro carrier
        interno siempre es 'GIC Underwriters'. term_length se calcula de Eff/Exp Date del
        propio statement, sin necesitar Compass."""
        rows = preparar_statement(libro(), '2026-08')
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row['policy_number'], '04-CIM-000062634')
        self.assertEqual(row['insured_name'], 'LBBROTHERS LLC')
        self.assertEqual(Decimal(row['premium_amount']), Decimal('635'))
        self.assertEqual(Decimal(row['commission_amount']), Decimal('63.5'))
        self.assertEqual(Decimal(row['del_toro_percent']), Decimal('0.1'))
        self.assertEqual(row['del_toro_percent_source'], 'statement')
        self.assertEqual(row['effective_date'], date(2026, 4, 22))
        self.assertEqual(row['expiration_date'], date(2027, 4, 22))
        self.assertEqual(row['term_length'], 12)
        self.assertNotIn('carrier', row)

    def test_rate_accepts_fraction_or_whole_percentage(self):
        self.assertEqual(porcentaje_fraccion('0.10'), Decimal('0.10'))
        self.assertEqual(porcentaje_fraccion('10'), Decimal('0.10'))
        self.assertEqual(porcentaje_fraccion('10%'), Decimal('0.10'))
        with self.assertRaises(ValueError):
            porcentaje_fraccion('abc')

    def test_policy_number_na_becomes_blank(self):
        rows = preparar_statement(libro(policy='N/A'), '2026-08')
        self.assertEqual(rows[0]['policy_number'], '')

    def test_missing_required_column_is_rejected(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['Policy Number', 'Insured Name', 'Eff Date', 'Premium', 'Comm', 'Comm %'])
        ws.append(['P1', 'JOHN DOE', datetime(2026, 8, 1), 1000.0, 100.0, 0.1])
        buf = BytesIO()
        wb.save(buf)
        with self.assertRaises(ValueError):
            preparar_statement(buf.getvalue(), '2026-08')

    def test_detectar_encabezado_finds_header_row_below_notes(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['Some note before the real header'])
        ws.append(['Policy Number', 'Insured Name', 'Eff Date', 'Exp Date', 'Premium', 'Comm', 'Comm %'])
        buf = BytesIO()
        wb.save(buf)
        book_bytes = buf.getvalue()
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(book_bytes), read_only=True, data_only=True)
        try:
            self.assertEqual(detectar_encabezado(book.worksheets[0]), 2)
        finally:
            book.close()


class GicFranquiciaTests(unittest.TestCase):
    def test_policy_found_in_compass_resolves_directly(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        rows = [dict(policy_number='04-CIM-000062634')]
        with patch('gic_underwriters.compass.obtener_token', return_value='tok'), \
             patch('gic_underwriters.compass.buscar_poliza',
                   return_value={'office_id': 'abc', 'policy_id': 'p1', 'status_id': 'active'}) as buscar:
            completar_franquicias(connection, rows)
        buscar.assert_called_once_with('tok', '04-CIM-000062634')
        cursor.fetchone.assert_not_called()
        self.assertEqual(rows[0]['franchise_number'], 'DTF0120')
        self.assertEqual(rows[0]['franchise_number_source'], 'compass')
        self.assertEqual(rows[0]['state'], 'FL')
        self.assertEqual(rows[0]['compass_policy_id'], 'p1')
        self.assertEqual(rows[0]['policy_status'], 'active')
        self.assertEqual(rows[0]['producer_name'], 'Del Toro Insurance Agency Inc')
        self.assertIsNone(rows[0]['code_lookup_alert'])

    def test_policy_not_in_compass_falls_back_to_historic(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = {'franchise': 'DTF0120'}
        rows = [dict(policy_number='CPL2664452C')]
        with patch('gic_underwriters.compass.obtener_token', return_value='tok'), \
             patch('gic_underwriters.compass.buscar_poliza', return_value=None):
            completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['franchise_number'], 'DTF0120')
        self.assertEqual(rows[0]['franchise_number_source'], 'historico')
        self.assertEqual(rows[0]['producer_name'], 'Del Toro Insurance Agency Inc')
        self.assertIsNone(rows[0]['code_lookup_alert'])

    def test_compass_error_falls_back_to_historic(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = {'franchise': 'DTF0120'}
        rows = [dict(policy_number='CPL2664452C')]
        with patch('gic_underwriters.compass.obtener_token', side_effect=RequestException('sin red')):
            completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['franchise_number'], 'DTF0120')
        self.assertEqual(rows[0]['franchise_number_source'], 'historico')

    def test_policy_not_in_compass_nor_historic_falls_back_to_client_name_search(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = None
        rows = [dict(policy_number='CPL9999999', insured_name='Omar Gonzalez')]
        with patch('gic_underwriters.compass.obtener_token', return_value='tok'), \
             patch('gic_underwriters.compass.buscar_poliza', return_value=None), \
             patch('gic_underwriters.buscar_franquicias_por_nombre', return_value=[
                 {'request_id': '0', 'client_name': 'Omar Gonzalez',
                  'offices': [{'label': 'Samy (120)', 'office_number': '120'}],
                  'matched_name': 'Omar Gonzalez', 'error': None}]) as buscar:
            completar_franquicias(connection, rows)
        buscar.assert_called_once_with([{'request_id': '0', 'client_name': 'Omar Gonzalez'}])
        self.assertEqual(rows[0]['franchise_number'], 'DTF0120')
        self.assertEqual(rows[0]['franchise_number_source'], 'compass_visual')
        self.assertEqual(rows[0]['producer_name'], 'Del Toro Insurance Agency Inc')
        self.assertIsNone(rows[0]['code_lookup_alert'])

    def test_ambiguous_client_name_search_is_flagged_not_guessed(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = None
        rows = [dict(policy_number='CPL9999999', insured_name='Omar Gonzalez')]
        with patch('gic_underwriters.compass.obtener_token', return_value='tok'), \
             patch('gic_underwriters.compass.buscar_poliza', return_value=None), \
             patch('gic_underwriters.buscar_franquicias_por_nombre', return_value=[
                 {'request_id': '0', 'client_name': 'Omar Gonzalez',
                  'offices': [{'label': 'A (120)', 'office_number': '120'}, {'label': 'B (73)', 'office_number': '73'}],
                  'matched_name': 'Omar Gonzalez', 'error': None}]):
            completar_franquicias(connection, rows)
        self.assertIsNone(rows[0]['franchise_number'])
        self.assertIn('varias franquicias', rows[0]['code_lookup_alert'])

    def test_policy_not_found_anywhere_is_flagged_not_guessed(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = None
        rows = [dict(policy_number='CPL0000000', insured_name='Nadie')]
        with patch('gic_underwriters.compass.obtener_token', return_value='tok'), \
             patch('gic_underwriters.compass.buscar_poliza', return_value=None), \
             patch('gic_underwriters.buscar_franquicias_por_nombre', return_value=[
                 {'request_id': '0', 'client_name': 'Nadie', 'offices': [], 'matched_name': None,
                  'error': 'client_not_found'}]):
            completar_franquicias(connection, rows)
        self.assertIsNone(rows[0]['franchise_number'])
        self.assertTrue(rows[0]['code_lookup_alert'])

    def test_missing_policy_number_is_flagged_without_querying_anything(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        rows = [dict(policy_number='')]
        with patch('gic_underwriters.compass.obtener_token') as token:
            completar_franquicias(connection, rows)
        token.assert_not_called()
        cursor.fetchone.assert_not_called()
        self.assertIsNone(rows[0]['franchise_number'])
        self.assertTrue(rows[0]['code_lookup_alert'])


class GicGuardarTests(unittest.TestCase):
    def rows(self):
        return preparar_statement(libro(), '2026-08')

    def test_save_inserts_new_row_and_resolves_franchise(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.side_effect = [{'acquired': 1}, {'franchise': 'DTF0120'}, None]
        cursor.fetchall.side_effect = [[], [oficina()]]
        with patch('gic_underwriters.conectar', return_value=connection), \
             patch('gic_underwriters.compass.obtener_token', side_effect=RequestException('sin red')):
            self.assertEqual(guardar(self.rows(), '2026-08', 'GIC.xlsx', b'contenido-gic', 1, None), (1, 0))
            insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
            self.assertEqual(insert.args[1][11], 'DTF0120')  # franchise_number resuelto por historico

    def test_retry_with_same_content_does_not_duplicate(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.side_effect = [{'acquired': 1}, {'franchise': 'DTF0120'}, None]
        cursor.fetchall.side_effect = [[], [oficina()]]
        with patch('gic_underwriters.conectar', return_value=connection), \
             patch('gic_underwriters.compass.obtener_token', side_effect=RequestException('sin red')):
            rows = self.rows()
            guardar(rows, '2026-08', 'GIC.xlsx', b'contenido-gic', 1, None)
            insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
            saved_payload = insert.args[1][15]
            source_row = rows[0]['source_row']

            cursor.execute.reset_mock()
            cursor.fetchone.side_effect = [{'acquired': 1}, None]
            cursor.fetchall.side_effect = [[{'id': 1, 'source_row': source_row, 'source_data': saved_payload}]]
            self.assertEqual(guardar(rows, '2026-08', 'GIC.xlsx', b'contenido-gic', 1, None), (0, 0))

    def test_changed_row_updates_instead_of_duplicating(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.side_effect = [{'acquired': 1}, {'franchise': 'DTF0120'}, None]
        cursor.fetchall.side_effect = [[], [oficina()]]
        with patch('gic_underwriters.conectar', return_value=connection), \
             patch('gic_underwriters.compass.obtener_token', side_effect=RequestException('sin red')):
            rows = self.rows()
            guardar(rows, '2026-08', 'GIC.xlsx', b'contenido-gic', 1, None)
            insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
            saved_payload = insert.args[1][15]
            source_row = rows[0]['source_row']

            cursor.execute.reset_mock()
            cursor.fetchone.side_effect = [{'acquired': 1}, {'franchise': 'DTF0120'}, None]
            cursor.fetchall.side_effect = [[{'id': 1, 'source_row': source_row, 'source_data': saved_payload}], [oficina()]]
            changed = preparar_statement(libro(insured='OTHER NAME'), '2026-08')
            self.assertEqual(guardar(changed, '2026-08', 'GIC.xlsx', b'contenido-gic', 1, None), (0, 1))
            update = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('UPDATE'))
            self.assertEqual(update.args[1][1], 'OTHER NAME')

    def test_save_uses_content_hash_as_file_id_not_the_file_name(self):
        """Regresion: el mismo statement subido con otro nombre de archivo debe reconocerse
        como el mismo import (mismo file_id, calculado del contenido), no como uno nuevo."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.side_effect = [{'acquired': 1}, {'franchise': 'DTF0120'}, None]
        cursor.fetchall.side_effect = [[], [oficina()]]
        with patch('gic_underwriters.conectar', return_value=connection), \
             patch('gic_underwriters.compass.obtener_token', side_effect=RequestException('sin red')):
            rows = self.rows()
            guardar(rows, '2026-08', 'nombre-original.xlsx', b'mismo-contenido', 1, None)
            insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
            file_id_guardado = insert.args[1][-5]
            saved_payload = insert.args[1][15]
            source_row = rows[0]['source_row']

            cursor.execute.reset_mock()
            cursor.fetchone.side_effect = [{'acquired': 1}, None]
            cursor.fetchall.side_effect = [[{'id': 1, 'source_row': source_row, 'source_data': saved_payload}]]
            self.assertEqual(
                guardar(rows, '2026-08', 'nombre-reenviado (1).xlsx', b'mismo-contenido', 1, None), (0, 0))
            select = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('SELECT id'))
            self.assertEqual(select.args[1][0], file_id_guardado)

    def test_save_rejects_invalid_month(self):
        with self.assertRaises(ValueError):
            guardar(self.rows(), '2026-13', 'GIC.xlsx', b'contenido-gic', 1, None)

    def test_save_rejects_empty_rows(self):
        with self.assertRaises(ValueError):
            guardar([], '2026-08', 'GIC.xlsx', b'contenido-gic', 1, None)


class GicCommissionSplitTests(unittest.TestCase):
    def commission_row(self, **overrides):
        return dict({'id': 1, 'policy_number': '04-CIM-000062634', 'insured_name': 'X',
                     'effective_date': date(2026, 4, 22), 'expiration_date': date(2027, 4, 22),
                     'franchise_number': 'DTF0120', 'accounting_month': date(2026, 8, 1),
                     'premium_amount': Decimal('635.00'), 'del_toro_percent': Decimal('0.100000'),
                     'commission_amount': Decimal('63.50'), 'term_length': 12, 'state': 'FL'}, **overrides)

    def test_franchise_gets_exact_same_percent_as_del_toro_pass_through(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.return_value = [self.commission_row()]
        snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 8, 1)}]}
        result = cargar_comisiones_excel(connection, snapshot)[0]
        self.assertEqual(result['franchise_percent'], Decimal('0.100000'))
        self.assertEqual(result['franchise_commission'], Decimal('63.50'))
        self.assertEqual(result['del_toro_commission'], Decimal('63.50'))
        self.assertIsNone(result['commission_alert'])
        self.assertEqual(result['carrier'], 'GIC Underwriters')

    def test_unresolved_franchise_is_alerted(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.return_value = [self.commission_row(franchise_number=None)]
        snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 8, 1)}]}
        result = cargar_comisiones_excel(connection, snapshot)[0]
        self.assertIsNone(result['franchise_number'])
        self.assertTrue(result['commission_alert'])

    def test_no_rows_raises(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.return_value = []
        snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 8, 1)}]}
        with self.assertRaises(ValueError):
            cargar_comisiones_excel(connection, snapshot)


class GicPdfExtractionTests(unittest.TestCase):
    def _blocks(self):
        headers = [('Policy', .02), ('Number', .04), ('Insured', .12), ('Name', .14), ('Eff', .22), ('Date', .24),
                   ('Exp', .32), ('Date', .34), ('Premium', .42), ('Comm', .50), ('Comm%', .58)]
        blocks = [dict(text=text, x=x, y=.1, height=.01, confidence=.99) for text, x in headers]
        values = [('04-CIM-000062634', .03, .2), ('LBBROTHERS', .12, .2), ('LLC', .14, .2),
                  ('4/22/2026', .23, .2), ('4/22/2027', .33, .2), ('635.00', .42, .2),
                  ('63.50', .50, .2), ('0.1', .58, .2)]
        blocks.extend(dict(text=text, x=x, y=y, height=.01, confidence=.97) for text, x, y in values)
        return blocks

    def test_extraer_tabla_pagina_distinguishes_comm_from_comm_percent(self):
        """'Comm' y 'Comm %' normalizan a la misma palabra si se descarta el simbolo '%'; el
        extractor de GIC conserva '%' precisamente para poder distinguirlas."""
        filas = extraer_tabla_pagina(self._blocks())
        self.assertEqual(len(filas), 1)
        fila = filas[0]
        self.assertEqual(fila['policy_number'], '04-CIM-000062634')
        self.assertEqual(fila['commission_amount'], '63.50')
        self.assertEqual(fila['del_toro_percent'], '0.1')
        self.assertEqual(fila['premium_amount'], '635.00')
        self.assertIn('LBBROTHERS', fila['insured_name'])

    def test_no_header_returns_no_rows(self):
        self.assertEqual(extraer_tabla_pagina([dict(text='random', x=.1, y=.1, height=.01)]), [])

    def test_alertas_lectura_flags_missing_and_invalid_fields(self):
        row = dict.fromkeys(FIELDS, '')
        row.update(policy_number='P1', insured_name='X', effective_date='2/30/2026',
                   expiration_date='8/1/2027', premium_amount='NaN', commission_amount='100',
                   del_toro_percent='0.10')
        alerts = alertas_lectura(row)
        self.assertIn('effective_date', alerts)
        self.assertIn('premium_amount', alerts)

    def test_validar_and_construir_filas_produce_canonical_rows(self):
        row = dict(policy_number='04-CIM-000062634', insured_name='LBBROTHERS LLC',
                   effective_date='4/22/2026', expiration_date='4/22/2027', premium_amount='635.00',
                   commission_amount='63.50', del_toro_percent='0.1')
        validar([row], '2026-08')
        filas = construir_filas([row], '2026-08')
        self.assertEqual(filas[0]['policy_number'], '04-CIM-000062634')
        self.assertEqual(filas[0]['effective_date'], date(2026, 4, 22))
        self.assertEqual(filas[0]['expiration_date'], date(2027, 4, 22))
        self.assertEqual(filas[0]['term_length'], 12)
        self.assertEqual(Decimal(filas[0]['del_toro_percent']), Decimal('0.1'))

    def test_validar_rejects_missing_required_field(self):
        row = dict(policy_number='', insured_name='X', effective_date='4/22/2026',
                   expiration_date='4/22/2027', premium_amount='635.00', commission_amount='63.50',
                   del_toro_percent='0.1')
        with self.assertRaises(ValueError):
            validar([row], '2026-08')
