import unittest
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch
from openpyxl import Workbook
from requests import RequestException
from carriers import CARRIERS
from the_general import (preparar_statement, detectar_encabezado, completar_franquicias,
                         guardar, cargar_comisiones_excel, _normalizar_codigo,
                         variantes_poliza_the_general)

HEADERS = ['Year', 'Month', 'Agent Master Number', 'Agent #', 'Agent Name', 'Customer Name',
           'Combined Policy #', 'Company Number', 'Policy Prefix', 'Policy Number',
           'Policy Prefix + Policy #', 'State', 'Eff Date', 'Exp Date',
           '# of Days Active During Month', 'Premium', 'Premium Commission %',
           'Premium Commission', 'Add On Product Fee Amount', 'Add On Product Fee Comm %',
           'Add On Product Fee Comm Amt', 'Total Comm Amount']


def _fila(agent_number='91479', agent_master='91409', agent_name='SUMI ALL INSURANCE LLC',
          customer='LAZARO M RAMOS', policy_prefix='FL', policy_number='8606397',
          state='FL', eff=datetime(2026, 5, 13), exp=datetime(2026, 11, 13),
          premium=223.48, premium_rate=0.10, premium_comm=22.35,
          fee_amount=0.0, fee_rate=0.0, fee_comm=0.0, total=22.35):
    combined = policy_prefix + policy_number
    return [2026, 8, agent_master, agent_number, agent_name, customer,
            '1Z-' + policy_prefix + '-' + policy_number, '1Z', policy_prefix, policy_number,
            combined, state, eff, exp, 26, premium, premium_rate, premium_comm,
            fee_amount, fee_rate, fee_comm, total]


def libro(filas=None, mvr=None):
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for fila in (filas if filas is not None else [_fila()]):
        ws.append(fila)
    if mvr is not None:
        hoja = wb.create_sheet('MVR')
        hoja.append(['Quote #', 'Policy #', 'Insured Name', 'DL', 'Chargeback', 'MVR Cost', 'Code', 'Agency name'])
        hoja.append([None, None, None, 'State', 'Applies', None, None, None])
        for fila in mvr:
            hoja.append(fila)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def libro_encabezado_multifila():
    wb = Workbook()
    ws = wb.active
    ws.append(['EARNED REPORT'] + [None] * 20)
    ws.append([None] * 12 + ['Add On'] + [None] * 9)
    ws.append(HEADERS)
    ws.append(_fila())
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def oficina(**over):
    base = dict(office_id='abc', office_number='15', office_name='Sumi All Insurance LLC', state='FL')
    base.update(over)
    return base


class TheGeneralExcelTests(unittest.TestCase):
    def test_carrier_registered_for_raw_commission_loading(self):
        self.assertEqual(CARRIERS['THE GENERAL']['table'], 'staging_hub.st_the_general_raw')

    def test_single_row_header_parses_by_fixed_position(self):
        rows = preparar_statement(libro(), '2026-08')
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row['agent_number'], '91479')
        self.assertEqual(row['customer_name'], 'LAZARO M RAMOS')
        self.assertEqual(row['policy_number'], '8606397')
        self.assertEqual(row['policy_prefix_plus_policy_number'], 'FL8606397')
        self.assertEqual(row['state_code'], 'FL')
        self.assertEqual(row['eff_date'], date(2026, 5, 13))
        self.assertEqual(row['exp_date'], date(2026, 11, 13))
        self.assertEqual(Decimal(row['total_commission_amount']), Decimal('22.35'))
        self.assertEqual(Decimal(row['del_toro_percent']), Decimal('0.1'))
        self.assertEqual(row['term_length'], 6)
        self.assertFalse(row['is_chargeback'])

    def test_multi_row_grouped_header_is_detected_by_year_month_anchor(self):
        rows = preparar_statement(libro_encabezado_multifila(), '2026-08')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['policy_prefix_plus_policy_number'], 'FL8606397')

    def test_agent_number_leading_zero_is_normalized_away(self):
        """El Agent #/Agent Master Number del statement trae cero a la izquierda
        (ej. '090903'), pero franchises_carrier_codes y los datos ya cargados lo guardan sin
        ese cero (ej. '90903'); deben quedar iguales para que coincidan."""
        rows = preparar_statement(libro([_fila(agent_number='090903', agent_master='090870')]), '2026-08')
        self.assertEqual(rows[0]['agent_number'], '90903')
        self.assertEqual(rows[0]['agent_master_number'], '90870')

    def test_premium_and_rate_zero_falls_back_to_product_fee(self):
        rows = preparar_statement(libro([_fila(premium=0, premium_rate=0, premium_comm=0,
                                               fee_amount=12.5, fee_rate=0.5, fee_comm=6.25, total=6.25)]), '2026-08')
        row = rows[0]
        self.assertEqual(Decimal(row['del_toro_percent']), Decimal('0.5'))
        self.assertEqual(Decimal(row['total_commission_amount']), Decimal('6.25'))

    def test_missing_combined_policy_number_row_is_skipped(self):
        filas = [_fila(policy_prefix='', policy_number=''), _fila()]
        rows = preparar_statement(libro(filas), '2026-08')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['policy_prefix_plus_policy_number'], 'FL8606397')

    def test_mvr_sheet_rows_become_chargebacks_only_when_applies_is_yes(self):
        mvr = [
            (709900, None, 'AMANDA MELENDEZ', 'FL', 'Yes', 6.44, '\xa0091483', 'LUCKYCOBRA.LLC'),
            (709901, None, 'SOMEONE ELSE', 'FL', 'No', 4.0, '091483', 'LUCKYCOBRA.LLC'),
        ]
        rows = preparar_statement(libro(mvr=mvr), '2026-08')
        chargebacks = [r for r in rows if r['is_chargeback']]
        self.assertEqual(len(chargebacks), 1)
        row = chargebacks[0]
        self.assertEqual(row['producer_code'], '91483')  # nbsp y cero a la izquierda quitados
        self.assertEqual(row['insured_name'], 'AMANDA MELENDEZ')
        self.assertEqual(Decimal(row['total_commission_amount']), Decimal('-6.44'))
        self.assertIsNone(row['policy_prefix_plus_policy_number'])

    def test_mvr_and_main_sheet_never_share_a_source_row(self):
        rows = preparar_statement(libro(filas=[_fila()], mvr=[(1, None, 'X', 'FL', 'Yes', 1.0, '91483', 'Y')]), '2026-08')
        self.assertEqual(len({r['source_row'] for r in rows}), len(rows))

    def test_detectar_encabezado_finds_year_month_anchor(self):
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(libro_encabezado_multifila()), read_only=True, data_only=True)
        try:
            self.assertEqual(detectar_encabezado(book.worksheets[0]), 3)
        finally:
            book.close()


class TheGeneralNormalizacionTests(unittest.TestCase):
    def test_normalizar_codigo_strips_leading_zero_and_nbsp(self):
        self.assertEqual(_normalizar_codigo('090903'), '90903')
        self.assertEqual(_normalizar_codigo('\xa0091483'), '91483')
        self.assertEqual(_normalizar_codigo('0'), '0')

    def test_variantes_poliza_tries_prefix_then_plain(self):
        self.assertEqual(variantes_poliza_the_general('FL8606397', '8606397'), ['FL8606397', '8606397'])
        self.assertEqual(variantes_poliza_the_general('FL8606397', ''), ['FL8606397'])


class TheGeneralFranquiciaTests(unittest.TestCase):
    def _codigos(self, master=False, franchise='DTF0015'):
        return [{'code': '91479', 'franchise': franchise, 'is_master_code': master}]

    def test_non_master_code_resolves_directly_without_compass(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [self._codigos(), [oficina()]]
        rows = [dict(producer_code='91479', policy_prefix_plus_policy_number='FL8606397', policy_number='8606397',
                     insured_name='X', state=None, is_chargeback=False)]
        with patch('the_general.compass.obtener_token') as token:
            completar_franquicias(connection, rows)
        token.assert_not_called()
        self.assertEqual(rows[0]['franchise_number'], 'DTF0015')
        self.assertEqual(rows[0]['franchise_number_source'], 'codigos')
        self.assertEqual(rows[0]['dt_or_dtf'], 'DTF')
        self.assertEqual(rows[0]['producer_name'], 'Sumi All Insurance LLC')

    def test_master_code_tries_compass_with_prefix_first_then_without(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [self._codigos(master=True), [oficina()]]
        rows = [dict(producer_code='91479', policy_prefix_plus_policy_number='FL8606397', policy_number='8606397',
                     insured_name='X', state=None, is_chargeback=False)]
        with patch('the_general.compass.obtener_token', return_value='tok'), \
             patch('the_general.compass.buscar_poliza',
                   side_effect=[None, {'office_id': 'abc', 'policy_id': 'p1', 'status_id': 'active'}]) as buscar:
            completar_franquicias(connection, rows)
        self.assertEqual([c.args[1] for c in buscar.call_args_list], ['FL8606397', '8606397'])
        self.assertEqual(rows[0]['franchise_number'], 'DTF0015')
        self.assertEqual(rows[0]['franchise_number_source'], 'compass')
        self.assertEqual(rows[0]['dt_or_dtf'], 'DT')

    def test_master_code_not_in_compass_falls_back_to_historico(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [self._codigos(master=True), [oficina()]]
        cursor.fetchone.return_value = {'franchise': 'DTF0015'}
        rows = [dict(producer_code='91479', policy_prefix_plus_policy_number='FL8606397', policy_number='8606397',
                     insured_name='X', state=None, is_chargeback=False)]
        with patch('the_general.compass.obtener_token', return_value='tok'), \
             patch('the_general.compass.buscar_poliza', return_value=None):
            completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['franchise_number'], 'DTF0015')
        self.assertEqual(rows[0]['franchise_number_source'], 'historico')

    def test_master_code_error_falls_back_to_client_name_search(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [self._codigos(master=True),
                                      [oficina(office_number='120', office_name='Del Toro Insurance Agency Inc')], []]
        cursor.fetchone.return_value = None
        rows = [dict(producer_code='91479', policy_prefix_plus_policy_number='FL8606397', policy_number='8606397',
                     insured_name='Lazaro Ramos', state=None, is_chargeback=False)]
        with patch('the_general.compass.obtener_token', side_effect=RequestException('sin red')), \
             patch('the_general.buscar_franquicias_por_nombre', return_value=[
                 {'request_id': '0', 'offices': [{'office_number': '120'}], 'error': None}]) as buscar:
            completar_franquicias(connection, rows)
        buscar.assert_called_once_with([{'request_id': '0', 'client_name': 'Lazaro Ramos'}])
        self.assertEqual(rows[0]['franchise_number'], 'DTF0120')
        self.assertEqual(rows[0]['franchise_number_source'], 'compass_visual')

    def test_unresolved_row_is_flagged_not_guessed(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [self._codigos(master=True), [oficina()], []]
        cursor.fetchone.return_value = None
        rows = [dict(producer_code='91479', policy_prefix_plus_policy_number='FL8606397', policy_number='8606397',
                     insured_name='Nadie', state=None, is_chargeback=False)]
        with patch('the_general.compass.obtener_token', side_effect=RequestException('sin red')), \
             patch('the_general.buscar_franquicias_por_nombre', return_value=[
                 {'request_id': '0', 'offices': [], 'error': 'client_not_found'}]):
            completar_franquicias(connection, rows)
        self.assertIsNone(rows[0]['franchise_number'])
        self.assertTrue(rows[0]['code_lookup_alert'])

    def test_chargeback_row_resolves_only_by_its_own_code_no_compass(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [self._codigos(franchise='DTF0044'), [oficina()]]
        rows = [dict(producer_code='91479', policy_prefix_plus_policy_number=None, policy_number=None,
                     insured_name='Amanda Melendez', state='FL', is_chargeback=True)]
        with patch('the_general.compass.obtener_token') as token, \
             patch('the_general.buscar_franquicias_por_nombre') as nombre:
            completar_franquicias(connection, rows)
        token.assert_not_called()
        nombre.assert_not_called()
        self.assertEqual(rows[0]['franchise_number'], 'DTF0044')
        self.assertIsNone(rows[0].get('franchise_percent'))

    def test_chargeback_row_with_master_code_uses_its_registered_franchise(self):
        """A diferencia de las filas con poliza (que SIEMPRE van a Compass para el codigo
        master, porque cada poliza puede tocar una oficina real distinta), un MVR no tiene
        poliza que buscar: si el codigo es el master, usa directo la franquicia que el master
        tiene registrada en la tabla de codigos (ej. 90903 -> DTF0120), sin ir a Compass."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [self._codigos(master=True, franchise='DTF0120'), [oficina(office_number='120')]]
        rows = [dict(producer_code='91479', policy_prefix_plus_policy_number=None, policy_number=None,
                     insured_name='Amanda Melendez', state='FL', is_chargeback=True)]
        with patch('the_general.compass.obtener_token') as token, \
             patch('the_general.buscar_franquicias_por_nombre') as nombre:
            completar_franquicias(connection, rows)
        token.assert_not_called()
        nombre.assert_not_called()
        self.assertEqual(rows[0]['franchise_number'], 'DTF0120')
        self.assertEqual(rows[0]['franchise_number_source'], 'codigos')

    def test_chargeback_row_without_registered_code_is_flagged(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = []
        rows = [dict(producer_code='00000', policy_number=None, policy_number_plain=None,
                     insured_name='Nadie', state='FL', is_chargeback=True)]
        completar_franquicias(connection, rows)
        self.assertIsNone(rows[0]['franchise_number'])
        self.assertIn('sin franquicia única registrada', rows[0]['code_lookup_alert'])


class TheGeneralGuardarTests(unittest.TestCase):
    def rows(self):
        return preparar_statement(libro(), '2026-08')

    def test_save_inserts_new_row_and_resolves_franchise(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.side_effect = [{'acquired': 1}, None]
        cursor.fetchall.side_effect = [[], [{'code': '91479', 'franchise': 'DTF0015', 'is_master_code': 0}], [oficina()], []]
        with patch('the_general.conectar', return_value=connection):
            self.assertEqual(guardar(self.rows(), '2026-08', 'THE_GENERAL.xlsx', b'contenido'), (1, 0, 0))
        insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
        self.assertIn('DTF0015', insert.args[1])

    def test_retry_with_same_content_does_not_duplicate(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.side_effect = [{'acquired': 1}, None]
        cursor.fetchall.side_effect = [[], [{'code': '91479', 'franchise': 'DTF0015', 'is_master_code': 0}], [oficina()], []]
        with patch('the_general.conectar', return_value=connection):
            rows = self.rows()
            guardar(rows, '2026-08', 'THE_GENERAL.xlsx', b'contenido')
            insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
            saved_payload = insert.args[1][-5]
            source_row = rows[0]['source_row']

            cursor.execute.reset_mock()
            cursor.fetchone.side_effect = [{'acquired': 1}, None]
            cursor.fetchall.side_effect = [[{'id': 1, 'source_row': source_row, 'source_data': saved_payload}], []]
            self.assertEqual(guardar(rows, '2026-08', 'THE_GENERAL.xlsx', b'contenido'), (0, 0, 0))

    def test_save_uses_content_hash_as_file_id_not_the_file_name(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.side_effect = [{'acquired': 1}, None]
        cursor.fetchall.side_effect = [[], [{'code': '91479', 'franchise': 'DTF0015', 'is_master_code': 0}], [oficina()], []]
        with patch('the_general.conectar', return_value=connection):
            rows = self.rows()
            guardar(rows, '2026-08', 'original.xlsx', b'mismo-contenido')
            insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
            saved_payload = insert.args[1][-5]
            source_row = rows[0]['source_row']

            cursor.execute.reset_mock()
            cursor.fetchone.side_effect = [{'acquired': 1}, None]
            cursor.fetchall.side_effect = [[{'id': 1, 'source_row': source_row, 'source_data': saved_payload}], []]
            self.assertEqual(guardar(rows, '2026-08', 'reenviado (1).xlsx', b'mismo-contenido'), (0, 0, 0))

    def test_different_file_same_content_same_month_is_not_duplicated(self):
        """El mismo statement puede reenviarse con otro archivo (otro file_id, ej. corregido o
        con otro nombre); si el contenido de una fila (poliza+monto) ya esta guardado ese mes
        contable, no se inserta de nuevo: solo se marca updated_at en la fila existente."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        rows = self.rows()
        cursor.fetchone.side_effect = [{'acquired': 1}, None]
        cursor.fetchall.side_effect = [
            [],  # existentes para este file_id: ninguno, es "otro archivo"
            [{'code': '91479', 'franchise': 'DTF0015', 'is_master_code': 0}], [oficina()],
            [{'id': 55, 'policy_prefix_plus_policy_number': rows[0]['policy_prefix_plus_policy_number'],
              'total_commission_amount': Decimal(rows[0]['total_commission_amount']),
              'customer_name': None, 'agent_number': None, 'is_chargeback': 0}],
        ]
        with patch('the_general.conectar', return_value=connection):
            self.assertEqual(guardar(rows, '2026-08', 'otro_archivo.xlsx', b'contenido-distinto'), (0, 0, 1))
        toque = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('UPDATE staging_hub.st_the_general_raw SET updated_at'))
        self.assertEqual(toque.args[1], (55,))
        self.assertFalse(any(c.args[0].startswith('INSERT') for c in cursor.execute.call_args_list))

    def test_save_rejects_invalid_month(self):
        with self.assertRaises(ValueError):
            guardar(self.rows(), '2026-13', 'THE_GENERAL.xlsx', b'contenido')

    def test_save_rejects_empty_rows(self):
        with self.assertRaises(ValueError):
            guardar([], '2026-08', 'THE_GENERAL.xlsx', b'contenido')


class TheGeneralCommissionSplitTests(unittest.TestCase):
    def commission_row(self, **overrides):
        return dict({'id': 1, 'policy_number': 'FL8606397', 'insured_name': 'X', 'is_chargeback': 0,
                     'effective_date': date(2026, 5, 13), 'expiration_date': date(2026, 11, 13),
                     'franchise_number': 'DTF0015', 'accounting_month': date(2026, 8, 1),
                     'premium_amount': Decimal('223.48'), 'del_toro_percent': Decimal('0.100000'),
                     'commission_amount': Decimal('22.35'), 'term_length': 6, 'state': 'FL'}, **overrides)

    def test_rate_table_split_not_pass_through(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.side_effect = [
            [self.commission_row()],
            [{'state': 'FL', 'carrier': 'THE GENERAL', 'transaction_type': 'ALL', 'business_line': '',
              'del_toro_percent': 0.10, 'franchise_percent': 0.08}],
        ]
        snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 8, 1)}]}
        result = cargar_comisiones_excel(connection, snapshot)[0]
        self.assertEqual(result['franchise_percent'], Decimal('0.08'))
        self.assertEqual(result['franchise_commission'], Decimal('223.48') * Decimal('0.08'))
        self.assertIsNone(result['commission_alert'])
        self.assertEqual(result['carrier'], 'THE GENERAL')

    def test_chargeback_row_is_100_percent_pass_through_to_the_franchise(self):
        """Los MVR se cargan integros a la franquicia (no se reparten): Franchise % = 1,
        Franchise $ = exactamente la comision que entra, Transaction Type = 'MVR', y Producer
        Name muestra el Agency Name de la hoja MVR (no la oficina resuelta)."""
        connection = MagicMock()
        connection.cursor.return_value.fetchall.side_effect = [
            [self.commission_row(policy_number=None, is_chargeback=1, franchise_number='DTF0044',
                                 commission_amount=Decimal('-6.44'), agent_name='Luckycobra.Llc')],
            [],
        ]
        snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 8, 1)}]}
        result = cargar_comisiones_excel(connection, snapshot)[0]
        self.assertEqual(result['franchise_number'], 'DTF0044')
        self.assertEqual(result['del_toro_commission'], Decimal('-6.44'))
        self.assertEqual(result['franchise_percent'], Decimal('1'))
        self.assertEqual(result['franchise_commission'], Decimal('-6.44'))
        self.assertEqual(result['commission_difference'], Decimal('0'))
        self.assertEqual(result['comm_percent'], Decimal('1'))
        self.assertEqual(result['transaction_type'], 'MVR')
        self.assertEqual(result['producer_name'], 'Luckycobra.Llc')
        self.assertIsNone(result['commission_alert'])

    def test_unresolved_franchise_is_alerted(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.side_effect = [[self.commission_row(franchise_number=None)], []]
        snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 8, 1)}]}
        result = cargar_comisiones_excel(connection, snapshot)[0]
        self.assertIsNone(result['franchise_number'])
        self.assertTrue(result['commission_alert'])

    def test_no_rows_raises(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.return_value = []
        snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 8, 1)}]}
        with self.assertRaises(ValueError):
            cargar_comisiones_excel(connection, snapshot)

    def test_select_includes_accounting_month(self):
        """Regresion: mostrar_conciliacion vuelve a leer accounting_month de la PRIMERA fila
        ya calculada (snapshot['rows'][0]['accounting_month']) para el resto de la pantalla;
        si el SELECT no la trae, revienta con KeyError en la interfaz aunque los tests con
        filas simuladas a mano no lo detecten."""
        connection = MagicMock()
        connection.cursor.return_value.fetchall.side_effect = [[self.commission_row()], []]
        snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 8, 1)}]}
        cargar_comisiones_excel(connection, snapshot)
        select_sql = connection.cursor.return_value.execute.call_args_list[0].args[0]
        self.assertIn('accounting_month', select_sql)


if __name__ == '__main__':
    unittest.main()
