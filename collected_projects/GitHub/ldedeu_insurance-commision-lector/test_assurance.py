import unittest
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch
from openpyxl import Workbook
from carriers import CARRIERS
from assurance import (preparar_statement, detectar_encabezado, guardar_statement,
                        completar_franquicias, cargar_comisiones_excel, variantes_poliza_assurance)


def libro(policy='RFL0001', producer='FL14295', amount=29.52, premium=246.0, type_='Endorse'):
    wb = Workbook()
    ws = wb.active
    ws.append(['StatementDate', 'StatementPeriod', 'Producer', 'AddDate', 'InsuredName', 'PolicyNumber',
               'Premiums', 'Type', 'TransEffDate', 'PolicyExpDate', 'Amount'])
    ws.append([datetime(2026, 2, 2), datetime(2026, 1, 1), producer, None, 'JOSE PEREZ', policy,
               premium, type_, datetime(2026, 1, 22), datetime(2026, 3, 15), amount])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def libro_periodos_mixtos():
    wb = Workbook()
    ws = wb.active
    ws.append(['StatementDate', 'StatementPeriod', 'Producer', 'AddDate', 'InsuredName', 'PolicyNumber',
               'Premiums', 'Type', 'TransEffDate', 'PolicyExpDate', 'Amount'])
    ws.append([datetime(2026, 9, 2), datetime(2026, 8, 1), 'FL14295', None, 'A', 'RFL0001',
               100.0, 'Endorse', datetime(2026, 8, 5), datetime(2026, 12, 1), 12.0])
    ws.append([datetime(2026, 9, 2), datetime(2026, 9, 1), 'FL14295', None, 'B', 'RFL0002',
               200.0, 'Endorse', datetime(2026, 9, 5), datetime(2026, 12, 1), 24.0])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class AssuranceTests(unittest.TestCase):
    def test_accounting_month_is_the_typed_month_not_each_row_statement_period(self):
        """Regresion: un statement puede traer filas de StatementPeriod distinto (ajustes
        tardios). accounting_month debe ser el mes contable escrito al importar, siempre,
        para que el reparto encuentre todas las filas del archivo sin importar StatementPeriod."""
        rows = preparar_statement(libro_periodos_mixtos(), '2026-09')
        self.assertEqual({r['statement_period'] for r in rows}, {date(2026, 8, 1), date(2026, 9, 1)})
        self.assertTrue(all(r['accounting_month'] == '2026-09' for r in rows))

    def test_carrier_registered_for_raw_commission_loading(self):
        self.assertEqual(CARRIERS['ASSURANCE']['table'], 'staging_hub.st_assurance_raw')

    def test_header_detection_and_parsing_computes_rate_from_amount_over_premiums(self):
        rows = preparar_statement(libro(), '2026-01')
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row['policy_number'], 'RFL0001')
        self.assertEqual(row['producer_code'], 'FL14295')
        self.assertEqual(row['premium_amount'], '246')
        self.assertEqual(row['commission_amount'], '29.52')
        self.assertEqual(Decimal(row['del_toro_percent']), Decimal('29.52') / Decimal('246'))
        self.assertEqual(row['del_toro_percent_source'], 'statement')

    def test_na_policy_becomes_empty_and_zero_premium_adjustment_gets_rate_one(self):
        rows = preparar_statement(libro(policy='N/A', producer='FL13221', amount=-8.0, premium=0.0,
                                        type_='Commission - U/W Rprt Chargeback (Adjustment)'), '2026-01')
        row = rows[0]
        self.assertEqual(row['policy_number'], '')
        self.assertEqual(row['del_toro_percent'], '1.000000')

    def test_header_detection_skips_leading_notes_rows(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['Generated report - confidential'])
        ws.append(['StatementDate', 'StatementPeriod', 'Producer', 'AddDate', 'InsuredName', 'PolicyNumber',
                   'Premiums', 'Type', 'TransEffDate', 'PolicyExpDate', 'Amount'])
        self.assertEqual(detectar_encabezado(ws), 2)

    def test_rejects_invalid_month_and_missing_columns(self):
        with self.assertRaises(ValueError):
            preparar_statement(libro(), '2026-13')
        wb = Workbook()
        wb.active.append(['StatementDate', 'Producer'])
        buf = BytesIO()
        wb.save(buf)
        with self.assertRaises(ValueError):
            preparar_statement(buf.getvalue(), '2026-01')

    def test_non_master_code_resolves_directly_from_codes_table(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [
            [{'code': 'FL14295', 'franchise': 'DTF0106', 'is_master_code': 0}],
            [{'office_id': 'abc', 'office_number': '106', 'state': 'FL'}],
        ]
        cursor.fetchone.return_value = None
        rows = [dict(producer_code='FL14295', policy_number='RFL0001', insured_name='X')]
        completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['franchise_number'], 'DTF0106')
        self.assertEqual(rows[0]['franchise_number_source'], 'codigos')
        self.assertEqual(rows[0]['state'], 'FL')
        self.assertIsNone(rows[0]['code_lookup_alert'])

    def test_code_missing_from_table_falls_back_to_compass_by_policy(self):
        """Regresion: FL13726 no esta en franchises_carrier_codes (no es master, simplemente
        no se cargo todavia). Antes esto se quedaba sin franquicia sin intentar Compass;
        ahora cualquier codigo sin franquicia unica por tabla cae a Compass/historico."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [
            [{'code': 'FL11679', 'franchise': '', 'is_master_code': 1}],
            [{'office_id': 'c64e0e16', 'office_number': '149', 'state': 'FL'}],
        ]
        cursor.fetchone.return_value = None
        rows = [dict(producer_code='FL13726', policy_number='RFL2761543', insured_name='X')]
        poliza = {'office_id': 'c64e0e16', 'policy_id': 'pid', 'status_id': 'active'}
        with patch('assurance.compass.obtener_token', return_value='tok'), \
             patch('assurance.compass.buscar_poliza', return_value=poliza):
            completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['franchise_number'], 'DTF0149')
        self.assertEqual(rows[0]['franchise_number_source'], 'compass')
        self.assertEqual(rows[0]['state'], 'FL')
        self.assertIsNone(rows[0]['code_lookup_alert'])

    def test_unresolved_policy_borrows_franchise_from_same_code_resolved_earlier_in_batch(self):
        """Ejemplo real: TX15792 es master. RTX2799427 (misma tanda) ya resolvio a DTF0155
        via Compass; RTX2809214 no aparece en Compass ni en historico. En vez de quedar sin
        franquicia, se toma DTF0155 porque es la unica que ese codigo ha mostrado."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [
            [{'code': 'TX15792', 'franchise': '', 'is_master_code': 1}],
            [{'office_id': 'oficina-155', 'office_number': '155', 'state': 'TX'}],
            [],  # _franquicia_observada: nada guardado antes en la base para este codigo
        ]
        cursor.fetchone.return_value = None
        rows = [
            dict(producer_code='TX15792', policy_number='RTX2799427', insured_name='ROBERTO'),
            dict(producer_code='TX15792', policy_number='RTX2809214', insured_name='JOSE'),
        ]
        poliza_resuelta = {'office_id': 'oficina-155', 'policy_id': 'p1', 'status_id': 'active'}
        def buscar_poliza(token, numero):
            return poliza_resuelta if numero == 'RTX2799427' else None
        with patch('assurance.compass.obtener_token', return_value='tok'), \
             patch('assurance.compass.buscar_poliza', side_effect=buscar_poliza), \
             patch('assurance.buscar_franquicia_historica', return_value=None):
            completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['franchise_number'], 'DTF0155')
        self.assertEqual(rows[0]['franchise_number_source'], 'compass')
        self.assertEqual(rows[1]['franchise_number'], 'DTF0155')
        self.assertEqual(rows[1]['franchise_number_source'], 'codigo_observado')
        self.assertEqual(rows[1]['state'], 'TX')
        self.assertIsNone(rows[1]['code_lookup_alert'])

    def test_unresolved_policy_borrows_franchise_from_previously_saved_rows(self):
        """Mismo respaldo pero cuando la fila 'ancla' ya resuelta viene de una importacion
        anterior (guardada en la base), no de esta misma tanda."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [
            [{'code': 'TX15792', 'franchise': '', 'is_master_code': 1}],
            [{'office_id': 'oficina-155', 'office_number': '155', 'state': 'TX'}],
            [{'franchise_number': 'DTF0155'}],  # ya guardado antes en st_assurance_raw
        ]
        cursor.fetchone.return_value = None
        rows = [dict(producer_code='TX15792', policy_number='RTX2809214', insured_name='JOSE')]
        with patch('assurance.compass.obtener_token', return_value='tok'), \
             patch('assurance.compass.buscar_poliza', return_value=None), \
             patch('assurance.buscar_franquicia_historica', return_value=None):
            completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['franchise_number'], 'DTF0155')
        self.assertEqual(rows[0]['franchise_number_source'], 'codigo_observado')
        self.assertEqual(rows[0]['state'], 'TX')
        self.assertIsNone(rows[0]['code_lookup_alert'])

    def test_code_seen_with_conflicting_franchises_is_flagged_not_guessed(self):
        """Si el mismo codigo aparece con dos franquicias distintas en filas ya guardadas,
        no se adivina ninguna: se marca como ambiguo para revision manual."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [
            [{'code': 'TX15792', 'franchise': '', 'is_master_code': 1}],
            [{'office_id': 'oficina-155', 'office_number': '155', 'state': 'TX'}],
            [{'franchise_number': 'DTF0155'}, {'franchise_number': 'DTF0133'}],
        ]
        cursor.fetchone.return_value = None
        rows = [dict(producer_code='TX15792', policy_number='RTX9999999', insured_name='X')]
        with patch('assurance.compass.obtener_token', return_value='tok'), \
             patch('assurance.compass.buscar_poliza', return_value=None), \
             patch('assurance.buscar_franquicia_historica', return_value=None):
            completar_franquicias(connection, rows)
        self.assertIsNone(rows[0]['franchise_number'])
        self.assertIn('franquicias distintas', rows[0]['code_lookup_alert'])
        self.assertIn('DTF0133', rows[0]['code_lookup_alert'])
        self.assertIn('DTF0155', rows[0]['code_lookup_alert'])

    def test_master_code_without_policy_is_alerted_without_calling_compass(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [
            [{'code': 'FL11679', 'franchise': '', 'is_master_code': 1}],
            [],
            [],
        ]
        cursor.fetchone.return_value = None
        rows = [dict(producer_code='FL11679', policy_number='', insured_name='X')]
        with patch('assurance.compass.obtener_token') as token:
            completar_franquicias(connection, rows)
            token.assert_not_called()
        self.assertIn('sin póliza para consultar Compass', rows[0]['code_lookup_alert'])

    def test_master_code_falls_back_to_historic_when_compass_has_no_policy(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [
            [{'code': 'FL11679', 'franchise': '', 'is_master_code': 1}],
            [{'office_id': 'xyz', 'office_number': '120', 'state': 'FL'}],
        ]
        cursor.fetchone.return_value = {'agent_name': 'Jael Perez', 'term_length': '6'}
        rows = [dict(producer_code='FL11679', policy_number='RFL2334092-2', insured_name='NELSON')]
        with patch('assurance.compass.obtener_token', return_value='tok'), \
             patch('assurance.compass.buscar_poliza', return_value=None), \
             patch('assurance.buscar_franquicia_historica', return_value='DTF0120'):
            completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['franchise_number'], 'DTF0120')
        self.assertEqual(rows[0]['franchise_number_source'], 'historico')
        self.assertEqual(rows[0]['agent_name'], 'Jael Perez')
        self.assertEqual(rows[0]['term_length'], '6')

    def test_non_master_code_matches_office_by_franchise_alias_sub_location(self):
        """DTF0188-0112 no es un office_number real, es el alias de la oficina 188
        (sub-location). El match debe caer al alias cuando el numero plano no calza."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [
            [{'code': 'FL99999', 'franchise': 'DTF0188-0112', 'is_master_code': 0}],
            [{'office_id': 'oficina-188', 'office_number': '188', 'state': 'FL', 'franchise_alias': '0188-0112'}],
        ]
        cursor.fetchone.return_value = None
        rows = [dict(producer_code='FL99999', policy_number='RFL0001', insured_name='X')]
        completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['office_id'], 'oficina-188')
        self.assertEqual(rows[0]['state'], 'FL')
        self.assertIsNone(rows[0]['code_lookup_alert'])

    def test_policy_number_variants_try_no_suffix_and_zero_padded_editions(self):
        variantes = variantes_poliza_assurance('PFL2502455-2')
        self.assertEqual(variantes[0], 'PFL2502455-2')
        self.assertIn('PFL2502455', variantes)
        self.assertIn('PFL2502455-01', variantes)
        self.assertIn('PFL2502455-02', variantes)
        self.assertIn('PFL2502455-1', variantes)
        self.assertEqual(len(variantes), len(set(variantes)))

    def test_policy_number_variants_decrement_before_bare_before_unpadded(self):
        """Ejemplo real: llega con dos digitos (XXXXX-02). Orden esperado: tal cual, decrementando
        con el mismo formato de dos digitos hasta 0, sin sufijo, y despues decrementando otra vez
        sin ceros a la izquierda -antes de intentar incrementar-."""
        variantes = variantes_poliza_assurance('XXXXX-02')
        self.assertEqual(variantes[:5], ['XXXXX-02', 'XXXXX-01', 'XXXXX-00', 'XXXXX', 'XXXXX-2'])
        self.assertEqual(len(variantes), len(set(variantes)))

    def test_master_code_tries_variants_in_order_and_stops_at_first_compass_match(self):
        """PFL2502455-2 no existe en Compass tal cual, pero PFL2502455-02 si; debe
        encontrarla sin necesidad del historico y sin validar el nombre del asegurado."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [
            [{'code': 'FL11679', 'franchise': '', 'is_master_code': 1}],
            [{'office_id': '54d735b4', 'office_number': '106', 'state': 'FL'}],
        ]
        cursor.fetchone.return_value = None
        rows = [dict(producer_code='FL11679', policy_number='PFL2502455-2', insured_name='ALGUIEN')]
        poliza = {'office_id': '54d735b4', 'policy_id': 'p1', 'status_id': 'active'}
        intentos = []
        def buscar_poliza(token, numero):
            intentos.append(numero)
            return poliza if numero == 'PFL2502455-02' else None
        with patch('assurance.compass.obtener_token', return_value='tok'), \
             patch('assurance.compass.buscar_poliza', side_effect=buscar_poliza):
            completar_franquicias(connection, rows)
        self.assertEqual(intentos, variantes_poliza_assurance('PFL2502455-2')[:len(intentos)])
        self.assertEqual(rows[0]['franchise_number'], 'DTF0106')
        self.assertEqual(rows[0]['franchise_number_source'], 'compass')
        self.assertIsNone(rows[0]['code_lookup_alert'])

    def test_save_inserts_skips_identical_and_updates_changed_row(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.return_value = {'acquired': 1}
        cursor.fetchall.side_effect = [[], [{'code': 'FL14295', 'franchise': 'DTF0106', 'is_master_code': 0}], []]
        with patch('assurance.conectar', return_value=connection):
            rows = preparar_statement(libro(), '2026-01')
            self.assertEqual(guardar_statement(rows, 'Statement.xlsx', b'contenido-assurance'), (1, 0))
            insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
            self.assertIn('accounting_month', insert.args[0])
            self.assertEqual(insert.args[1][4], date(2026, 1, 1))
            saved_payload = insert.args[1][16]

            cursor.fetchone.side_effect = None
            cursor.fetchone.return_value = {'acquired': 1}
            cursor.fetchall.side_effect = [[{'id': 1, 'source_row': 2, 'source_data': saved_payload}]]
            same_rows = preparar_statement(libro(), '2026-01')
            self.assertEqual(guardar_statement(same_rows, 'Statement.xlsx', b'contenido-assurance'), (0, 0))

            cursor.fetchall.side_effect = [[{'id': 1, 'source_row': 2, 'source_data': saved_payload}],
                [{'code': 'FL14295', 'franchise': 'DTF0106', 'is_master_code': 0}], []]
            changed_rows = preparar_statement(libro(policy='RFL9999'), '2026-01')
            self.assertEqual(guardar_statement(changed_rows, 'Statement.xlsx', b'contenido-assurance'), (0, 1))
            update = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('UPDATE'))
            self.assertEqual(update.args[1][4], 'RFL9999')
            self.assertEqual(update.args[1][-1], 1)

    def test_save_blocks_when_file_has_fewer_rows_than_saved(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.return_value = {'acquired': 1}
        cursor.fetchall.side_effect = [[{'id': 1, 'source_row': 5, 'source_data': '{}'}]]
        rows = preparar_statement(libro(), '2026-01')
        with patch('assurance.conectar', return_value=connection), self.assertRaises(ValueError):
            guardar_statement(rows, 'Statement.xlsx', b'contenido-assurance')

    def test_save_uses_content_hash_as_file_id_not_the_file_name(self):
        """Regresion: el mismo statement subido con otro nombre de archivo debe reconocerse
        como el mismo import (mismo file_id, calculado del contenido), no como uno nuevo."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.return_value = {'acquired': 1}
        cursor.fetchall.side_effect = [[], [{'code': 'FL14295', 'franchise': 'DTF0106', 'is_master_code': 0}], []]
        with patch('assurance.conectar', return_value=connection):
            rows = preparar_statement(libro(), '2026-01')
            guardar_statement(rows, 'nombre-original.xlsx', b'mismo-contenido')
            insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
            file_id_guardado = insert.args[1][2]
            saved_payload = insert.args[1][16]

            cursor.fetchall.side_effect = [[{'id': 1, 'source_row': 2, 'source_data': saved_payload}]]
            same_rows = preparar_statement(libro(), '2026-01')
            self.assertEqual(guardar_statement(same_rows, 'nombre-reenviado (1).xlsx', b'mismo-contenido'), (0, 0))
            select = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('SELECT id'))
            self.assertEqual(select.args[1][0], file_id_guardado)

    def test_commission_split_computes_franchise_percent_from_matching_rate(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        row = dict(id=1, policy_number='RFL0001', producer_code='FL14295', insured_name='X',
                   transaction_type='Endorse', franchise_number='DTF0106', state='FL',
                   premium_amount=Decimal('246.00'), commission_amount=Decimal('29.52'),
                   del_toro_percent=Decimal('0.120000'), accounting_month=date(2026, 1, 1))
        cursor.fetchall.side_effect = [
            [row],
            [{'state': 'FL', 'carrier': 'ASSURANCE', 'transaction_type': 'ALL', 'business_line': '',
              'del_toro_percent': Decimal('0.12'), 'franchise_percent': Decimal('0.10')}],
        ]
        snapshot = {'file_id': 'Statement.xlsx', 'rows': [{'accounting_month': date(2026, 1, 1)}]}
        result = cargar_comisiones_excel(connection, snapshot)
        self.assertIsNone(result[0]['commission_alert'])
        self.assertEqual(result[0]['franchise_percent'], Decimal('0.10'))
        self.assertEqual(result[0]['franchise_commission'], Decimal('246.00') * Decimal('0.10'))

    def test_commission_split_treats_na_policy_as_chargeback_passthrough(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        row = dict(id=2, policy_number='N/A', producer_code='FL11679', insured_name='X',
                   transaction_type='Commission - U/W Rprt Chargeback (Adjustment)', franchise_number=None,
                   state=None, premium_amount=Decimal('0'), commission_amount=Decimal('-8.00'),
                   del_toro_percent=Decimal('1.000000'), accounting_month=date(2026, 1, 1))
        cursor.fetchall.side_effect = [[row], []]
        snapshot = {'file_id': 'Statement.xlsx', 'rows': [{'accounting_month': date(2026, 1, 1)}]}
        result = cargar_comisiones_excel(connection, snapshot)
        self.assertIsNone(result[0]['policy_number'])
        self.assertIsNone(result[0]['franchise_percent'])
        self.assertEqual(result[0]['commission_amount'], Decimal('-8.00'))

    def test_commission_split_without_file_id_combines_every_file_for_the_month(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        row_a = dict(id=1, policy_number='RFL0001', producer_code='FL14295', insured_name='A',
                     transaction_type='Endorse', franchise_number='DTF0106', state='FL',
                     premium_amount=Decimal('100.00'), commission_amount=Decimal('12.00'),
                     del_toro_percent=Decimal('0.12'), accounting_month=date(2026, 8, 1))
        row_b = dict(id=2, policy_number='RFL0002', producer_code='FL14295', insured_name='B',
                     transaction_type='Endorse', franchise_number='DTF0106', state='FL',
                     premium_amount=Decimal('200.00'), commission_amount=Decimal('24.00'),
                     del_toro_percent=Decimal('0.12'), accounting_month=date(2026, 8, 1))
        cursor.fetchall.side_effect = [[row_a, row_b], []]
        snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 8, 1)}]}
        result = cargar_comisiones_excel(connection, snapshot)
        select_sql = cursor.execute.call_args_list[0].args[0]
        self.assertNotIn('file_id', select_sql)
        self.assertEqual(len(result), 2)
        self.assertEqual({r['policy_number'] for r in result}, {'RFL0001', 'RFL0002'})


if __name__ == '__main__':
    unittest.main()
