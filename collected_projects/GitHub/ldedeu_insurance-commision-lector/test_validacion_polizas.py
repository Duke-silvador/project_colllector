import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from validacion_polizas import validar_polizas


class ValidacionPolizasTests(unittest.TestCase):
    def archivo(self, principal, junio, julio):
        libro = Workbook()
        libro.active.title = "Compass"
        libro.create_sheet("June")
        libro.create_sheet("Julio")
        for hoja, filas in zip(libro.worksheets, (principal, junio, julio)):
            for fila in filas:
                hoja.append(fila)
        salida = BytesIO()
        libro.save(salida)
        return salida.getvalue()

    def test_union_faltantes_vacios_y_conservacion(self):
        data = self.archivo(
            [("Policy Number", "Cliente"), (" A-1 ", "Ana"), ("B2", "Luis"),
             ("003", "Eva"), (None, "Sin número"), (None, None), (123, "Sol"), ("A-1", "Repetida")],
            [("Policy",), ("a-1",), ("123",)], [("Policy",), ("B2",), ("3",)])
        salida, detalle = validar_polizas(data)
        self.assertEqual([r["Estado"] for r in detalle],
                         ["Encontrada", "Encontrada", "No encontrada", "Sin Policy Number", "Encontrada", "Encontrada"])
        self.assertEqual(detalle[0]["Hojas encontradas"], "June")
        self.assertEqual(detalle[1]["Hojas encontradas"], "Julio")
        libro = load_workbook(BytesIO(salida))
        self.assertEqual(libro.sheetnames, ["Compass", "June", "Julio", "Reporte de validación"])
        for fila in (4, 5):
            for celda in libro.active[fila]:
                self.assertEqual(celda.fill.fgColor.rgb, "FFFFC7CE")
        self.assertIsNone(libro.active["A2"].fill.fill_type)
        self.assertEqual(libro.active["A4"].value, "003")
        self.assertEqual(libro["Julio"]["A3"].value, "3")

    def test_encabezados_configurables_y_ambos_meses(self):
        data = self.archivo([("Título",), ("Policy Number",), ("P1",)],
                            [("Policy",), ("P1",)], [("Policy",), ("P1",)])
        _, detalle = validar_polizas(data, (2, 1, 1))
        self.assertEqual(detalle[0]["Hojas encontradas"], "June, Julio")
        self.assertEqual(detalle[0]["Fila Excel"], 3)

    def test_columnas_invalidas(self):
        for encabezado in [("Otra",), ("Policy", "Policy")]:
            with self.subTest(encabezado=encabezado), self.assertRaisesRegex(ValueError, "June"):
                validar_polizas(self.archivo([("Policy Number",)], [encabezado], [("Policy",)]))

    def test_cantidad_hojas_invalida(self):
        libro = Workbook()
        salida = BytesIO()
        libro.save(salida)
        with self.assertRaisesRegex(ValueError, "al menos dos hojas"):
            validar_polizas(salida.getvalue())

    def test_doce_hojas_busca_todas_y_conserva_reglas(self):
        libro = Workbook()
        libro.active.title = "Compass"
        for fila in [("Policy Number", "Carrier", "Line of Business"),
                     ("SOLO_FINAL", "Otro", "Auto"), ("REPETIDA", "Otro", "Auto"),
                     ("UAO000704127", "United Auto FL", "Auto"),
                     ("2030105262 - 01", "National General FL", "Auto"),
                     ("FALTA", "Otro", "Auto")]:
            libro.active.append(fila)
        for i in range(1, 12):
            libro.create_sheet(f"Producción {i}").append(["Policy", "LOB"])
        libro["Producción 1"].append(["REPETIDA"])
        libro["Producción 10"].append(["REPETIDA"])
        libro["Producción 11"].append(["SOLO_FINAL"])
        libro["Producción 11"].append([704127, "UAO"])
        libro["Producción 9"].append(["2030105262 01"])
        data = BytesIO()
        libro.save(data)
        salida, detalle = validar_polizas(data.getvalue())
        self.assertEqual([f["Estado"] for f in detalle], ["Encontrada"] * 4 + ["No encontrada"])
        self.assertEqual(detalle[0]["Hojas encontradas"], "Producción 11")
        self.assertEqual(detalle[1]["Hojas encontradas"], "Producción 1, Producción 10")
        self.assertEqual(detalle[2]["Hojas encontradas"], "Producción 11")
        self.assertEqual(detalle[3]["Hojas encontradas"], "Producción 9")
        resultado = load_workbook(BytesIO(salida))
        self.assertEqual(len(resultado.worksheets), 13)
        self.assertEqual(resultado.worksheets[-1]["B4"].value, 12)

    def test_match_con_otro_mes(self):
        # Las dos primeras hojas de producción (June, July) son los meses
        # esperados; May, al ser la tercera, cuenta como "otro mes".
        libro = Workbook()
        libro.active.title = "Compass"
        libro.active.append(["Policy Number"])
        libro.active.append(["A1"])
        libro.active.append(["A2"])
        libro.active.append(["A3"])
        for nombre in ("June", "July", "May"):
            libro.create_sheet(nombre).append(["Policy"])
        libro["May"].append(["A1"])
        libro["June"].append(["A2"])
        libro["June"].append(["A3"])
        libro["May"].append(["A3"])
        data = BytesIO()
        libro.save(data)
        _, detalle = validar_polizas(data.getvalue())
        self.assertEqual([f["Estado"] for f in detalle], ["Encontrada"] * 3)
        # Solo aparece en la tercera hoja (May): se avisa junto con el mes.
        self.assertEqual(detalle[0]["Hojas encontradas"], "May")
        self.assertEqual(detalle[0]["Observación"], "Match con otro mes: May")
        # Aparece en un mes esperado (June): sin aviso.
        self.assertEqual(detalle[1]["Hojas encontradas"], "June")
        self.assertEqual(detalle[1]["Observación"], "")
        # Aparece en un mes esperado Y en otro: tampoco se avisa (ya está en junio).
        self.assertEqual(detalle[2]["Hojas encontradas"], "June, May")
        self.assertEqual(detalle[2]["Observación"], "")

    def test_dos_hojas_y_encabezados_independientes(self):
        libro = Workbook()
        libro.active.append(["Policy Number"])
        libro.active.append(["A"])
        produccion = libro.create_sheet("Mayo")
        produccion.append(["Título"])
        produccion.append(["Policy"])
        produccion.append(["A"])
        data = BytesIO()
        libro.save(data)
        _, detalle = validar_polizas(data.getvalue(), (1, 2))
        self.assertEqual(detalle[0]["Hojas encontradas"], "Mayo")
        with self.assertRaisesRegex(ValueError, "2 hojas"):
            validar_polizas(data.getvalue(), (1, 2, 1))

    def test_lotes_multiples_y_progreso(self):
        data = self.archivo(
            [("Policy Number",)] + [(f"P{i}",) for i in range(10005)],
            [("Policy",)] + [(f"P{i}",) for i in range(6000)],
            [("Policy",)] + [(f"P{i}",) for i in range(6000, 10004)])
        eventos = []
        salida, detalle = validar_polizas(data, progreso=lambda texto, avance: eventos.append((texto, avance)))
        self.assertEqual(len(detalle), 10005)
        self.assertEqual(sum(f["Estado"] == "Encontrada" for f in detalle), 10004)
        self.assertEqual(detalle[-1]["Estado"], "No encontrada")
        self.assertEqual(len([e for e in eventos if e[0].startswith("Hoja Compass:")]), 3)
        self.assertTrue(all(0 <= avance <= 1 for _, avance in eventos))
        self.assertEqual(eventos[-1], ("Validación completada", 1.0))
        libro = load_workbook(BytesIO(salida))
        self.assertEqual(libro.active["A10006"].fill.fgColor.rgb, "FFFFC7CE")

    def test_lote_invalido(self):
        with self.assertRaisesRegex(ValueError, "lote"):
            validar_polizas(b"", tamano_lote=0)

    def test_reporte_colores_corresponden_al_estado(self):
        data = self.archivo([("Policy Number", "Cliente"), ("FALTA1", "A"),
                             ("OK1", "B"), ("FALTA2", "C"), ("OK2", "D"), (None, "E")],
                            [("Policy",), ("OK1",)], [("Policy",), ("OK2",)])
        salida, detalle = validar_polizas(data)
        hoja = load_workbook(BytesIO(salida)).worksheets[-1]
        encabezado = next(fila[0].row for fila in hoja.iter_rows() if fila[0].value == "Fila Excel")
        self.assertEqual(hoja.auto_filter.ref, f"A{encabezado}:I{encabezado + len(detalle)}")
        for numero, resultado in enumerate(detalle, encabezado + 1):
            self.assertEqual(hoja.cell(numero, 3).value, resultado["Estado"])
            for columna in range(1, 10):
                celda = hoja.cell(numero, columna)
                if resultado["Estado"] == "Encontrada":
                    self.assertIsNone(celda.fill.fill_type)
                else:
                    self.assertEqual(celda.fill.fgColor.rgb, "FFFFC7CE")

    def test_united_lob_y_numero_en_misma_fila(self):
        data = self.archivo(
            [("Policy Number", "Carrier"), ("UAO000704127", "United   Auto FL"),
             ("UAO000800", "UNITED"), ("UAO000900", "United"),
             ("UAO000704127", "Otro carrier"), ("UAO0001002", "United"),
             ("UAO000000", "United"), ("UAO-123", "United"), ("UAO", "United")],
            [("Policy", "LOB"), (704127, "UAO"), (800, "OTRO"), (999, "UAO"),
             ("UAO000900", "UAO"), (0, "UAO"), ("UAO-123", "UAO")],
            [("Police", "LOB"), ("0001002", " uao ")])
        salida, detalle = validar_polizas(data, tamano_lote=2)
        self.assertEqual([r["Estado"] for r in detalle],
                         ["Encontrada", "No encontrada", "No encontrada", "No encontrada",
                          "Encontrada", "Encontrada", "No encontrada", "No encontrada"])
        self.assertEqual(detalle[0]["LOB buscado"], "UAO")
        self.assertEqual(detalle[0]["Policy buscado"], "704127")
        self.assertEqual(detalle[4]["Hojas encontradas"], "Julio")
        # El guion ya no invalida el formato: ambos valores se normalizan antes de
        # separar letras y número, así que "UAO-123" se interpreta como UAO + 123.
        self.assertEqual(detalle[6]["LOB buscado"], "UAO")
        self.assertEqual(detalle[6]["Policy buscado"], "123")
        self.assertEqual(detalle[6]["Observación"], "")
        # Sin ningún dígito, el formato sigue siendo inválido.
        self.assertIn("inválido", detalle[-1]["Observación"])
        libro = load_workbook(BytesIO(salida))
        self.assertEqual(libro.active["A2"].value, "UAO000704127")
        self.assertIsNone(libro.active["A2"].fill.fill_type)
        self.assertEqual(libro.active["A3"].fill.fgColor.rgb, "FFFFC7CE")

    def test_united_respaldo_sin_prefijo_de_letras(self):
        data = self.archivo(
            [("Policy Number", "Carrier"),
             ("MC984008227", "United Motorcycle"),
             ("MC000123", "United"),
             ("UAO704127", "United")],
            [("Policy", "LOB"), ("984008227", ""), (704127, "UAO")],
            [("Policy", "LOB"), ("nada", "")])
        salida, detalle = validar_polizas(data)
        self.assertEqual([r["Estado"] for r in detalle], ["Encontrada", "No encontrada", "Encontrada"])
        # "MC" no es un LOB de producción: se quita el prefijo y se busca "984008227"
        # en cualquier hoja (regla específica de United sobre la unificación general).
        self.assertEqual(detalle[0]["Regla"], "United: sin prefijo de letras")
        self.assertEqual(detalle[0]["Hojas encontradas"], "June")
        self.assertEqual(detalle[0]["LOB buscado"], "")
        self.assertEqual(detalle[0]["Policy buscado"], "984008227")
        self.assertIn("quitando el prefijo de letras", detalle[0]["Observación"])
        # Ni por LOB + Policy ni quitando el prefijo aparece "000123" en ningún lado.
        self.assertEqual(detalle[1]["Regla"], "United: LOB + Policy")
        self.assertEqual(detalle[1]["Observación"], "")
        # El respaldo no se activa cuando LOB + Policy ya encuentra la póliza.
        self.assertEqual(detalle[2]["Regla"], "United: LOB + Policy")
        self.assertEqual(detalle[2]["LOB buscado"], "UAO")
        self.assertEqual(detalle[2]["Policy buscado"], "704127")
        self.assertEqual(detalle[2]["Observación"], "")

    def test_united_necesita_lob(self):
        data = self.archivo([("Policy Number", "Carrier"), ("UAO001", "United Auto FL")],
                            [("Policy",), (1,)], [("Policy", "LOB"), (1, "UAO")])
        with self.assertRaisesRegex(ValueError, 'June.*LOB'):
            validar_polizas(data)

    def test_progressive_valor_normalizado_y_respaldo_united(self):
        data = self.archivo(
            [("Policy Number", "Carrier"),
             ("PR-000123", "Progressive Direct"),
             ("UAO000704127", "Progressive Advantage"),
             ("UAM000900", "Progressive"),
             ("XYZ999999", "Progressive Select")],
            [("Policy", "LOB"), ("PR000123", ""), (704127, "UAO")],
            [("Policy",), ("nada",)])
        salida, detalle = validar_polizas(data)
        self.assertEqual([r["Estado"] for r in detalle],
                         ["Encontrada", "Encontrada", "No encontrada", "No encontrada"])
        # Coincidencia directa por el valor normalizado (el guion de "PR-000123" no
        # impide igualarlo con "PR000123"): sin respaldo United.
        self.assertEqual(detalle[0]["Regla"], "Progressive: valor normalizado")
        self.assertEqual(detalle[0]["Hojas encontradas"], "June")
        self.assertEqual(detalle[0]["Policy buscado"], "pr000123")
        self.assertEqual(detalle[0]["Observación"], "")
        # Sin coincidencia por el valor normalizado pero prefijo UAO: respaldo United
        # encuentra la póliza.
        self.assertEqual(detalle[1]["Regla"], "Progressive → United: LOB + Policy")
        self.assertEqual(detalle[1]["Hojas encontradas"], "June")
        self.assertEqual(detalle[1]["LOB buscado"], "UAO")
        self.assertEqual(detalle[1]["Policy buscado"], "704127")
        self.assertIn("No encontrada con la búsqueda Progressive por el valor normalizado", detalle[1]["Observación"])
        self.assertIn("No se pudo comprobar United en hojas sin LOB: Julio", detalle[1]["Observación"])
        # Prefijo UAM: se intenta el respaldo United pero tampoco aparece.
        self.assertEqual(detalle[2]["Regla"], "Progressive → United: LOB + Policy")
        self.assertEqual(detalle[2]["Hojas encontradas"], "")
        self.assertEqual(detalle[2]["LOB buscado"], "UAM")
        self.assertEqual(detalle[2]["Policy buscado"], "900")
        self.assertIn("No se pudo comprobar United en hojas sin LOB: Julio", detalle[2]["Observación"])
        # Sin coincidencia por el valor normalizado y sin prefijo UAO/UAM: no se
        # intenta el respaldo.
        self.assertEqual(detalle[3]["Regla"], "Progressive: valor normalizado")
        self.assertEqual(detalle[3]["Observación"], "")
        libro = load_workbook(BytesIO(salida))
        self.assertIsNone(libro.active["A2"].fill.fill_type)
        self.assertIsNone(libro.active["A3"].fill.fill_type)
        self.assertEqual(libro.active["A4"].fill.fgColor.rgb, "FFFFC7CE")
        self.assertEqual(libro.active["A5"].fill.fgColor.rgb, "FFFFC7CE")

    def test_progressive_respaldo_sin_prefijo_de_letras(self):
        # Caso reportado: prefijos que no son un LOB de United (MC, MT) y en
        # producción solo aparece el número, sin ningún prefijo.
        data = self.archivo(
            [("Policy Number", "Carrier"),
             ("MC984008227", "Progressive Motorcycle"),
             ("MT939904610", "Progressive")],
            [("Policy",), ("984008227",)],
            [("Policy",), ("939904610",)])
        salida, detalle = validar_polizas(data)
        self.assertEqual([r["Estado"] for r in detalle], ["Encontrada", "Encontrada"])
        self.assertEqual(detalle[0]["Regla"], "Progressive: sin prefijo de letras")
        self.assertEqual(detalle[0]["Hojas encontradas"], "June")
        self.assertEqual(detalle[0]["Policy buscado"], "984008227")
        self.assertIn("quitando el prefijo de letras", detalle[0]["Observación"])
        self.assertEqual(detalle[1]["Regla"], "Progressive: sin prefijo de letras")
        self.assertEqual(detalle[1]["Hojas encontradas"], "Julio")
        self.assertEqual(detalle[1]["Policy buscado"], "939904610")

    def test_national_general_y_excepcion_flood(self):
        for encabezado in ("Line of Bussinees", "Line of Business"):
            with self.subTest(encabezado=encabezado):
                data = self.archivo(
                    [("Policy Number", "Carrier", encabezado),
                     ("2036745969-00", "National General Auto", "Auto"),
                     ("2030105262 - 01", "Carrier NATIONAL GENERAL FL", "Home"),
                     ("2036745969-00", "National General", " Flood "),
                     ("2036745969-00", "Otro", "Auto"),
                     ("000123-01", "National General", "Auto"),
                     ("999-02", "National General", "Flood")],
                    [("Policy",), ("2036745969 00",), ("000123 01",), ("999-02",)],
                    [("Policy",), ("2030105262 01",)])
                salida, detalle = validar_polizas(data)
                # Con la normalización unificada, Flood ya no es más estricto que el
                # resto (fila 3) y un carrier "Otro" también se beneficia (fila 4):
                # ambos comparan igual que cualquier otra regla.
                self.assertEqual([r["Estado"] for r in detalle], ["Encontrada"] * 6)
                self.assertEqual(detalle[0]["Policy buscado"], "203674596900")
                self.assertEqual(detalle[1]["Policy buscado"], "203010526201")
                self.assertEqual(detalle[4]["Policy buscado"], "00012301")
                self.assertEqual(detalle[2]["Regla"], "Policy Number")
                libro = load_workbook(BytesIO(salida))
                self.assertEqual(libro.active["A2"].value, "2036745969-00")

    def test_national_general_requiere_linea_para_excluir_flood(self):
        data = self.archivo([("Policy Number", "Carrier"), ("123-00", "National General")],
                            [("Policy",), ("123 00",)], [("Policy",)])
        with self.assertRaisesRegex(ValueError, "Flood"):
            validar_polizas(data)

    def test_national_unifica_separadores_en_ambas_direcciones(self):
        variantes = ["203314922801", "2033149228-01", "2033149228 - 01", "2033149228 01",
                     203314922801, "2033149228\u00a001", "2033149228–01"]
        for origen in variantes:
            for destino in variantes:
                with self.subTest(origen=origen, destino=destino):
                    data = self.archivo([("Policy Number", "Carrier", "Line of Business"),
                                         (origen, "National Florida Auto", "Auto")],
                                        [("Policy",), (destino,)], [("Policy",)])
                    _, detalle = validar_polizas(data)
                    self.assertEqual(detalle[0]["Estado"], "Encontrada")
                    self.assertEqual(detalle[0]["Policy buscado"], "203314922801")

    def test_national_no_confunde_sufijos_ceros_ni_otros_carriers(self):
        data = self.archivo(
            [("Policy Number", "Carrier", "Line of Business"),
             ("203314922802", "National General", "Auto"),
             ("00203314922801", "National General", "Auto"),
             ("203314922801", "National Florida", "Flood"),
             ("203314922801", "Otro carrier", "Auto"),
             ("12-01", "National General", "Auto"),
             (" - ", "National General", "Auto")],
            [("Policy",), ("2033149228 01",), ("1201",), ("-",)], [("Policy",)])
        _, detalle = validar_polizas(data)
        # Los sufijos/ceros distintos (filas 1-2) siguen sin confundirse. Las filas
        # 3 y 4 ahora sí aparecen: con la normalización unificada, Flood y un
        # carrier "Otro" comparan igual que el resto. La fila 6 (" - ") normaliza a
        # vacío en ambos lados y no debe "empatar" con la celda vacía de producción.
        self.assertEqual([f["Estado"] for f in detalle],
                         ["No encontrada", "No encontrada", "Encontrada", "Encontrada",
                          "Encontrada", "No encontrada"])


if __name__ == "__main__":
    unittest.main()
