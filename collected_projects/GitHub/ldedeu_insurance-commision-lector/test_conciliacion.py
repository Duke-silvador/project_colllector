import unittest
from copy import copy
from datetime import date
from decimal import Decimal
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from conciliacion import (buscar_transacciones, buscar_transacciones_por_monto, buscar_transacciones_por_montos,
                          generar_excel, generar_excel_solo_reporte, grand_total_pdf,
                          LDA_COLUMNS_DEFAULT, LDA_COLUMNS_AMWINS, CARRIERS_FORMATO_AMWINS,
                          _mes_produccion_mm_aaaa)
from openpyxl.utils import get_column_letter
from franquicias import _normalizar


def bank_bytes():
    book = Workbook()
    ws = book.active
    ws.title = 'Statement'
    ws.append(['Date', 'Description', 'Amount', 'Class'])
    ws.append(['8/17/2026', 'COMMONWEALTH CAS DES:ACH', 64.85, 'Commissions'])
    ws.append(['8/18/2026', 'OTHER', 64.85, 'Commissions'])
    ws.append(['8/19/2026', 'commonwealth ACH', 10, 'Commissions'])
    ws.append(['8/20/2026', 'COMMONWEALTH duplicate', 64.85, 'Commissions'])
    ws['C2'].number_format = '"USD "#,##0.0000'
    ws['C2'].font = Font(name='Arial', bold=True, color='123456')
    ws['C2'].fill = PatternFill('solid', fgColor='FFEEDD')
    output = BytesIO()
    book.save(output)
    return output.getvalue()


class ReconciliationTests(unittest.TestCase):
    def test_standalone_report_preserves_cached_values_and_format_for_all_carriers(self):
        from carriers import CARRIERS
        from zipfile import ZipFile
        for carrier in CARRIERS:
            with self.subTest(carrier=carrier):
                bank = Workbook()
                bank.active.title = 'Statement'
                bank.active.append(['Description', 'Amount'])
                bank.active.append([carrier, 64.85])
                raw = BytesIO()
                bank.save(raw)
                bank.close()
                rows = [dict(policy_number='=P001', producer_code='0012', premium_amount='1000',
                             commission_amount='70', franchise_number='DTF0082', franchise_percent='.070123'),
                        dict(policy_number='', producer_code='', transaction_type='Report fee',
                             commission_amount='-5.15', comm_percent='1', _chargeback=True)]
                original = generar_excel(rows, Decimal('64.85'), raw.getvalue(), 'Statement', 1, 2, carrier, {})
                exported = generar_excel_solo_reporte(original)
                source = load_workbook(BytesIO(original), data_only=True)
                result = load_workbook(BytesIO(exported))
                schema = LDA_COLUMNS_AMWINS if carrier in CARRIERS_FORMATO_AMWINS else LDA_COLUMNS_DEFAULT
                def col(field):
                    return get_column_letter(next(i for i, (_, f) in enumerate(schema, 1) if f == field))
                try:
                    self.assertEqual(source.sheetnames, ['Data', 'Pivot', 'Bank', 'Report_LDA'])
                    self.assertEqual(result.sheetnames, ['Report_LDA'])
                    ws = result.active
                    self.assertEqual(list(ws.values), list(source['Report_LDA'].values))
                    policy_cell = ws[col('policy_number') + '2']
                    self.assertEqual(policy_cell.value, '=P001')
                    self.assertEqual(policy_cell.data_type, 's')
                    self.assertEqual(ws[col('producer_code') + '2'].value, '0012')
                    self.assertEqual(ws[col('del_toro_commission') + '2'].value, 70)
                    fran_col, dif_col = col('franchise_commission'), col('commission_difference')
                    self.assertEqual(ws[fran_col + '2'].value, 70.123)
                    self.assertAlmostEqual(ws[dif_col + '2'].value, -.123)
                    self.assertIsNone(ws[fran_col + '3'].value)
                    self.assertEqual(ws[col('del_toro_commission') + '3'].value, -5.15)
                    self.assertEqual(ws.freeze_panes, 'A2')
                    self.assertEqual(ws.auto_filter.ref, source['Report_LDA'].auto_filter.ref)
                    self.assertEqual(ws[fran_col + '2'].number_format, source['Report_LDA'][fran_col + '2'].number_format)
                    self.assertEqual(copy(ws['A1'].fill), copy(source['Report_LDA']['A1'].fill))
                    self.assertFalse(any(cell.data_type == 'f' for row in ws for cell in row))
                    with ZipFile(BytesIO(exported)) as archive:
                        self.assertFalse(any(name.startswith('xl/externalLinks/') for name in archive.namelist()))
                finally:
                    source.close()
                    result.close()

    def test_office_alias_replaces_franchise_in_both_output_sheets(self):
        rows = [dict(policy_number='P1', producer_code='12IA', premium_amount='100',
                     commission_amount='64.85', franchise_number='DTF0082', office_id='uid')]
        aliases = [dict(office_id='uid', office_number='82', franchise_alias='99')]
        result = generar_excel(rows, Decimal('64.85'), bank_bytes(), 'Statement', 1, 2, 'ORCHID', {}, alias_franquicias=aliases)
        book = load_workbook(BytesIO(result))
        lda_col = next(c.column for c in book['Report_LDA'][1] if c.value in ('franchise', 'Franchise'))
        self.assertEqual(book['Data']['A2'].value, 'DTF0099')
        self.assertEqual(book['Report_LDA'].cell(2, lda_col).value, 'DTF0099')
        self.assertEqual(rows[0]['franchise_number'], 'DTF0082')
        book.close()
        aliases[0]['franchise_alias'] = '144-0024'
        result = generar_excel(rows, Decimal('64.85'), bank_bytes(), 'Statement', 1, 2, 'ORCHID', {}, alias_franquicias=aliases)
        book = load_workbook(BytesIO(result))
        self.assertEqual(book['Data']['A2'].value, 'DTF144-0024')
        self.assertEqual(book['Report_LDA'].cell(2, lda_col).value, 'DTF144-0024')
        book.close()

    def test_office_alias_also_applies_to_chargeback_rows(self):
        """Regresion: el alias se aplicaba solo a filas con poliza; una fila de chargeback
        (ej. MVR de THE GENERAL, o un cargo de COMMONWEALTH) resuelta a una oficina con alias
        configurado debia mostrar el alias igual, no el numero de franquicia crudo."""
        rows = [dict(policy_number='', producer_code='91483', insured_name='Amanda Melendez',
                     transaction_type='MVR', commission_amount='-6.44', comm_percent='1',
                     franchise_number='DTF0044', office_id='uid-mvr', franchise_percent='1',
                     franchise_commission='-6.44', commission_difference='0')]
        aliases = [dict(office_id='uid-mvr', office_number='44', franchise_alias='7')]
        result = generar_excel(rows, Decimal('-6.44'), None, None, None, None, 'THE GENERAL', {},
                               alias_franquicias=aliases)
        book = load_workbook(BytesIO(result))
        try:
            lda_col = next(c.column for c in book['Report_LDA'][1] if c.value in ('franchise', 'Franchise'))
            self.assertEqual(book['Data']['A2'].value, 'DTF0007')
            self.assertEqual(book['Report_LDA'].cell(2, lda_col).value, 'DTF0007')
        finally:
            book.close()

    def test_florida_bank_sums_three_carriers_and_exports_all_rows(self):
        book = Workbook()
        ws = book.active
        ws.title = 'Statement'
        ws.append(['Description', 'Amount'])
        for name, amount in [('ACH FLORIDA PENINSULA PAYMENT', 10), ('ovation home ACH', 20), ('EDISON deposit', 30), ('OTHER', 60)]:
            ws.append([name, amount])
        raw = BytesIO()
        book.save(raw)
        terminos = ('florida peninsul', 'ovation home', 'edison')
        candidates = buscar_transacciones(raw.getvalue(), 'Statement', 1, 'FLORIDA PENINSULA', terminos)
        self.assertEqual([r['row'] for r in candidates], [2, 3, 4])
        rows = [dict(producer_code='45595', commission_amount='60', comm_percent='.12')]
        result = generar_excel(rows, Decimal('60'), raw.getvalue(), 'Statement', 1, [2, 3, 4], 'FLORIDA PENINSULA', {},
                                terminos_banco=terminos)
        output = load_workbook(BytesIO(result))
        cached = load_workbook(BytesIO(result), data_only=True)
        self.assertEqual(output['Bank'].max_row, 4)
        self.assertIn('=SUM(Bank!B2:B4)', [cell.value for row in output['Pivot'] for cell in row])
        self.assertEqual(cached['Pivot'].cell(cached['Pivot'].max_row, 2).value, 0)
        output.close()
        cached.close()
        # Si el banco no suma exacto (aqui solo 2 de las 3 filas), ya no bloquea: se genera
        # igual y la hoja Pivot registra la diferencia real.
        parcial = generar_excel(rows, Decimal('60'), raw.getvalue(), 'Statement', 1, [2, 3], 'FLORIDA PENINSULA', {},
                                terminos_banco=terminos)
        cached_parcial = load_workbook(BytesIO(parcial), data_only=True)
        self.assertEqual(cached_parcial['Pivot'].cell(cached_parcial['Pivot'].max_row, 2).value, 30)
        cached_parcial.close()
        book.close()

    def test_formulas_tables_and_cached_results_support_manual_rows(self):
        from openpyxl.formula.translate import Translator
        rows = [dict(policy_number='P1', producer_code='12IA', premium_amount='1000',
                     commission_amount='70', franchise_number='DTF0082', franchise_percent='.070123'),
                dict(policy_number='', producer_code='', transaction_type='Report fee',
                     commission_amount='-5.15', comm_percent='1', _chargeback=True)]
        blob = generar_excel(rows, Decimal('64.85'), bank_bytes(), 'Statement', 1, 2, 'ORCHID', {})
        book = load_workbook(BytesIO(blob))
        cached = load_workbook(BytesIO(blob), data_only=True)

        def col(field):
            return get_column_letter(next(i for i, (_, f) in enumerate(LDA_COLUMNS_AMWINS, 1) if f == field))
        pol, prem = col('policy_number'), col('premium_amount')
        toro_pct, toro_com = col('del_toro_percent'), col('del_toro_commission')
        fran_pct, fran_com, dif = col('franchise_percent'), col('franchise_commission'), col('commission_difference')
        try:
            ws = book['Data']
            self.assertEqual(ws['N2'].value, '=IF(C2="","",IF(F2=0,IF(G2=0,0,NA()),G2/F2))')
            self.assertEqual(ws['O2'].value, '=IF(C2="","",G2)')
            self.assertEqual(book['Report_LDA'][fran_com + '2'].value,
                             f'=IF(OR({pol}2="",{fran_pct}2=""),"",{prem}2*{fran_pct}2)')
            self.assertEqual(ws['G2'].data_type, 'n')
            self.assertEqual(book['Report_LDA'][fran_pct + '2'].value, .070123)
            self.assertEqual(ws.auto_filter.ref, 'A1:P3')
            self.assertFalse(ws.tables)
            self.assertEqual(Translator(book['Report_LDA'][fran_com + '2'].value, origin=fran_com + '2').translate_formula(fran_com + '4'),
                             f'=IF(OR({pol}4="",{fran_pct}4=""),"",{prem}4*{fran_pct}4)')
            self.assertEqual(book['Report_LDA'][toro_com + '2'].value, '=Data!G2')
            self.assertEqual(book['Report_LDA'][toro_pct + '2'].data_type, 'f')
            self.assertFalse(book['Report_LDA'].tables)
            from zipfile import ZipFile
            with ZipFile(BytesIO(blob)) as archive:
                self.assertFalse(any(name.startswith('xl/tables/') for name in archive.namelist()))
            self.assertEqual(book['Pivot']['B6'].value, '=SUM(Data!G:G)')
            self.assertIn('SUMIF(Data!H:H,', book['Pivot']['B4'].value)
            self.assertEqual(book['Pivot']['B7'].value, '=Bank!C2')
            self.assertEqual(book['Pivot']['B8'].value, '=B6-B7')
            self.assertEqual(cached['Data']['N2'].value, .07)
            self.assertEqual(cached['Report_LDA'][fran_com + '2'].value, 70.123)
            self.assertAlmostEqual(cached['Report_LDA'][dif + '2'].value, -.123)
            self.assertNotIn('Franchise %', [cell.value for cell in ws[1]])
            self.assertNotIn('Franchise $', [cell.value for cell in ws[1]])
            self.assertIsNone(cached['Report_LDA'][fran_com + '3'].value)
            self.assertEqual(cached['Report_LDA'][toro_pct + '3'].value, 1)
            self.assertEqual(cached['Pivot']['B8'].value, 0)
            self.assertTrue(book.calculation.fullCalcOnLoad)
        finally:
            book.close()
            cached.close()

    def test_matching_and_duplicate_candidates(self):
        rows = buscar_transacciones(bank_bytes(), 'Statement', 1, 'COMMONWEALTH')
        self.assertEqual([row['row'] for row in rows], [2, 4, 5])
        self.assertEqual(sum(row['amount'] == Decimal('64.85') for row in rows), 2)

    def test_florida_truncated_bank_description(self):
        book = Workbook()
        ws = book.active
        ws.append(['Description', 'Amount'])
        ws.append(['ACH FLORIDA PENINSUL', 10])
        ws.append(['Florida Peninsula deposit', 20])
        ws.append(['OTHER', 30])
        raw = BytesIO()
        book.save(raw)
        matches = buscar_transacciones(raw.getvalue(), ws.title, 1, 'FLORIDA PENINSULA',
                                       ('florida peninsul', 'ovation home', 'edison'))
        self.assertEqual([r['row'] for r in matches], [2, 3])
        book.close()

    def test_configured_terms_add_to_carrier_name_instead_of_replacing_it(self):
        """Regresion: los terminos configurados se suman a la busqueda por nombre del carrier,
        no la reemplazan. THE GENERAL deposita las franquicias con 'The General' en Description
        pero el codigo master con otro texto (ej. 'DTF PAYROLL'); con un termino configurado para
        ese texto, ambas variantes deben encontrarse juntas."""
        book = Workbook()
        ws = book.active
        ws.append(['Description', 'Amount'])
        ws.append(['AAMGA OPERATING ACH', 10])
        ws.append(['ASSURANCE deposit', 20])
        raw = BytesIO()
        book.save(raw)
        con_termino = buscar_transacciones(raw.getvalue(), ws.title, 1, 'ASSURANCE', ['AAMGA OPERATING'])
        self.assertEqual(sorted(r['row'] for r in con_termino), [2, 3])
        sin_termino = buscar_transacciones(raw.getvalue(), ws.title, 1, 'ASSURANCE')
        self.assertEqual([r['row'] for r in sin_termino], [3])
        book.close()

    def test_buscar_por_monto_matches_any_description_by_exact_amount(self):
        """Ultimo recurso cuando el carrier no aparece en Description ni por terminos
        configurados: se busca en todo el banco por el importe exacto del reporte."""
        book = Workbook()
        ws = book.active
        ws.append(['Description', 'Amount'])
        ws.append(['WIRE TRANSFER XYZ', 64.85])
        ws.append(['OTHER PAYMENT', 10.00])
        ws.append(['EMPTY ROW', None])
        raw = BytesIO()
        book.save(raw)
        rows = buscar_transacciones_por_monto(raw.getvalue(), ws.title, 1, Decimal('64.85'))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['description'], 'WIRE TRANSFER XYZ')
        book.close()

    def test_buscar_por_monto_returns_every_match_when_amount_repeats(self):
        book = Workbook()
        ws = book.active
        ws.append(['Description', 'Amount'])
        ws.append(['WIRE TRANSFER A', 64.85])
        ws.append(['WIRE TRANSFER B', 64.85])
        ws.append(['UNRELATED', 10.00])
        raw = BytesIO()
        book.save(raw)
        rows = buscar_transacciones_por_monto(raw.getvalue(), ws.title, 1, Decimal('64.85'))
        self.assertEqual({r['description'] for r in rows}, {'WIRE TRANSFER A', 'WIRE TRANSFER B'})
        book.close()

    def test_buscar_por_montos_matches_each_piece_separately(self):
        """Cuando el total del reporte es la suma de varios cheques que se depositan por
        separado (BASS, CRC GROUP), ninguna fila individual del banco coincide con el total
        combinado; hay que buscar cada pedazo por su cuenta."""
        book = Workbook()
        ws = book.active
        ws.append(['Description', 'Amount'])
        ws.append(['CHECK DEPOSIT 1', 243.90])
        ws.append(['CHECK DEPOSIT 2', 97.60])
        ws.append(['UNRELATED', 500.00])
        raw = BytesIO()
        book.save(raw)
        rows = buscar_transacciones_por_montos(raw.getvalue(), ws.title, 1, [Decimal('243.90'), Decimal('97.60')])
        self.assertEqual({r['description'] for r in rows}, {'CHECK DEPOSIT 1', 'CHECK DEPOSIT 2'})
        book.close()

    def test_workbook_and_signed_total(self):
        records = [dict(producer_code='12IA', insured_name='=literal', commission_amount=Decimal('74.95'),
                         comm_percent=Decimal('.12'), producer_name='Del Toro Insurance'),
                   dict(producer_code='12IA', commission_amount=Decimal('-10.10'), comm_percent=Decimal('1'),
                        producer_name='')]
        mapa = {(_normalizar('COMMONWEALTH'), _normalizar('Del Toro Insurance')): 'DTF0001'}
        data = generar_excel(records, Decimal('64.85'), bank_bytes(), 'Statement', 1, 2, 'COMMONWEALTH', mapa)
        book = load_workbook(BytesIO(data), data_only=True)
        self.assertEqual(book.sheetnames, ['Data', 'Pivot', 'Bank', 'Report_LDA'])
        self.assertEqual(book['Pivot']['A4'].value, '12IA')
        self.assertEqual(book['Pivot']['B4'].value, 64.85)
        self.assertEqual(book['Pivot']['B7'].value, 0)
        self.assertEqual(book['Bank']['B2'].value, 'COMMONWEALTH CAS DES:ACH')
        self.assertEqual(book['Bank']['C2'].number_format, '"USD "#,##0.0000')
        self.assertEqual(book['Bank']['C2'].font.name, 'Arial')
        self.assertTrue(book['Bank']['C2'].font.bold)
        self.assertEqual(book['Bank']['C2'].fill.fgColor.rgb, '00FFEEDD')
        self.assertEqual(book['Data']['A1'].value, 'Franchise')
        self.assertIsNone(book['Data']['A2'].value)
        self.assertEqual(book['Data']['A2'].fill.fgColor.rgb, '00FFC7CE')
        self.assertIsNone(book['Data']['A3'].value)
        self.assertEqual(book['Data']['B1'].value, 'Name Insured')
        self.assertEqual(book['Data']['B2'].value, '=literal')
        book.close()

    def test_missing_commission_or_franchise_marks_red_including_chargebacks(self):
        records = [dict(producer_code='12IA', producer_name='Company', policy_number='P1', premium_amount='250', commission_amount='30'),
                   dict(producer_code='OTHER', producer_name='Company', policy_number='P2', premium_amount='400', commission_amount='40'),
                   dict(producer_code='', policy_number='', transaction_type='Report fee', _chargeback=True, commission_amount='-5.15')]
        data = generar_excel(records, Decimal('64.85'), bank_bytes(), 'Statement', 1, 2,
                             'COMMONWEALTH', {}, lambda number: None, {},
                             {('COMMONWEALTH', 'COMPANY', '12IA')})
        book = load_workbook(BytesIO(data), data_only=True)
        sheet = book['Data']
        self.assertEqual(sheet['A2'].value, 'DT120')
        self.assertEqual(sheet['A2'].fill.fgColor.rgb, '00FFC7CE')
        self.assertIsNone(sheet['A3'].value)
        self.assertTrue(all(cell.fill.fgColor.rgb == '00FFC7CE' for cell in sheet[3]))
        self.assertIsNone(sheet['A4'].value)
        self.assertEqual(sheet['A4'].fill.fgColor.rgb, '00FFC7CE')
        book.close()

    def test_lda_headers_lob_and_chargeback_values(self):
        """Formato default (LDA_COLUMNS_DEFAULT): usado como respaldo para cualquier carrier
        que todavia no este en CARRIERS_FORMATO_AMWINS. 'GENERIC CARRIER' no es un carrier real
        del proyecto -a proposito-, para no acoplar esta prueba a uno que en algun momento
        pueda migrarse al formato Amwins."""
        rows = [dict(policy_number='P1', insured_name='Test', producer_code='12IA',
                     commission_amount=Decimal('70'), premium_amount=Decimal('1000'),
                     franchise_number='DTF0082', state='TX', transaction_type='NEW_BUSINESS',
                     line_business_id='renters', term_length=12, agent_name='Must stay blank', del_toro_percent=Decimal('.12'), del_toro_commission=Decimal('120')),
                dict(policy_number='', producer_code='', transaction_type='Report fee',
                     commission_amount=Decimal('-5.15'), comm_percent=Decimal('1'), _chargeback=True)]
        book = load_workbook(BytesIO(generar_excel(rows, Decimal('64.85'), bank_bytes(),
                                 'Statement', 1, 2, 'GENERIC CARRIER', {})), data_only=True)
        ws = book['Report_LDA']
        self.assertEqual([c.value for c in ws[1]], ['franchise', 'state_code', 'producer_code',
            'producer_name', 'policy', 'insured_name', 'transaction_date', 'transaction_type',
            'premium', 'del_toro_rate', 'del_toro_commission', 'Franchise %', 'Franchise $',
            'Difference', 'agent_name', 'lob', 'term_length'])
        self.assertEqual(ws['P2'].value, 'renters')
        self.assertEqual(ws['J2'].value, .07)
        self.assertEqual(ws['J2'].number_format, '0.###')
        self.assertIsNone(ws['O2'].value)
        self.assertEqual(ws['Q2'].value, 12)
        self.assertIsNone(ws['N2'].value)
        self.assertEqual(ws['H3'].value, 'Report fee')
        self.assertEqual(ws['J3'].value, 1)
        self.assertEqual(ws['K3'].value, -5.15)
        self.assertIsNone(ws['F3'].value)
        book.close()

    def test_mes_produccion_is_one_month_before_accounting_month(self):
        """'Month' en el formato Amwins es el mes de PRODUCCION real, no el accounting_month
        guardado: el statement de un mes de produccion se contabiliza al mes siguiente, asi que
        Month siempre debe ser un mes antes (incluyendo el caso de cruzar de enero a diciembre
        del año anterior)."""
        self.assertEqual(_mes_produccion_mm_aaaa(date(2026, 9, 1)), '08-2026')
        self.assertEqual(_mes_produccion_mm_aaaa('2026-09'), '08-2026')
        self.assertEqual(_mes_produccion_mm_aaaa('2026-09-01'), '08-2026')
        self.assertEqual(_mes_produccion_mm_aaaa(date(2026, 1, 1)), '12-2025')
        self.assertEqual(_mes_produccion_mm_aaaa('2026-01'), '12-2025')

    def test_lda_headers_lob_and_chargeback_values_amwins_format(self):
        """COMMONWEALTH (piloto) usa el formato de columnas pedido por Amwins: mismos valores
        que el formato default, pero con otros nombres/orden de columna y sin State."""
        rows = [dict(policy_number='P1', insured_name='Test', producer_code='12IA',
                     commission_amount=Decimal('70'), premium_amount=Decimal('1000'),
                     franchise_number='DTF0082', state='TX', transaction_type='NEW_BUSINESS',
                     line_business_id='renters', term_length=12, agent_name='Must stay blank', del_toro_percent=Decimal('.12'), del_toro_commission=Decimal('120')),
                dict(policy_number='', producer_code='', transaction_type='Report fee',
                     commission_amount=Decimal('-5.15'), comm_percent=Decimal('1'), _chargeback=True)]
        book = load_workbook(BytesIO(generar_excel(rows, Decimal('64.85'), bank_bytes(),
                                 'Statement', 1, 2, 'COMMONWEALTH', {})), data_only=True)
        ws = book['Report_LDA']
        self.assertEqual([c.value for c in ws[1]], ['Date', 'Insured Name', 'Agency Code',
            'Transaction', 'Franchise', 'Premium', 'Del Toro %', 'Del Toro $', 'Policy',
            'Company', 'Franchise %', 'Franchise $', 'Difference', 'Month', 'LOB',
            'Agent Name', 'Term Length'])

        def col(field):
            return get_column_letter(next(i for i, (_, f) in enumerate(LDA_COLUMNS_AMWINS, 1) if f == field))

        self.assertEqual(ws[col('line_business_id') + '2'].value, 'renters')
        self.assertEqual(ws[col('del_toro_percent') + '2'].value, .07)
        self.assertEqual(ws[col('del_toro_percent') + '2'].number_format, '0.###')
        self.assertIsNone(ws[col('agent_name') + '2'].value)
        self.assertEqual(ws[col('term_length') + '2'].value, 12)
        self.assertIsNone(ws[col('commission_difference') + '2'].value)
        self.assertEqual(ws[col('carrier') + '2'].value, 'COMMONWEALTH')
        self.assertEqual(ws[col('transaction_type') + '3'].value, 'Report fee')
        self.assertEqual(ws[col('del_toro_percent') + '3'].value, 1)
        self.assertEqual(ws[col('del_toro_commission') + '3'].value, -5.15)
        self.assertIsNone(ws[col('insured_name') + '3'].value)
        book.close()

    def test_hundred_percent_shows_as_plain_1_not_1_dot(self):
        """Regresion: '0.###' deja un punto colgando ('1.') cuando Del Toro %/Franchise % es
        exactamente 1 (100%, como en los MVR de THE GENERAL, que no se reparten). Debe mostrarse
        como '1' entero, sin punto, tanto en Data como en Report_LDA."""
        rows = [dict(policy_number='', producer_code='91483', insured_name='Amanda Melendez',
                     transaction_type='MVR', commission_amount=Decimal('-6.44'), comm_percent=Decimal('1'),
                     del_toro_percent=Decimal('1'), franchise_number='DTF0044', franchise_percent=Decimal('1'),
                     franchise_commission=Decimal('-6.44'), commission_difference=Decimal('0'))]
        result = generar_excel(rows, Decimal('-6.44'), None, None, None, None, 'THE GENERAL', {})
        book = load_workbook(BytesIO(result))
        try:
            lda = book['Report_LDA']
            toro_col = next(c.column for c in lda[1] if c.value in ('del_toro_rate', 'Del Toro %'))
            fran_col = next(c.column for c in lda[1] if c.value == 'Franchise %')
            self.assertEqual(book['Data']['N2'].number_format, '0')
            self.assertEqual(lda.cell(2, toro_col).number_format, '0')
            self.assertEqual(lda.cell(2, fran_col).number_format, '0')
        finally:
            book.close()

    def test_fractional_percent_keeps_the_normal_format(self):
        """El fix del '1.' no debe afectar porcentajes fraccionarios normales (ej. 0.1, 0.08)."""
        rows = [dict(policy_number='P1', producer_code='12IA', premium_amount='1000',
                     commission_amount='70', franchise_number='DTF0082', franchise_percent='.08',
                     del_toro_percent='.10', del_toro_percent_source='statement')]
        result = generar_excel(rows, Decimal('70'), None, None, None, None, 'COMMONWEALTH', {})
        book = load_workbook(BytesIO(result))
        try:
            self.assertEqual(book['Data']['N2'].number_format, '0.###')
        finally:
            book.close()

    def test_chargeback_with_its_own_franchise_split_shows_in_lda_instead_of_blank(self):
        """Regresion: a diferencia de COMMONWEALTH (Franchise $/Difference siempre en blanco
        para chargebacks), THE GENERAL calcula su propio reparto para las filas MVR (100% a
        cargo de la franquicia) ANTES de generar_excel(); esas columnas no deben quedar en
        blanco solo porque la fila no tiene poliza."""
        rows = [dict(policy_number='FL1', insured_name='X', producer_code='91479',
                     commission_amount=Decimal('85.46'), premium_amount=Decimal('1000'),
                     franchise_number='DTF0015', state='FL', del_toro_percent=Decimal('.10'),
                     del_toro_percent_source='statement', franchise_percent=Decimal('.08')),
                dict(policy_number='', producer_code='91483', insured_name='Amanda Melendez',
                     transaction_type='MVR', commission_amount=Decimal('-6.44'), comm_percent=Decimal('1'),
                     del_toro_percent=Decimal('1'), del_toro_percent_source='statement',
                     del_toro_commission=Decimal('-6.44'), franchise_number='DTF0044', franchise_percent=Decimal('1'),
                     franchise_commission=Decimal('-6.44'), commission_difference=Decimal('0'),
                     # THE GENERAL copia el nombre del agente a producer_name para sus MVR (ver
                     # the_general.py); en filas reales ambos campos llegan con el mismo valor.
                     producer_name='LUCKYCOBRA.LLC', agent_name='LUCKYCOBRA.LLC')]
        result = generar_excel(rows, Decimal('79.02'), None, None, None, None, 'THE GENERAL', {})
        book = load_workbook(BytesIO(result), data_only=True)
        try:
            ws = book['Report_LDA']
            cols = {c.value: c.column for c in ws[1]}
            self.assertEqual(ws.cell(3, cols['Agent Name']).value, 'LUCKYCOBRA.LLC')
            self.assertEqual(ws.cell(3, cols['Franchise $']).value, -6.44)
            self.assertEqual(ws.cell(3, cols['Difference']).value, 0)
            # Regresion: el nombre del asegurado (viene del statement MVR original) quedaba
            # siempre en blanco para chargebacks en Report_LDA, aunque la fila ya lo trajera.
            self.assertEqual(ws.cell(3, cols['Insured Name']).value, 'Amanda Melendez')
            # Regresion: 'Del Toro $' en Data tambien debe mostrar el negativo, no quedar en
            # blanco solo porque la fila no tiene poliza (completar_del_toro_statement se salta
            # para chargebacks; THE GENERAL debe traer su propio del_toro_commission ya calculado).
            self.assertEqual(book['Data']['O3'].value, -6.44)
        finally:
            book.close()

    def test_internal_total_mismatch_still_blocks_output(self):
        """El total interno (suma de commission_amount) debe coincidir con el Grand Total del
        reporte; eso es un problema de datos, no del banco, y sigue bloqueando la salida."""
        with self.assertRaises(ValueError):
            generar_excel([dict(producer_code='12IA', commission_amount='64.84')], Decimal('64.85'), bank_bytes(), 'Statement', 1, 2, 'COMMONWEALTH', {})

    def test_bank_rows_can_come_from_different_sheets(self):
        """Regresion: a veces las transacciones aparecen en hojas distintas del mismo banco.
        bank_row acepta dicts {'sheet':, 'row':} ademas de numeros de fila (que asumen la
        hoja `sheet`, uso previo); el tab Bank y el importe deben combinar ambas hojas."""
        book = Workbook()
        ws1 = book.active
        ws1.title = 'Sheet1'
        ws1.append(['Description', 'Amount'])
        ws1.append(['CRC DEPOSIT 1', 243.90])
        ws2 = book.create_sheet('Sheet2')
        ws2.append(['Description', 'Amount'])
        ws2.append(['CRC DEPOSIT 2', 97.60])
        raw = BytesIO()
        book.save(raw)
        rows = [dict(producer_code='45595', commission_amount='341.50', comm_percent='.12')]
        bank_row = [{'sheet': 'Sheet1', 'row': 2}, {'sheet': 'Sheet2', 'row': 2}]
        result = generar_excel(rows, Decimal('341.50'), raw.getvalue(), None, 1, bank_row, 'CRC GROUP', {})
        output = load_workbook(BytesIO(result))
        cached = load_workbook(BytesIO(result), data_only=True)
        self.assertEqual(output['Bank'].max_row, 3)
        self.assertEqual(output['Bank'].cell(2, 1).value, 'CRC DEPOSIT 1')
        self.assertEqual(output['Bank'].cell(3, 1).value, 'CRC DEPOSIT 2')
        self.assertEqual(cached['Pivot'].cell(cached['Pivot'].max_row - 1, 2).value, 341.5)
        self.assertEqual(cached['Pivot'].cell(cached['Pivot'].max_row, 2).value, 0)
        output.close()
        cached.close()
        book.close()

    def test_pass_through_franchise_amount_matches_del_toro_exactly_no_rounding_penny(self):
        """Regresion: en carriers pass-through (ORCHID/SLIDE/GRANADA: franchise_percent ==
        del_toro_percent), Franchise $ recalculaba premium*porcentaje en vez de usar el mismo
        importe exacto del statement, y podia diferir un centavo por el redondeo del
        porcentaje mostrado (8.00% en vez de la tasa exacta 141.32/1766.40)."""
        rows = [dict(policy_number='H3FL000489030', producer_code='9991006', premium_amount='1766.40',
                     commission_amount='141.32', del_toro_percent='0.08', del_toro_percent_source='statement',
                     franchise_number='DTF0006', franchise_percent='0.08')]
        blob = generar_excel(rows, Decimal('141.32'), bank_bytes(), 'Statement', 1, 4, 'ORCHID', {})
        cached = load_workbook(BytesIO(blob), data_only=True)
        lda = cached['Report_LDA']
        cols = {c.value: c.column for c in lda[1]}
        self.assertEqual(lda.cell(2, cols['Del Toro $']).value, 141.32)
        self.assertEqual(lda.cell(2, cols['Franchise $']).value, 141.32)
        self.assertEqual(lda.cell(2, cols['Difference']).value, 0)
        cached.close()

    def test_pivot_groups_by_producer_name_when_there_is_no_producer_code(self):
        """Regresion: CRC GROUP no tiene codigo de agente (solo Vendor Name en producer_name);
        el Pivot no debe fallar con KeyError, debe agrupar por producer_name en ese caso."""
        rows = [dict(producer_name='SAMY INSURANCE, INC', commission_amount='64.85')]
        result = generar_excel(rows, Decimal('64.85'), bank_bytes(), 'Statement', 1, 2, 'CRC GROUP', {})
        cached = load_workbook(BytesIO(result), data_only=True)
        pivot = cached['Pivot']
        self.assertEqual(pivot.cell(4, 1).value, 'SAMY INSURANCE, INC')
        self.assertEqual(pivot.cell(4, 2).value, 64.85)
        cached.close()

    def test_orchid_and_slide_accept_multiple_bank_rows_like_the_other_multi_deposit_carriers(self):
        """Regresion: ORCHID y SLIDE se agregaron al grupo que suma varias transacciones
        bancarias en conciliacion_ui.py, pero generar_excel() todavia solo lo permitia para
        FLORIDA PENINSULA/ASSURANCE/CRC GROUP y rechazaba con 'Seleccion bancaria invalida'."""
        book = Workbook()
        ws = book.active
        ws.title = 'Statement'
        ws.append(['Description', 'Amount'])
        ws.append(['ORCHID ACH 1', 100])
        ws.append(['ORCHID ACH 2', 50])
        raw = BytesIO()
        book.save(raw)
        rows = [dict(producer_code='AGY9416', producer_name='Agency', commission_amount='150')]
        for carrier in ('ORCHID', 'SLIDE'):
            result = generar_excel(rows, Decimal('150'), raw.getvalue(), 'Statement', 1, [2, 3], carrier, {})
            self.assertTrue(result)

    def test_the_general_accepts_multiple_bank_rows(self):
        """Misma regresion que ORCHID/SLIDE: THE GENERAL tambien se agrego al grupo que suma
        varias transacciones bancarias en conciliacion_ui.py (puede llegar en varios depositos
        separados), y generar_excel() debe aceptarlo en vez de 'Seleccion bancaria invalida'."""
        book = Workbook()
        ws = book.active
        ws.title = 'Statement'
        ws.append(['Description', 'Amount'])
        ws.append(['The General ACH 1', 100])
        ws.append(['The General ACH 2', 50])
        raw = BytesIO()
        book.save(raw)
        rows = [dict(producer_code='91479', commission_amount='150')]
        result = generar_excel(rows, Decimal('150'), raw.getvalue(), 'Statement', 1, [2, 3], 'THE GENERAL', {})
        self.assertTrue(result)
        book.close()

    def test_pivot_lists_every_bank_transaction_when_more_than_one_instead_of_only_the_sum(self):
        """Si son varias transacciones bancarias, Pivot no debe mostrar solo la suma: debe
        listar cada una (para ubicar a que fila corresponde una diferencia), y despues el
        total del banco y la diferencia, igual que antes."""
        book = Workbook()
        ws = book.active
        ws.title = 'Statement'
        ws.append(['Description', 'Amount'])
        for name, amount in [('ACH FLORIDA PENINSULA PAYMENT', 10), ('ovation home ACH', 20), ('EDISON deposit', 30)]:
            ws.append([name, amount])
        raw = BytesIO()
        book.save(raw)
        rows = [dict(producer_code='45595', commission_amount='60', comm_percent='.12')]
        result = generar_excel(rows, Decimal('60'), raw.getvalue(), 'Statement', 1, [2, 3, 4], 'FLORIDA PENINSULA', {})
        cached = load_workbook(BytesIO(result), data_only=True)
        pivot = cached['Pivot']
        self.assertEqual(pivot.cell(pivot.max_row - 5, 1).value, 'Grand Total')
        self.assertEqual(pivot.cell(pivot.max_row - 5, 2).value, 60)
        self.assertEqual(pivot.cell(pivot.max_row - 4, 1).value, 'Statement · Fila 2 · ACH FLORIDA PENINSULA PAYMENT')
        self.assertEqual(pivot.cell(pivot.max_row - 4, 2).value, 10)
        self.assertEqual(pivot.cell(pivot.max_row - 3, 1).value, 'Statement · Fila 3 · ovation home ACH')
        self.assertEqual(pivot.cell(pivot.max_row - 3, 2).value, 20)
        self.assertEqual(pivot.cell(pivot.max_row - 2, 1).value, 'Statement · Fila 4 · EDISON deposit')
        self.assertEqual(pivot.cell(pivot.max_row - 2, 2).value, 30)
        self.assertEqual(pivot.cell(pivot.max_row - 1, 1).value, 'Bank')
        self.assertEqual(pivot.cell(pivot.max_row - 1, 2).value, 60)
        self.assertEqual(pivot.cell(pivot.max_row, 1).value, 'Difference')
        self.assertEqual(pivot.cell(pivot.max_row, 2).value, 0)
        cached.close()
        book.close()

    def test_bank_mismatch_does_not_block_and_records_difference_in_pivot(self):
        """Regresion: si la transaccion bancaria elegida no coincide exactamente con el total
        del reporte, ya no se bloquea la generacion; se sigue el flujo normal y la diferencia
        real queda en la fila 'Difference' de Pivot."""
        rows = [dict(producer_code='12IA', commission_amount='64.85')]
        result = generar_excel(rows, Decimal('64.85'), bank_bytes(), 'Statement', 1, 4, 'COMMONWEALTH', {})
        cached = load_workbook(BytesIO(result), data_only=True)
        self.assertEqual(cached['Pivot'].cell(cached['Pivot'].max_row - 1, 2).value, 10)
        self.assertEqual(cached['Pivot'].cell(cached['Pivot'].max_row, 2).value, 54.85)
        cached.close()

    def test_pdf_total(self):
        self.assertEqual(grand_total_pdf('Grand Total DEL TORO FRANCHISING CORP 64.85'), Decimal('64.85'))

    def test_the_general_pivot_shows_bank_amount_and_difference_per_agent_code(self):
        """THE GENERAL agrupa el Pivot por franquicia (no por codigo crudo), mostrando debajo de
        cada una todos sus registros, el subtotal y la transaccion del banco que le corresponde
        (buscada por 'PGA<Agent #>' en las hojas elegidas, incluida Commons para el master). Si
        no se encuentra en el banco, 'Bank' queda vacio para esa franquicia. Al final, las
        transacciones del banco que no corresponden a ningun codigo nuestro se listan aparte,
        solo con el monto, y no se cuentan en el 'Bank'/'Difference' generales."""
        book = Workbook()
        statement = book.active
        statement.title = 'Statement'
        statement.append(['Description', 'Amount'])
        statement.append(['The General DES:COMMISSION ID:1 INDN:SUMI ALL INSURANCE LLC CO ID:SENTRYINS2 CCD PMT INFO:PGA091479\\', 85.46])
        commons = book.create_sheet('Commons')
        commons.append(['Description', 'Amount'])
        commons.append(['The General DES:COMMISSION ID:2 INDN:NOBODY WE KNOW CO ID:SENTRYINS2 CCD PMT INFO:PGA999999\\', 12.34])
        raw = BytesIO()
        book.save(raw)
        book.close()
        rows = [dict(policy_number='FL1', producer_code='91479', premium_amount='1000', commission_amount='85.46',
                     franchise_number='DTF0015', insured_name='Lazaro Ramos', dt_or_dtf='DTF'),
                dict(policy_number='FL2', producer_code='90903', premium_amount='500', commission_amount='40.00',
                     franchise_number='DTF0120', insured_name='Ariel Lastra', dt_or_dtf='DT')]
        result = generar_excel(rows, Decimal('125.46'), raw.getvalue(), 'Statement', 1,
                               [{'sheet': 'Statement', 'row': 2}], 'THE GENERAL', {},
                               bank_sheets=['Statement', 'Commons'])
        cached = load_workbook(BytesIO(result), data_only=True)
        try:
            pivot = cached['Pivot']
            self.assertEqual([c.value for c in pivot[3]][:2], ['Row Labels', 'Sum of Commission'])
            self.assertEqual([c.value for c in pivot[4]], [None, 'Name Insured', 'Policy #', 'Commission'])
            self.assertEqual(pivot.cell(5, 1).value, 'DTF0015 (Código 91479, No master)')
            self.assertEqual([c.value for c in pivot[6]], [None, 'Lazaro Ramos', 'FL1', 85.46])
            self.assertEqual([c.value for c in pivot[7]][:2], ['Total DTF0015 (Código 91479, No master)', 85.46])
            self.assertEqual([c.value for c in pivot[8]][:2], ['Bank', 85.46])
            self.assertEqual(pivot.cell(10, 1).value, 'DTF0120 (Código 90903, Master)')
            self.assertEqual([c.value for c in pivot[11]], [None, 'Ariel Lastra', 'FL2', 40])
            self.assertEqual([c.value for c in pivot[12]][:2], ['Total DTF0120 (Código 90903, Master)', 40])
            self.assertEqual(pivot.cell(13, 1).value, 'Bank')
            self.assertIsNone(pivot.cell(13, 2).value)  # el master (90903) no aparecio en el banco
            self.assertEqual([c.value for c in pivot[15]][:2], ['Grand Total', 125.46])
            self.assertEqual(pivot.cell(16, 1).value, 'PGA999999 en el banco, no corresponde a ningún registro')
            self.assertEqual(pivot.cell(16, 2).value, 12.34)
            self.assertEqual([c.value for c in pivot[17]][:2], ['Bank', 85.46])  # solo lo que corresponde a 91479
            self.assertEqual([c.value for c in pivot[18]][:2], ['Difference', 40])  # 125.46 - 85.46
            bank = cached['Bank']
            self.assertEqual([c.value for c in bank[1]][-2:], ['Franchise', 'Difference'])
            # bank_rows queda ordenado por (hoja, fila): 'Commons' antes que 'Statement'.
            self.assertIn('999999', bank.cell(2, 1).value)
            self.assertIsNone(bank.cell(2, 3).value)
            self.assertIsNone(bank.cell(2, 4).value)
            self.assertIn('091479', bank.cell(3, 1).value)
            self.assertEqual(bank.cell(3, 3).value, 'DTF0015')
            self.assertEqual(bank.cell(3, 4).value, 0)
        finally:
            cached.close()

    def test_the_general_master_code_splits_into_one_block_per_real_franchise(self):
        """Regresion: el codigo master no equivale a una sola franquicia real (cada poliza se
        resuelve por su cuenta en Compass); si sus filas tocan varias oficinas distintas
        (ej. DTF0106, DTF0120, DTF0179), el Pivot debe separarlas en un bloque por franquicia,
        cada una con su propio subtotal, y mostrar un solo 'Bank' agregado para todo el codigo
        (no repetido ni adivinado por bloque, porque el deposito real es uno solo combinado)."""
        book = Workbook()
        statement = book.active
        statement.title = 'Statement'
        statement.append(['Description', 'Amount'])
        statement.append(['The General DES:COMMISSION ID:1 INDN:LAZARO CO ID:SENTRYINS2 CCD PMT INFO:PGA091479\\', 10])
        commons = book.create_sheet('Commons')
        commons.append(['Description', 'Amount'])
        commons.append(['THE GENERAL MASTER PAYOUT ACH', 25.19])
        raw = BytesIO()
        book.save(raw)
        book.close()
        rows = [dict(policy_number='FL1', producer_code='91479', commission_amount='10', premium_amount='100',
                     franchise_number='DTF0015', insured_name='Lazaro', dt_or_dtf='DTF'),
                dict(policy_number='FL2', producer_code='90903', commission_amount='6.25', premium_amount='62.5',
                     franchise_number='DTF0106', insured_name='Martiny', dt_or_dtf='DT'),
                dict(policy_number='FL3', producer_code='90903', commission_amount='9.59', premium_amount='95.9',
                     franchise_number='DTF0120', insured_name='Eveling', dt_or_dtf='DT'),
                dict(policy_number='FL4', producer_code='90903', commission_amount='9.35', premium_amount='93.5',
                     franchise_number='DTF0179', insured_name='Dina', dt_or_dtf='DT')]
        result = generar_excel(rows, Decimal('35.19'), raw.getvalue(), 'Statement', 1,
                               [{'sheet': 'Statement', 'row': 2}], 'THE GENERAL', {},
                               bank_sheets=['Statement', 'Commons'])
        cached = load_workbook(BytesIO(result), data_only=True)
        try:
            pivot = cached['Pivot']
            self.assertEqual(pivot.cell(5, 1).value, 'DTF0015 (Código 91479, No master)')
            self.assertEqual([c.value for c in pivot[8]][:2], ['Bank', 10])
            self.assertEqual(pivot.cell(10, 1).value, 'DTF0106 (Código 90903, Master)')
            self.assertEqual([c.value for c in pivot[11]], [None, 'Martiny', 'FL2', 6.25])
            self.assertEqual([c.value for c in pivot[12]][:2], ['Total DTF0106 (Código 90903, Master)', 6.25])
            self.assertEqual(pivot.cell(14, 1).value, 'DTF0120 (Código 90903, Master)')
            self.assertEqual([c.value for c in pivot[16]][:2], ['Total DTF0120 (Código 90903, Master)', 9.59])
            self.assertEqual(pivot.cell(18, 1).value, 'DTF0179 (Código 90903, Master)')
            self.assertEqual([c.value for c in pivot[20]][:2], ['Total DTF0179 (Código 90903, Master)', 9.35])
            # No hay 'Bank' individual entre los bloques de 90903 (filas 13 y 17 son blancas).
            self.assertIsNone(pivot.cell(13, 1).value)
            self.assertIsNone(pivot.cell(17, 1).value)
            self.assertEqual([c.value for c in pivot[22]][:2], ['Bank (Código 90903)', 25.19])
            self.assertEqual([c.value for c in pivot[24]][:2], ['Grand Total', 35.19])
            self.assertEqual([c.value for c in pivot[25]][:2], ['Bank', 35.19])
            self.assertEqual([c.value for c in pivot[26]][:2], ['Difference', 0])
            bank = cached['Bank']
            self.assertIn('MASTER PAYOUT', bank.cell(2, 1).value)
            self.assertEqual(bank.cell(2, 3).value, 'Varias (DTF0106, DTF0120, DTF0179)')
        finally:
            cached.close()

    def test_the_general_master_code_transaction_without_pga_gets_attributed_by_name(self):
        """El codigo master no deposita con 'PGA<codigo>' (aparece con otro texto, sin ningun
        codigo). Si hay un solo codigo master en el statement, esa transaccion (encontrada por
        el nombre del carrier) se atribuye a el: aparece en el Pivot de esa franquicia y en la
        hoja Bank, en vez de perderse."""
        book = Workbook()
        statement = book.active
        statement.title = 'Statement'
        statement.append(['Description', 'Amount'])
        statement.append(['The General DES:COMMISSION ID:1 INDN:SUMI ALL INSURANCE LLC CO ID:SENTRYINS2 CCD PMT INFO:PGA091479\\', 85.46])
        commons = book.create_sheet('Commons')
        commons.append(['Description', 'Amount'])
        commons.append(['THE GENERAL MASTER PAYOUT ACH', 40.00])
        raw = BytesIO()
        book.save(raw)
        book.close()
        rows = [dict(policy_number='FL1', producer_code='91479', premium_amount='1000', commission_amount='85.46',
                     franchise_number='DTF0015', insured_name='Lazaro Ramos', dt_or_dtf='DTF'),
                dict(policy_number='FL2', producer_code='90903', premium_amount='500', commission_amount='40.00',
                     franchise_number='DTF0120', insured_name='Ariel Lastra', dt_or_dtf='DT')]
        result = generar_excel(rows, Decimal('125.46'), raw.getvalue(), 'Statement', 1,
                               [{'sheet': 'Statement', 'row': 2}], 'THE GENERAL', {},
                               bank_sheets=['Statement', 'Commons'])
        cached = load_workbook(BytesIO(result), data_only=True)
        try:
            pivot = cached['Pivot']
            self.assertEqual([c.value for c in pivot[13]][:2], ['Bank', 40])  # master ya no queda vacio
            self.assertEqual([c.value for c in pivot[15]][:2], ['Grand Total', 125.46])
            self.assertEqual([c.value for c in pivot[16]][:2], ['Bank', 125.46])  # 85.46 + 40 del master
            self.assertEqual([c.value for c in pivot[17]][:2], ['Difference', 0])
            bank = cached['Bank']
            self.assertEqual(bank.max_row, 3)
            self.assertIn('MASTER PAYOUT', bank.cell(2, 1).value)
            self.assertEqual(bank.cell(2, 3).value, 'DTF0120')
            self.assertEqual(bank.cell(2, 4).value, 0)
        finally:
            cached.close()

    def test_the_general_pivot_has_no_detail_rows_when_bank_omitted(self):
        """Sin banco (omitir_banco), Pivot se sigue agrupando por franquicia y mostrando el
        detalle, pero sin filas 'Bank' con monto ni la lista de transacciones sin corresponder."""
        rows = [dict(policy_number='FL1', producer_code='91479', premium_amount='1000', commission_amount='85.46',
                     franchise_number='DTF0015', insured_name='Lazaro Ramos')]
        result = generar_excel(rows, Decimal('85.46'), None, None, None, None, 'THE GENERAL', {})
        cached = load_workbook(BytesIO(result), data_only=True)
        try:
            pivot = cached['Pivot']
            self.assertEqual(pivot.cell(5, 1).value, 'DTF0015 (Código 91479, No master)')
            self.assertEqual([c.value for c in pivot[7]][:2], ['Total DTF0015 (Código 91479, No master)', 85.46])
            self.assertEqual([c.value for c in pivot[8]][:2], ['Bank', None])
            self.assertEqual([c.value for c in pivot[10]][:2], ['Grand Total', 85.46])
            self.assertEqual([c.value for c in pivot[11]][:2], ['Bank', None])
            self.assertEqual([c.value for c in pivot[12]][:2], ['Difference', None])
            self.assertEqual(cached['Bank'].max_row, 1)  # hoja Bank vacia, solo la fila creada por defecto
        finally:
            cached.close()


if __name__ == '__main__':
    unittest.main()
