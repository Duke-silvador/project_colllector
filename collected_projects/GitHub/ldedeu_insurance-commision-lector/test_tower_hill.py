import unittest
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch
from openpyxl import Workbook
from requests import RequestException
from carriers import CARRIERS
from tower_hill import (preparar_statement, detectar_encabezado, porcentaje_fraccion,
                        extraer_tabla_pagina, alertas_lectura, validar, construir_filas,
                        completar_franquicias, guardar, cargar_comisiones_excel, FIELDS)


def libro(policy='W014240270', rate=0.10, written_premium=4922.1, premium=4410.0, commission=441.0,
          term='REN', form='H3', producing='Maria Orduz', submitted='Camila Orduz', insured='Tafur Samir Walid'):
    wb = Workbook()
    ws = wb.active
    ws.append(['Policy Number', 'Effective Date', 'Named Insured', 'Form', 'Term', 'Written Premium',
               'Commission Premium', 'Rate', 'Commission Amount', 'Producing Agent', 'Submitted Agent'])
    ws.append([policy, datetime(2026, 8, 5), insured, form, term, written_premium, premium, rate,
              commission, producing, submitted])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def oficina(**over):
    base = dict(office_id='abc', office_number='120', office_name='Del Toro Insurance Agency Inc', state='FL')
    base.update(over)
    return base


class TowerHillExcelTests(unittest.TestCase):
    def test_carrier_registered_for_raw_commission_loading(self):
        self.assertEqual(CARRIERS['TOWER HILL']['table'], 'staging_hub.st_tower_hill_raw')

    def test_header_detection_and_parsing_uses_commission_premium_not_written_premium(self):
        """Rate = Commission Amount / Commission Premium (no Written Premium, que es una
        cifra distinta): 441/4410 = 0.10, no 441/4922.1."""
        rows = preparar_statement(libro(), '2026-08')
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row['policy_number'], 'W014240270')
        self.assertEqual(row['insured_name'], 'Tafur Samir Walid')
        self.assertEqual(Decimal(row['written_premium']), Decimal('4922.1'))
        self.assertEqual(Decimal(row['premium_amount']), Decimal('4410'))
        self.assertEqual(Decimal(row['commission_amount']), Decimal('441'))
        self.assertEqual(Decimal(row['del_toro_percent']), Decimal('0.1'))
        self.assertEqual(row['del_toro_percent_source'], 'statement')
        self.assertEqual(row['transaction_type'], 'REN')
        self.assertEqual(row['transaction_type_source'], 'statement')
        self.assertEqual(row['form'], 'H3')
        self.assertEqual(row['producing_agent'], 'Maria Orduz')
        self.assertEqual(row['submitted_agent'], 'Camila Orduz')
        self.assertEqual(row['effective_date'], date(2026, 8, 5))

    def test_no_agency_code_or_office_columns_are_required(self):
        """El statement real de TOWER HILL no trae Code ni Office; preparar_statement no debe
        exigirlos."""
        rows = preparar_statement(libro(), '2026-08')
        self.assertNotIn('producer_code', rows[0])
        self.assertNotIn('office_raw', rows[0])

    def test_rate_accepts_fraction_or_whole_percentage(self):
        self.assertEqual(porcentaje_fraccion('0.10'), Decimal('0.10'))
        self.assertEqual(porcentaje_fraccion('10'), Decimal('0.10'))
        self.assertEqual(porcentaje_fraccion('10%'), Decimal('0.10'))
        with self.assertRaises(ValueError):
            porcentaje_fraccion('abc')

    def test_policy_number_na_becomes_blank(self):
        rows = preparar_statement(libro(policy='N/A'), '2026-08')
        self.assertEqual(rows[0]['policy_number'], '')

    def test_missing_required_column_is_rejected(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['Policy Number', 'Effective Date', 'Named Insured', 'Term', 'Written Premium',
                   'Rate', 'Commission Amount'])
        ws.append(['W1', datetime(2026, 8, 1), 'JOHN DOE', 'REN', 1000.0, 0.10, 100.0])
        buf = BytesIO()
        wb.save(buf)
        with self.assertRaises(ValueError):
            preparar_statement(buf.getvalue(), '2026-08')

    def test_detectar_encabezado_finds_header_row_below_notes(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['Some note before the real header'])
        ws.append(['Policy Number', 'Effective Date', 'Named Insured', 'Term', 'Written Premium',
                   'Commission Premium', 'Rate', 'Commission Amount'])
        buf = BytesIO()
        wb.save(buf)
        book_bytes = buf.getvalue()
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(book_bytes), read_only=True, data_only=True)
        try:
            self.assertEqual(detectar_encabezado(book.worksheets[0]), 2)
        finally:
            book.close()


class TowerHillFranquiciaTests(unittest.TestCase):
    """_franquicia_historica hace su propia consulta directa a historic_data_commissions
    (para traer tambien term_length), asi que el respaldo por historico ya no se mockea via
    franquicias.buscar_franquicia_historica: se simula con cursor.fetchone."""

    def test_policy_found_in_compass_resolves_directly(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        rows = [dict(policy_number='W014240270')]
        with patch('tower_hill.compass.obtener_token', return_value='tok'), \
             patch('tower_hill.compass.buscar_poliza',
                   return_value={'office_id': 'abc', 'policy_id': 'p1', 'status_id': 'active',
                                'effective_date': '2026-08-05', 'expiration_date': '2027-08-05'}) as buscar:
            completar_franquicias(connection, rows)
        buscar.assert_called_once_with('tok', 'W014240270')
        cursor.fetchone.assert_not_called()  # No hizo falta caer al historico.
        self.assertEqual(rows[0]['franchise_number'], 'DTF0120')
        self.assertEqual(rows[0]['franchise_number_source'], 'compass')
        self.assertEqual(rows[0]['state'], 'FL')
        self.assertEqual(rows[0]['compass_policy_id'], 'p1')
        self.assertEqual(rows[0]['policy_status'], 'active')
        self.assertEqual(rows[0]['term_length'], 12)
        self.assertEqual(rows[0]['producer_name'], 'Del Toro Insurance Agency Inc')
        self.assertIsNone(rows[0]['code_lookup_alert'])

    def test_policy_not_in_compass_falls_back_to_historic(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = {'franchise': 'DTF0120', 'term_length': 12}
        rows = [dict(policy_number='W014240270')]
        with patch('tower_hill.compass.obtener_token', return_value='tok'), \
             patch('tower_hill.compass.buscar_poliza', return_value=None):
            completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['franchise_number'], 'DTF0120')
        self.assertEqual(rows[0]['franchise_number_source'], 'historico')
        self.assertEqual(rows[0]['term_length'], 12)
        self.assertEqual(rows[0]['producer_name'], 'Del Toro Insurance Agency Inc')
        self.assertIsNone(rows[0]['code_lookup_alert'])

    def test_compass_office_not_translatable_falls_back_to_historic(self):
        """Compass encuentra la poliza pero el office_id no traduce a ninguna oficina real."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = {'franchise': 'DTF0120', 'term_length': 12}
        rows = [dict(policy_number='W014240270')]
        with patch('tower_hill.compass.obtener_token', return_value='tok'), \
             patch('tower_hill.compass.buscar_poliza', return_value={'office_id': 'no-existe'}):
            completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['franchise_number'], 'DTF0120')
        self.assertEqual(rows[0]['franchise_number_source'], 'historico')

    def test_compass_error_falls_back_to_historic(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = {'franchise': 'DTF0120', 'term_length': 12}
        rows = [dict(policy_number='W014240270')]
        with patch('tower_hill.compass.obtener_token', side_effect=RequestException('sin red')):
            completar_franquicias(connection, rows)
        self.assertEqual(rows[0]['franchise_number'], 'DTF0120')
        self.assertEqual(rows[0]['franchise_number_source'], 'historico')

    def test_policy_not_in_compass_nor_historic_is_flagged_not_guessed(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = None
        rows = [dict(policy_number='W999999999')]
        with patch('tower_hill.compass.obtener_token', return_value='tok'), \
             patch('tower_hill.compass.buscar_poliza', return_value=None):
            completar_franquicias(connection, rows)
        self.assertIsNone(rows[0]['franchise_number'])
        self.assertTrue(rows[0]['code_lookup_alert'])

    def test_missing_policy_number_is_flagged_without_querying_compass_or_historic(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        rows = [dict(policy_number='')]
        with patch('tower_hill.compass.obtener_token') as token:
            completar_franquicias(connection, rows)
        token.assert_not_called()
        cursor.fetchone.assert_not_called()
        self.assertIsNone(rows[0]['franchise_number'])
        self.assertTrue(rows[0]['code_lookup_alert'])

    def test_client_name_search_resolves_when_policy_and_historic_both_fail(self):
        """Ultimo respaldo: si la poliza no aparece por numero ni en Compass ni en el
        historico, se busca al cliente por nombre en Company Search (busqueda visual)."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = None
        rows = [dict(policy_number='W030725262', insured_name='Olaves Gabriela')]
        with patch('tower_hill.compass.obtener_token', return_value='tok'), \
             patch('tower_hill.compass.buscar_poliza', return_value=None), \
             patch('tower_hill.buscar_franquicias_por_nombre', return_value=[
                 {'request_id': '0', 'client_name': 'Olaves Gabriela',
                  'offices': [{'label': 'Samy Insurance Inc. (120)', 'office_number': '120'}],
                  'matched_name': 'Gabriela Olaves', 'error': None}]) as buscar:
            completar_franquicias(connection, rows)
        buscar.assert_called_once_with([{'request_id': '0', 'client_name': 'Olaves Gabriela'}])
        self.assertEqual(rows[0]['franchise_number'], 'DTF0120')
        self.assertEqual(rows[0]['franchise_number_source'], 'compass_visual')
        self.assertEqual(rows[0]['state'], 'FL')
        self.assertEqual(rows[0]['producer_name'], 'Del Toro Insurance Agency Inc')
        self.assertIsNone(rows[0]['code_lookup_alert'])

    def test_client_name_search_with_multiple_offices_is_flagged_not_guessed(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = None
        rows = [dict(policy_number='W030725262', insured_name='Olaves Gabriela')]
        with patch('tower_hill.compass.obtener_token', return_value='tok'), \
             patch('tower_hill.compass.buscar_poliza', return_value=None), \
             patch('tower_hill.buscar_franquicias_por_nombre', return_value=[
                 {'request_id': '0', 'client_name': 'Olaves Gabriela',
                  'offices': [{'label': 'A (120)', 'office_number': '120'}, {'label': 'B (73)', 'office_number': '73'}],
                  'matched_name': 'Gabriela Olaves', 'error': None}]):
            completar_franquicias(connection, rows)
        self.assertIsNone(rows[0]['franchise_number'])
        self.assertIn('varias franquicias', rows[0]['code_lookup_alert'])

    def test_client_name_search_unavailable_keeps_previous_alert(self):
        """Mejor esfuerzo: si el bot no esta disponible, no bloquea el resto del proceso."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = None
        rows = [dict(policy_number='W030725262', insured_name='Olaves Gabriela')]
        with patch('tower_hill.compass.obtener_token', return_value='tok'), \
             patch('tower_hill.compass.buscar_poliza', return_value=None), \
             patch('tower_hill.buscar_franquicias_por_nombre', side_effect=ValueError('bot no configurado')):
            completar_franquicias(connection, rows)
        self.assertIsNone(rows[0]['franchise_number'])
        self.assertTrue(rows[0]['code_lookup_alert'])

    def test_rows_without_insured_name_skip_the_client_name_search(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.return_value = [oficina()]
        cursor.fetchone.return_value = None
        rows = [dict(policy_number='W999999999', insured_name='')]
        with patch('tower_hill.compass.obtener_token', return_value='tok'), \
             patch('tower_hill.compass.buscar_poliza', return_value=None), \
             patch('tower_hill.buscar_franquicias_por_nombre') as buscar:
            completar_franquicias(connection, rows)
        buscar.assert_not_called()
        self.assertIsNone(rows[0]['franchise_number'])


class TowerHillGuardarTests(unittest.TestCase):
    def rows(self):
        return preparar_statement(libro(), '2026-08')

    def test_save_inserts_new_row_and_resolves_franchise(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.side_effect = [{'acquired': 1}, {'franchise': 'DTF0120', 'term_length': 12}, None]
        cursor.fetchall.side_effect = [[], [oficina()]]
        with patch('tower_hill.conectar', return_value=connection), \
             patch('tower_hill.compass.obtener_token', side_effect=RequestException('sin red')):
            self.assertEqual(guardar(self.rows(), '2026-08', 'TowerHill.xlsx', b'contenido-th', 1, None), (1, 0))
            insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
            self.assertEqual(insert.args[1][15], 'DTF0120')  # franchise_number resuelto por historico

    def test_retry_with_same_content_does_not_duplicate(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.side_effect = [{'acquired': 1}, {'franchise': 'DTF0120', 'term_length': 12}, None]
        cursor.fetchall.side_effect = [[], [oficina()]]
        with patch('tower_hill.conectar', return_value=connection), \
             patch('tower_hill.compass.obtener_token', side_effect=RequestException('sin red')):
            rows = self.rows()
            guardar(rows, '2026-08', 'TowerHill.xlsx', b'contenido-th', 1, None)
            insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
            saved_payload = insert.args[1][19]
            source_row = rows[0]['source_row']

            cursor.execute.reset_mock()
            cursor.fetchone.side_effect = [{'acquired': 1}, None]
            cursor.fetchall.side_effect = [[{'id': 1, 'source_row': source_row, 'source_data': saved_payload}]]
            self.assertEqual(guardar(rows, '2026-08', 'TowerHill.xlsx', b'contenido-th', 1, None), (0, 0))

    def test_changed_row_updates_instead_of_duplicating(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.side_effect = [{'acquired': 1}, {'franchise': 'DTF0120', 'term_length': 12}, None]
        cursor.fetchall.side_effect = [[], [oficina()]]
        with patch('tower_hill.conectar', return_value=connection), \
             patch('tower_hill.compass.obtener_token', side_effect=RequestException('sin red')):
            rows = self.rows()
            guardar(rows, '2026-08', 'TowerHill.xlsx', b'contenido-th', 1, None)
            insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
            saved_payload = insert.args[1][19]
            source_row = rows[0]['source_row']

            cursor.execute.reset_mock()
            cursor.fetchone.side_effect = [{'acquired': 1}, {'franchise': 'DTF0120', 'term_length': 12}, None]
            cursor.fetchall.side_effect = [[{'id': 1, 'source_row': source_row, 'source_data': saved_payload}], [oficina()]]
            changed = preparar_statement(libro(insured='OTHER NAME'), '2026-08')
            self.assertEqual(guardar(changed, '2026-08', 'TowerHill.xlsx', b'contenido-th', 1, None), (0, 1))
            update = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('UPDATE'))
            self.assertEqual(update.args[1][1], 'OTHER NAME')

    def test_save_uses_content_hash_as_file_id_not_the_file_name(self):
        """Regresion: el mismo statement subido con otro nombre de archivo debe reconocerse
        como el mismo import (mismo file_id, calculado del contenido), no como uno nuevo."""
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchone.side_effect = [{'acquired': 1}, {'franchise': 'DTF0120', 'term_length': 12}, None]
        cursor.fetchall.side_effect = [[], [oficina()]]
        with patch('tower_hill.conectar', return_value=connection), \
             patch('tower_hill.compass.obtener_token', side_effect=RequestException('sin red')):
            rows = self.rows()
            guardar(rows, '2026-08', 'nombre-original.xlsx', b'mismo-contenido', 1, None)
            insert = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('INSERT'))
            file_id_guardado = insert.args[1][-5]
            saved_payload = insert.args[1][19]
            source_row = rows[0]['source_row']

            cursor.execute.reset_mock()
            cursor.fetchone.side_effect = [{'acquired': 1}, None]
            cursor.fetchall.side_effect = [[{'id': 1, 'source_row': source_row, 'source_data': saved_payload}]]
            self.assertEqual(
                guardar(rows, '2026-08', 'nombre-reenviado (1).xlsx', b'mismo-contenido', 1, None), (0, 0))
            select = next(c for c in cursor.execute.call_args_list if c.args[0].startswith('SELECT id'))
            self.assertEqual(select.args[1][0], file_id_guardado)

    def test_save_rejects_invalid_month(self):
        with self.assertRaises(ValueError):
            guardar(self.rows(), '2026-13', 'TowerHill.xlsx', b'contenido-th', 1, None)

    def test_save_rejects_empty_rows(self):
        with self.assertRaises(ValueError):
            guardar([], '2026-08', 'TowerHill.xlsx', b'contenido-th', 1, None)


class TowerHillCommissionSplitTests(unittest.TestCase):
    def commission_row(self, **overrides):
        return dict({'id': 1, 'policy_number': 'W014240270', 'insured_name': 'X', 'effective_date': date(2026, 8, 5),
                     'transaction_type': 'REN', 'franchise_number': 'DTF0120', 'form': 'H3',
                     'producing_agent': 'Maria Orduz', 'submitted_agent': 'Camila Orduz',
                     'accounting_month': date(2026, 8, 1), 'written_premium': Decimal('4922.10'),
                     'premium_amount': Decimal('4410.00'), 'del_toro_percent': Decimal('0.1000'),
                     'commission_amount': Decimal('441.00'), 'state': 'FL'}, **overrides)

    def test_franchise_percent_comes_from_commission_rates_matching_del_toro_rate(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [
            [self.commission_row()],
            [dict(state='FL', carrier='TOWER HILL', transaction_type='ALL', business_line='',
                  del_toro_percent=Decimal('0.1000'), franchise_percent=Decimal('0.0800'))]]
        snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 8, 1)}]}
        result = cargar_comisiones_excel(connection, snapshot)[0]
        self.assertEqual(result['franchise_percent'], Decimal('0.0800'))
        self.assertEqual(result['franchise_commission'], Decimal('352.8000'))
        self.assertIsNone(result['commission_alert'])
        self.assertEqual(result['agent_name'], 'Maria Orduz')
        self.assertEqual(result['line_business_id'], 'H3')
        self.assertEqual(result['carrier'], 'TOWER HILL')

    def test_unresolved_franchise_is_alerted(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value
        cursor.fetchall.side_effect = [[self.commission_row(franchise_number=None)], []]
        snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 8, 1)}]}
        result = cargar_comisiones_excel(connection, snapshot)[0]
        self.assertTrue(result['commission_alert'])

    def test_no_rows_raises(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.return_value = []
        snapshot = {'file_id': None, 'rows': [{'accounting_month': date(2026, 8, 1)}]}
        with self.assertRaises(ValueError):
            cargar_comisiones_excel(connection, snapshot)


class TowerHillPdfExtractionTests(unittest.TestCase):
    def _blocks(self):
        headers = [('Policy', .02), ('Number', .04), ('Effective', .12), ('Date', .14), ('Named', .22),
                   ('Insured', .24), ('Term', .32), ('Written', .40), ('Premium', .42), ('Commission', .50),
                   ('Premium', .52), ('Rate', .60), ('Commission', .68), ('Amount', .70)]
        blocks = [dict(text=text, x=x, y=.1, height=.01, confidence=.99) for text, x in headers]
        values = [('W014240270', .03, .2), ('8/5/2026', .13, .2), ('JOHN', .22, .2), ('DOE', .24, .2),
                  ('REN', .32, .2), ('4922.10', .41, .2), ('4410.00', .51, .2), ('0.10', .60, .2),
                  ('441.00', .69, .2)]
        blocks.extend(dict(text=text, x=x, y=y, height=.01, confidence=.97) for text, x, y in values)
        return blocks

    def test_extraer_tabla_pagina_reads_columns_by_header_position(self):
        filas = extraer_tabla_pagina(self._blocks())
        self.assertEqual(len(filas), 1)
        fila = filas[0]
        self.assertEqual(fila['policy_number'], 'W014240270')
        self.assertEqual(fila['transaction_type'], 'REN')
        self.assertEqual(fila['premium_amount'], '4410.00')
        self.assertEqual(fila['written_premium'], '4922.10')
        self.assertIn('JOHN', fila['insured_name'])
        self.assertIn('DOE', fila['insured_name'])

    def test_no_header_returns_no_rows(self):
        self.assertEqual(extraer_tabla_pagina([dict(text='random', x=.1, y=.1, height=.01)]), [])

    def test_alertas_lectura_flags_missing_and_invalid_fields(self):
        row = dict.fromkeys(FIELDS, '')
        row.update(policy_number='W1', insured_name='X', transaction_type='REN',
                   effective_date='2/30/2026', written_premium='1000', premium_amount='NaN',
                   commission_amount='100', del_toro_percent='0.10')
        alerts = alertas_lectura(row)
        self.assertIn('effective_date', alerts)
        self.assertIn('premium_amount', alerts)
        self.assertNotIn('form', alerts)  # informativo, puede faltar
        self.assertNotIn('producing_agent', alerts)

    def test_validar_and_construir_filas_produce_canonical_rows(self):
        row = dict(policy_number='W014240270', insured_name='JOHN DOE', transaction_type='REN',
                   effective_date='8/5/2026', written_premium='4922.10', premium_amount='4410.00',
                   commission_amount='441.00', del_toro_percent='0.10', form='H3',
                   producing_agent='Maria Orduz', submitted_agent='Camila Orduz')
        validar([row], '2026-08')
        filas = construir_filas([row], '2026-08')
        self.assertEqual(filas[0]['policy_number'], 'W014240270')
        self.assertEqual(filas[0]['effective_date'], date(2026, 8, 5))
        self.assertEqual(filas[0]['premium_amount'], '4410.00')
        self.assertEqual(Decimal(filas[0]['del_toro_percent']), Decimal('0.10'))

    def test_validar_rejects_missing_required_field(self):
        row = dict(policy_number='', insured_name='JOHN DOE', transaction_type='REN',
                   effective_date='8/5/2026', written_premium='4922.10', premium_amount='4410.00',
                   commission_amount='441.00', del_toro_percent='0.10')
        with self.assertRaises(ValueError):
            validar([row], '2026-08')
